"""GUI 界面 — PyQt5 桌面应用"""
import sys
import os
import asyncio
import subprocess
import threading
import re
import shutil
from datetime import datetime
from pathlib import Path
from typing import Dict, List, Set

from PyQt5.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QPushButton, QLabel, QDateEdit, QFileDialog, QTextEdit,
    QProgressBar, QMessageBox, QGroupBox, QFrame, QTableWidget,
    QTableWidgetItem, QTabWidget, QSplitter, QStatusBar,
    QLineEdit, QFormLayout, QRadioButton, QButtonGroup, QCheckBox, QComboBox,
    QInputDialog
)
from PyQt5.QtCore import QDate, QThread, pyqtSignal, Qt, QTimer, QLockFile
from PyQt5.QtGui import QFont, QColor

import yaml
import json

from utils.logger import setup_logger, GuiLogHandler
from utils.fuzzy_match import normalize_text
from modules.llm_intent import LLMIntentClient
from modules.project_normalizer import country_of_project
from utils.runtime_paths import (
    APP_ROOT,
    app_relative_path,
    copy_imported_file,
    copy_reference_file,
    resolve_runtime_path,
)

APP_DIR = str(APP_ROOT)
SESSION_FILE = os.path.join(APP_DIR, "session_state.json")
INTERNAL_EMAIL_CACHE_FILE = os.path.join(
    APP_DIR, "data", "internal_email_cache.xlsx"
)
WORKORDER_RETRY_QUEUE_FILE = os.path.join(
    APP_DIR, "storage", "workorder_retry_queue.json"
)


def _stage1_cache_root() -> Path:
    """阶段一历史缓存按邮件日期范围分目录保存。"""
    return Path(APP_DIR) / "output" / "stage1_email" / "cache"


def _cache_date_text(value) -> str:
    """把 GUI/线程传入的日期统一成可用于文件夹名的日期文本。"""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    text = str(value or "").strip()
    if len(text) >= 10:
        candidate = text[:10].replace("/", "-")
        if re.fullmatch(r"\d{4}-\d{2}-\d{2}", candidate):
            return candidate
    return ""


def _write_imap_summary(imap_total, date_from=None, date_to=None, loaded_total=None):
    """把最近一次 IMAP SEARCH 数量写入工作台可读取的脱敏摘要。"""
    try:
        total = int(imap_total or 0)
    except (TypeError, ValueError):
        total = 0
    output_dir = Path(APP_DIR) / "output" / "stage1_email"
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = {
        "schema": 1,
        "imap_read_total": max(0, total),
        "parsed_total": max(0, int(loaded_total or 0)),
        "date_from": _cache_date_text(date_from),
        "date_to": _cache_date_text(date_to),
        "created_at": datetime.now().isoformat(timespec="seconds"),
    }
    path = output_dir / "imap_summary.json"
    temp = path.with_suffix(".json.tmp")
    temp.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(path)
    return path


def _archive_stage1_cache(primary, review, filtered, date_from=None, date_to=None, imap_total=0):
    """保存一次阶段一结果，供操作人员按日期重新读取。

    稳定输出文件仍写在 ``output/stage1_email``，这里额外复制到日期范围目录，
    不会改变原有阶段二输入路径和工作台数据库。
    """
    primary_path = Path(str(primary or ""))
    if not primary_path.is_file():
        return None
    start = _cache_date_text(date_from) or "unknown-start"
    end = _cache_date_text(date_to) or start
    folder = _stage1_cache_root() / f"{start}_to_{end}"
    folder.mkdir(parents=True, exist_ok=True)
    files = {}
    for source, name in (
        (primary, "to_workorder_list.xlsx"),
        (review, "to_review_list.xlsx"),
        (filtered, "filtered_mail_record.xlsx"),
    ):
        source_path = Path(str(source or ""))
        if source_path.is_file():
            target = folder / name
            shutil.copy2(source_path, target)
            files[name] = name
    manifest = {
        "schema": 1,
        "date_from": start,
        "date_to": end,
        "imap_read_total": max(0, int(imap_total or 0)),
        "created_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "files": files,
    }
    with (folder / "manifest.json").open("w", encoding="utf-8") as handle:
        json.dump(manifest, handle, ensure_ascii=False, indent=2)
    return str(folder)


def _stage1_cache_options():
    """返回可供 GUI 读取的日期缓存，按结束日期和生成时间倒序。"""
    root = _stage1_cache_root()
    options = []
    if root.is_dir():
        for folder in root.iterdir():
            if not folder.is_dir():
                continue
            manifest_path = folder / "manifest.json"
            try:
                manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
            except (OSError, ValueError, TypeError):
                continue
            files = manifest.get("files") if isinstance(manifest, dict) else {}
            if not isinstance(files, dict):
                files = {}
            primary = folder / str(files.get("to_workorder_list.xlsx", "to_workorder_list.xlsx"))
            review = folder / str(files.get("to_review_list.xlsx", "to_review_list.xlsx"))
            filtered = folder / str(files.get("filtered_mail_record.xlsx", "filtered_mail_record.xlsx"))
            if not primary.is_file() and not review.is_file():
                continue
            date_from = str(manifest.get("date_from", "") or "")
            date_to = str(manifest.get("date_to", "") or date_from)
            created_at = str(manifest.get("created_at", "") or "")
            label = f"{date_from} 至 {date_to}"
            if created_at:
                label += f"（缓存于 {created_at}）"
            options.append({
                "label": label,
                "folder": str(folder),
                "primary": str(primary) if primary.is_file() else "",
                "review": str(review) if review.is_file() else "",
                "filtered": str(filtered) if filtered.is_file() else "",
                "date_from": date_from,
                "date_to": date_to,
            })
    options.sort(key=lambda item: (item.get("date_to", ""), item.get("label", "")), reverse=True)
    # 兼容升级前已经存在的稳定输出：新运行会自动归档，旧结果仍可先读取。
    if not options:
        for stage1_dir in (
            Path(APP_DIR) / "output" / "stage1_email",
            Path(APP_DIR) / "output",
        ):
            legacy_primary = stage1_dir / "to_workorder_list.xlsx"
            legacy_review = stage1_dir / "to_review_list.xlsx"
            if not (legacy_primary.is_file() or legacy_review.is_file()):
                continue
            legacy_filtered = stage1_dir / "filtered_mail_record.xlsx"
            options.append({
                "label": "当前最新阶段一结果（旧版未按日期归档）",
                "folder": str(stage1_dir),
                "primary": str(legacy_primary) if legacy_primary.is_file() else "",
                "review": str(legacy_review) if legacy_review.is_file() else "",
                "filtered": str(legacy_filtered) if legacy_filtered.is_file() else "",
                "date_from": "",
                "date_to": "",
            })
            break
    return options


def _workorder_retry_identity(row):
    """与工作台服务端一致的邮件+业务明细身份，用于自动重查回流。"""
    def text(value):
        return "" if value is None else str(value).strip()

    sender = text(row.get("发件人邮箱") or row.get("sender_email") or row.get("sender")).lower()
    date = text(row.get("发件日期") or row.get("date")).replace("T", " ")
    subject = text(row.get("邮件主题") or row.get("subject") or row.get("主题"))
    company = text(row.get("客户公司名称") or row.get("客户") or row.get("company"))
    project = text(row.get("标准化项目名称") or row.get("项目") or row.get("program"))
    request = text(row.get("需求") or row.get("request"))
    return "\x1f".join((sender, date, subject, company, project, request))


def _apply_workorder_retry_queue(rows, logger=None):
    """给待自动重查的阶段二行打标，并返回本轮命中的队列键。"""
    try:
        with open(WORKORDER_RETRY_QUEUE_FILE, "r", encoding="utf-8") as fh:
            payload = json.load(fh)
    except (FileNotFoundError, OSError, ValueError):
        return set()
    items = payload.get("items") if isinstance(payload, dict) else []
    pending = {
        str(item.get("key")): item
        for item in items or []
        if isinstance(item, dict) and item.get("status", "pending") == "pending"
    }
    matched = set()
    for row in rows or []:
        key = _workorder_retry_identity(row)
        if key in pending:
            row["_force_workorder_query"] = True
            row["_retry_queue_key"] = key
            matched.add(key)
    if matched and logger:
        logger.info(f"漏单自动重查队列命中 {len(matched)} 条，本轮将跳过旧缓存")
    return matched


def _filter_workbench_confirmed_rows(headers, rows, input_path):
    """工作台导出只把已确认明细交给阶段二；旧阶段一文件保持原口径。"""
    header_names = {str(header or "").strip() for header in headers or []}
    is_workbench_export = (
        "工作台状态" in header_names
        or os.path.basename(str(input_path or "")).lower().startswith("workbench_reviewed")
    )
    if not is_workbench_export or "工作台状态" not in header_names:
        return list(rows or []), is_workbench_export, False
    confirmed = {
        "confirmed", "已完成", "确认", "人工确认",
    }
    selected = [
        row for row in rows or []
        if str(row.get("工作台状态") or "").strip().lower() in confirmed
    ]
    return selected, True, True


def _ack_workorder_retry_queue(processed_rows, logger=None):
    """仅确认已得到明确结果的自动重查项；失败/未比对项保留待下次重试。"""
    try:
        with open(WORKORDER_RETRY_QUEUE_FILE, "r", encoding="utf-8") as fh:
            payload = json.load(fh)
    except (FileNotFoundError, OSError, ValueError):
        return 0
    if not isinstance(payload, dict) or not isinstance(payload.get("items"), list):
        return 0
    completed_keys = set()
    for row in processed_rows or []:
        key = row.get("_retry_queue_key")
        result = str(row.get("是否已录单") or "").strip()
        query_status = str(row.get("RPA查询状态") or "").strip()
        if key and result in {"是", "否", "待复核"} and query_status not in {
            "登录失败", "查询异常", "查询未执行", "查询条件填写失败", "查询按钮失败", "结果表解析失败",
        }:
            completed_keys.add(str(key))
    if not completed_keys:
        return 0
    before = len(payload["items"])
    payload["items"] = [
        item for item in payload["items"]
        if not (isinstance(item, dict) and str(item.get("key")) in completed_keys)
    ]
    payload["updated_at"] = datetime.now().isoformat(timespec="seconds")
    try:
        os.makedirs(os.path.dirname(WORKORDER_RETRY_QUEUE_FILE), exist_ok=True)
        tmp = WORKORDER_RETRY_QUEUE_FILE + ".tmp"
        with open(tmp, "w", encoding="utf-8") as fh:
            json.dump(payload, fh, ensure_ascii=False, indent=2)
        os.replace(tmp, WORKORDER_RETRY_QUEUE_FILE)
    except OSError as exc:
        if logger:
            logger.warning(f"自动重查队列回写失败，已保留待处理项: {exc}")
        return 0
    count = before - len(payload["items"])
    if logger and count:
        logger.info(f"自动重查队列已完成 {count} 条，剩余 {len(payload['items'])} 条")
    return count


def _ingest_workbench_database(rows, dataset, source_path, logger=None):
    """程序输出后同步到独立工作台数据库；失败不阻断原有 Excel 流程。"""
    try:
        from workbench_database import WorkbenchDatabase, default_database_path

        db = WorkbenchDatabase(default_database_path(APP_DIR))
        result = db.ingest_rows(rows or [], dataset=dataset, source_path=str(source_path or ""))
        if logger:
            logger.info(
                f"工作台数据库增量写入: {dataset} {result.get('rows', 0)} 条 "
                f"(新增 {result.get('inserted', 0)} / 更新 {result.get('updated', 0)})"
            )
        return result
    except Exception as exc:
        if logger:
            logger.warning(f"工作台数据库同步失败，已保留 Excel 输出: {exc}")
        return {"rows": 0, "inserted": 0, "updated": 0, "error": str(exc)}


def _cache_internal_email_table(source_path: str, cache_path: str = INTERNAL_EMAIL_CACHE_FILE) -> str:
    """把用户导入的内部邮箱表规范化到本地缓存。

    缓存保留邮箱全称、邮箱后缀、所属部门/用途和是否内部四类字段；
    即使原表只有“姓名 + 邮箱地址”，也会补齐后缀与“是”，供规则层
    稳定读取。返回缓存文件路径，失败时由调用方提示用户，不会替换旧缓存。
    """
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill

    source_path = os.path.abspath(source_path)
    cache_path = os.path.abspath(cache_path)
    source_wb = load_workbook(source_path, read_only=True, data_only=True)
    source_ws = source_wb.active
    rows = list(source_ws.iter_rows(values_only=True))
    source_wb.close()
    if not rows:
        raise ValueError("内部邮箱表为空")

    headers = [str(value or "").strip() for value in rows[0]]
    email_re = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
    def find_index(words):
        for index, header in enumerate(headers):
            if any(word in header.lower() for word in words):
                return index
        return None

    email_index = find_index(("邮箱", "email", "mail"))
    name_index = find_index(("姓名", "名称", "name"))
    dept_index = find_index(("部门", "用途", "dept", "use"))
    internal_index = find_index(("是否内部", "内部"))
    output = []
    seen = set()
    for raw in rows[1:]:
        values = ["" if value is None else str(value).strip() for value in raw]
        cells = [values[email_index]] if email_index is not None and email_index < len(values) else values
        emails = []
        for cell in cells:
            emails.extend(email_re.findall(cell))
        for email in emails:
            email = email.lower()
            if email in seen:
                continue
            seen.add(email)
            suffix = "@" + email.split("@", 1)[1]
            name = values[name_index] if name_index is not None and name_index < len(values) else ""
            dept = values[dept_index] if dept_index is not None and dept_index < len(values) else "内部邮箱"
            internal = values[internal_index] if internal_index is not None and internal_index < len(values) else "是"
            # “邮箱全称”按规则表语义保存完整地址；显示名另列，避免把姓名
            # 误当成可用于匹配的邮箱标识。
            output.append([email, email, suffix, dept or "内部邮箱", internal or "是", name])
    if not output:
        raise ValueError("未找到有效邮箱地址")

    os.makedirs(os.path.dirname(cache_path), exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "内部邮箱表"
    out_headers = ["邮箱全称", "邮箱地址", "邮箱后缀", "所属部门/用途", "是否内部", "联系人/显示名称"]
    ws.append(out_headers)
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="0F766E")
    for row in output:
        ws.append(row)
    widths = [34, 34, 28, 24, 12, 20]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + index)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(cache_path)
    return cache_path


class WorkerThread(QThread):
    """后台处理线程 — 支持独立运行 stage1(邮件解析) / stage2(工单查询) / all(全流程)"""
    log_signal = pyqtSignal(str)
    progress_signal = pyqtSignal(int, int)  # current, total
    finished_signal = pyqtSignal(str, str)  # primary_file, secondary_file
    cancelled_signal = pyqtSignal(str)
    error_signal = pyqtSignal(str)

    def __init__(self, config, mode="all", date_from=None, date_to=None,
                 agent_email_path=None, internal_email_path=None, project_table_path=None,
                 stage2_input_path=None, prefer_internal_email_path=False):
        super().__init__()
        self.config = config
        self.mode = mode  # "stage1" | "stage2" | "all"
        self.date_from = date_from
        self.date_to = date_to
        self.agent_email_path = agent_email_path
        self.internal_email_path = internal_email_path
        self.project_table_path = project_table_path
        self.stage2_input_path = stage2_input_path  # 阶段二输入：阶段一输出的 xlsx
        self.prefer_internal_email_path = bool(prefer_internal_email_path)
        self._stop = False
        self._stop_event = threading.Event()
        self._gui_log_handler = None
        self.imap_total = 0
        self.imap_loaded_total = 0

    def stop(self, source="运行控制"):
        self._stop = True
        self._stop_event.set()
        self._log(f"收到停止请求：来源={source}，阶段={getattr(self, '_stage1_phase', self.mode)}")

    def is_stop_requested(self):
        """供 RPA 子模块安全读取的跨线程停止状态。"""
        return self._stop_event.is_set()

    def run(self):
        try:
            if self.mode == "stage2":
                self._run_stage2()
            elif self.mode == "stage1":
                self._run_stage1()
            else:  # "all"
                self._run_all()
        except Exception as e:
            import traceback
            self.error_signal.emit(f"{e}\n{traceback.format_exc()}")
        else:
            if self._stop:
                self.cancelled_signal.emit(
                    "任务已停止，未完成的阶段不会覆盖结果文件。已完成的工单查询缓存保留。"
                )
        finally:
            self._detach_gui_log_handler()

    def _create_run_logger(self):
        """每次运行创建独立文件日志，并且只挂一个 GUI handler。"""
        import logging
        log_dir = os.path.join(APP_DIR, "logs")
        logger = setup_logger(log_dir=log_dir, level="DEBUG")
        handler = GuiLogHandler(self._log)
        handler.setLevel(logging.INFO)
        fmt = logging.Formatter("%(asctime)s [%(levelname)s] %(message)s", "%H:%M:%S")
        handler.setFormatter(fmt)
        logger.addHandler(handler)
        self._gui_log_handler = handler
        return logger

    def _detach_gui_log_handler(self):
        if self._gui_log_handler is None:
            return
        import logging
        logger = logging.getLogger("mail_audit")
        logger.removeHandler(self._gui_log_handler)
        self._gui_log_handler.close()
        self._gui_log_handler = None

    def _run_stage1(self):
        """阶段一: 邮件拉取→过滤→LLM意图分类→字段提取→OCR懒加载兜底→项目标准化→输出三份stage1产物"""
        logger = self._create_run_logger()

        logger.info("=" * 50)
        logger.info("[阶段一] 邮件过滤解析 (不含工单查询)")
        logger.info("=" * 50)

        result = self._run_pipeline_stage1(logger)
        if result is None:
            if not self._stop:
                raise RuntimeError("阶段一未返回有效结果，请检查运行日志；原结果文件未覆盖")
            logger.info(f"阶段一已停止：{getattr(self, '_stage1_phase', '阶段一')}；未覆盖结果文件")
            return
        all_rows, filtered_mails = result

        if self._stop:
            return

        # 完整成功后才按日期替换；新旧文件、数据库和人工状态一同发布。
        published = self._publish_stage1_results(all_rows, filtered_mails, logger)
        if published is not None:
            self.finished_signal.emit(*published)

    def _publish_stage1_results(self, all_rows, filtered_mails, logger):
        from utils.stage1_refresh import publish_stage1
        published = publish_stage1(
            APP_DIR, self.config["output"]["dir"],
            bool(self.config.get("output", {}).get("categorized", True)),
            all_rows, filtered_mails, self.date_from, self.date_to, logger,
            cancelled=self.is_stop_requested,
        )
        if published is None:
            return
        primary, secondary = published
        _write_imap_summary(
            self.imap_total,
            self.date_from,
            self.date_to,
            loaded_total=self.imap_loaded_total,
        )
        filtered_path = os.path.join(os.path.dirname(str(primary)), "filtered_mail_record.xlsx")
        cache_folder = _archive_stage1_cache(
            primary,
            secondary,
            filtered_path,
            date_from=self.date_from,
            date_to=self.date_to,
            imap_total=self.imap_total,
        )
        if cache_folder:
            logger.info(f"阶段一日期缓存已保存: {cache_folder}")
        logger.info("=" * 50)
        logger.info(f"[阶段一完成] 待查清单: {primary}")
        logger.info(f"[阶段一完成] 漏单复查: {secondary}")

        return primary, secondary

    def _run_stage2(self):
        """阶段二: 读取阶段一输出 → 前置校验 → 工单查询 → 输出 workorder_check_result.xlsx"""
        logger = self._create_run_logger()

        logger.info("=" * 50)
        logger.info("[阶段二] 工单系统比对 (读阶段一输出)")
        logger.info("=" * 50)

        if not self.stage2_input_path or not os.path.exists(self.stage2_input_path):
            logger.error(f"阶段二输入文件不存在: {self.stage2_input_path}")
            self.error_signal.emit(f"阶段二输入文件不存在: {self.stage2_input_path}")
            return

        from openpyxl import load_workbook
        wb = load_workbook(self.stage2_input_path, read_only=True, data_only=True)
        # 阶段一输出的工作表可能被用户手动打开后切换了 active sheet；
        # 阶段二应按工作表名称读取，而不是依赖 active 状态。
        ws = wb["工单待查"] if "工单待查" in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            logger.error("阶段二输入文件为空")
            return
        data_rows = [r for r in rows[1:] if any(v is not None and str(v).strip() for v in r)]
        if not data_rows:
            reason = "阶段二输入的‘工单待查’工作表没有数据行，请选择最新的阶段一输出"
            logger.error(reason)
            self.error_signal.emit(reason)
            return
        headers = [str(c) if c is not None else "" for c in rows[0]]
        required_headers = {"代理", "客户公司名称", "标准化项目名称"}
        if ws.title != "工单待查" or not required_headers.issubset(set(headers)):
            reason = (
                f"阶段二输入不是阶段一工单待查清单: 工作表={ws.title}, "
                f"缺少表头={sorted(required_headers.difference(set(headers)))}"
            )
            logger.error(reason)
            self.error_signal.emit(reason)
            return
        logger.info("阶段二输入校验通过: 阶段一工作表“工单待查”")
        # 表头名 → 内部字段名(客户公司名称→客户 等), 否则下游取值为空
        from modules.excel_writer import normalize_header_rows
        all_rows = normalize_header_rows(headers, data_rows)
        # 工作台导出包含每条明细的“工作台状态”。只有状态为 confirmed
        # 的明细才允许进入阶段二；原始阶段一文件没有该列，保持原有
        # 阶段二全部候选的行为，避免破坏旧流程。
        before_count = len(all_rows)
        all_rows, is_workbench_export, has_workbench_status = _filter_workbench_confirmed_rows(
            headers, all_rows, self.stage2_input_path
        )
        if is_workbench_export and has_workbench_status:
            logger.info(
                f"工作台确认口径：{before_count} 条中 {len(all_rows)} 条已确认，"
                f"仅将已确认明细送入阶段二"
            )
        elif is_workbench_export:
            logger.warning(
                "工作台导出缺少‘工作台状态’列，无法按确认状态过滤；"
                "将兼容旧文件继续处理全部行"
            )
        logger.info(f"已读取 {len(all_rows)} 行工单待查清单")
        _apply_workorder_retry_queue(all_rows, logger)

        if self._stop:
            return

        # 前置校验: 跳过代理空+置信度需人工确认; (公司,项目) 去重
        from modules.workorder_checker import WorkOrderChecker
        to_query, skipped = WorkOrderChecker.preprocess_rows(all_rows)
        repaired = [r for r in to_query if r.get("_stage2_note")]
        if repaired:
            logger.warning(f"前置校验已修复 {len(repaired)} 条异常公司字段:")
            for r in repaired[:5]:
                logger.warning(
                    f"  - 已修复: {r.get('客户', '')} — {r.get('_stage2_note', '')}"
                )
        if skipped:
            logger.warning(f"前置校验跳过 {len(skipped)} 条 (代理空+需人工确认 / 重复):")
            for s in skipped[:5]:
                logger.warning(f"  - 跳过: {s.get('客户','')} / {s.get('项目','')} — {s.get('_skip_reason','')}")
        logger.info(f"待RPA查询: {len(to_query)} 条 (原始 {len(all_rows)} 条)")

        if not to_query:
            logger.info("无可查询条目，跳过RPA")
            # 即使全跳空也输出空结果文件
            from modules.excel_writer import ExcelWriter
            writer = ExcelWriter(
                self.config["output"]["dir"], logger,
                categorized=bool(self.config.get("output", {}).get("categorized", True)),
            )
            primary, secondary = writer.write_workorder_check_result(all_rows, [], skipped)
            _ingest_workbench_database(all_rows, "workorder", primary, logger)
            self.finished_signal.emit(primary, secondary)
            return

        # M5: 工单系统比对
        checker = WorkOrderChecker(
            self.config["workorder"], logger,
            cancel_requested=self.is_stop_requested,
        )

        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        processed_rows = []
        try:
            login_ok = loop.run_until_complete(
                checker._await_or_cancel(checker.login(), False)
            )
            if self._stop:
                logger.warning("用户在登录/导航期间停止，未生成阶段二结果文件")
                return
            if not login_ok:
                logger.error("工单系统登录失败，跳过RPA比对")
                for r in to_query:
                    r["是否已录单"] = "未比对"
                    r["匹配状态"] = "工单系统未连接"
                    r["RPA查询状态"] = "登录失败"
                    r["查询时间戳"] = datetime.now().isoformat(timespec="seconds")
                processed_rows = list(to_query)
            else:
                try:
                    processed_rows = loop.run_until_complete(checker.search_batch(to_query)) or []
                except Exception as e:
                    logger.error(f"RPA查询异常: {e}", "error")
                    processed_rows = list(getattr(checker, "last_processed_rows", []))
                if self._stop:
                    logger.warning("用户已停止阶段二，未生成不完整的工单核对结果文件")
                    return
        finally:
            try:
                loop.run_until_complete(checker.close())
            except Exception as exc:
                logger.warning(f"停止后关闭浏览器异常: {exc}")
            finally:
                loop.close()

        # 输出 stage2 产物
        from modules.excel_writer import ExcelWriter
        writer = ExcelWriter(
            self.config["output"]["dir"], logger,
            categorized=bool(self.config.get("output", {}).get("categorized", True)),
        )
        primary, secondary = writer.write_workorder_check_result(
            all_rows, processed_rows, skipped
        )
        _ingest_workbench_database(processed_rows, "workorder", primary, logger)
        _ack_workorder_retry_queue(processed_rows, logger)
        pending_count = max(0, len(to_query) - len(processed_rows))
        logger.info(
            f"阶段二实际处理: {len(processed_rows)}/{len(to_query)} 条"
        )
        if pending_count:
            logger.warning(f"阶段二未执行 {pending_count} 条，结果表已标记为“未执行”")
        logger.info("=" * 50)
        logger.info(f"[阶段二完成] 工单核对结果: {primary}")

        self.finished_signal.emit(primary, secondary)

    def _run_all(self):
        """一站式: stage1 + stage2 (兼容原行为, 阶段一阶段二连续跑)"""
        logger = self._create_run_logger()

        logger.info("=" * 50)
        logger.info("[一站式] 邮件 + 工单全流程")
        logger.info("=" * 50)

        result = self._run_pipeline_stage1(logger)
        if result is None:
            if not self._stop:
                raise RuntimeError("阶段一未返回有效结果，请检查运行日志；未进入工单查询")
            logger.info(f"一站式已停止：{getattr(self, '_stage1_phase', '阶段一')}；未进入工单查询")
            return
        all_rows, filtered_mails = result

        if self._stop:
            return

        # 阶段二手动并行: 直接用 stage1 结果
        if self._publish_stage1_results(all_rows, filtered_mails, logger) is None:
            return
        from openpyxl import load_workbook
        from modules.workorder_checker import WorkOrderChecker
        from modules.excel_writer import ExcelWriter

        to_query, skipped = WorkOrderChecker.preprocess_rows(all_rows)
        _apply_workorder_retry_queue(to_query, logger)
        logger.info(f"前置校验: 待RPA {len(to_query)} 条, 跳过 {len(skipped)} 条")

        checker = WorkOrderChecker(
            self.config["workorder"], logger,
            cancel_requested=self.is_stop_requested,
        )
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        try:
            login_ok = loop.run_until_complete(
                checker._await_or_cancel(checker.login(), False)
            )
            if self._stop:
                logger.warning("用户在登录/导航期间停止，未生成一站式结果文件")
                return
            if not login_ok:
                logger.warning("工单系统登录失败")
                for r in all_rows:
                    r["是否已录单"] = "未比对"
                    r["匹配状态"] = "工单系统未连接"
                    r["查询时间戳"] = datetime.now().isoformat(timespec="seconds")
            else:
                try:
                    loop.run_until_complete(checker.search_batch(to_query))
                    if self._stop:
                        logger.warning("用户已停止一站式阶段二，未生成不完整的结果文件")
                        return
                    # 把查询结果合并回 all_rows
                    queried = {id(r): r for r in to_query}
                    for r in all_rows:
                        q = queried.get(id(r))
                        if q:
                            r["是否已录单"] = q.get("是否已录单", "")
                            r["工单日期"] = q.get("工单日期") or r.get("date")
                            r["下单日期"] = q.get("下单日期", "")
                            r["匹配状态"] = q.get("匹配状态", "")
                            r["工单记录"] = q.get("工单记录", [])
                            r["查询时间戳"] = q.get("查询时间戳")
                            r["RPA查询状态"] = q.get("RPA查询状态", "")
                            r["邮件日期筛选范围"] = q.get("邮件日期筛选范围", "")
                            r["日期筛选说明"] = q.get("日期筛选说明", "")
                except Exception as e:
                    logger.error(f"RPA查询异常: {e}", "error")
        finally:
            try:
                loop.run_until_complete(checker.close())
            except Exception as exc:
                logger.warning(f"停止后关闭浏览器异常: {exc}")
            finally:
                loop.close()

        writer = ExcelWriter(
            self.config["output"]["dir"], logger,
            categorized=bool(self.config.get("output", {}).get("categorized", True)),
        )
        _ack_workorder_retry_queue(to_query, logger)
        primary, secondary = writer.write_all_outputs(all_rows, filtered_mails)
        _write_imap_summary(
            self.imap_total,
            self.date_from,
            self.date_to,
            loaded_total=self.imap_loaded_total,
        )
        # 一站式模式也要把阶段二结论增量写入工作台数据库；否则自动重查
        # 后的结果只能留在本次 Excel，工作台刷新时无法与人工覆盖层统一。
        _ingest_workbench_database(all_rows, "workorder", primary, logger)
        logger.info("=" * 50)
        logger.info(f"[一站式完成] 漏单清单: {primary}")
        logger.info(f"[一站式完成] 过滤清单: {secondary}")
        self.finished_signal.emit(primary, secondary)

    def _run_pipeline_stage1(self, logger):
        """阶段一内部 pipeline (M1-M4), 返回 (rows, filtered_mails) 或 None=用户停止"""
        from openpyxl import load_workbook
        self._stage1_phase = "读取参考表"
        ref_cfg = self.config["reference_tables"]

        # 加载附件二: 内部邮箱
        # 本地规范化缓存优先；用户通过 GUI 更换表格时会覆盖该缓存。
        # 这样正式运行不会因为会话里残留旧路径而回读旧版本表格。
        if self.prefer_internal_email_path and self.internal_email_path:
            # 网页上传使用独立留存文件，不覆盖桌面缓存；明确选择必须优先。
            internal_email_path = self.internal_email_path
        elif os.path.exists(INTERNAL_EMAIL_CACHE_FILE):
            internal_email_path = INTERNAL_EMAIL_CACHE_FILE
        else:
            source_internal_path = self.internal_email_path or ref_cfg["internal_emails"]
            try:
                # 首次运行也把默认表规范化落到缓存，之后不再依赖来源路径。
                internal_email_path = _cache_internal_email_table(source_internal_path)
                logger.info(f"内部邮箱表已规范化缓存: {internal_email_path}")
            except Exception as exc:
                # 兼容旧版/非 xlsx 表格；读取失败只回退到原有读取逻辑，
                # 不让缓存功能改变阶段一的既有可用性。
                logger.warning(f"内部邮箱表缓存未生成，沿用来源表: {exc}")
                internal_email_path = source_internal_path
        internal_emails = self._load_internal_emails(internal_email_path)
        logger.info(f"内部邮箱表已加载: {len(internal_emails)} 条")

        # 加载附件三: 代理邮箱
        agent_email_path = self.agent_email_path or ref_cfg["agent_emails"]
        agent_map = self._load_agent_emails(agent_email_path)
        logger.info(f"代理邮箱表已加载: {len(agent_map)} 条")

        # 加载附件四: 项目名称
        project_table_path = self.project_table_path or ref_cfg["project_names"]
        project_table = self._load_project_names(project_table_path)
        logger.info(f"项目名称表已加载: {len(project_table)} 条")

        # 项目表是阶段一项目标准化的必要输入；缺失或为空时，
        # 继续跑下去会产生大量"项目标准化失败"且后续工单查询无法命中。
        # 此处明确提示并停止本次处理，由用户在 GUI 中选择有效文件后重跑。
        if not project_table:
            abs_path = os.path.abspath(project_table_path)
            logger.error(
                f"项目名称表无效(缺失或 0 条): {abs_path}。"
                f"请在 GUI 点击'项目表导入'选择有效文件后重跑，本次处理已终止。"
            )
            raise ValueError("项目名称表缺失或没有有效数据，请重新导入项目名称表后运行")

        # 创建 LLM 客户端
        self._stage1_phase = "LLM连通性预检"
        llm_client = None
        llm_cfg = self.config.get("llm", {})
        api_key_env = str(llm_cfg.get("api_key_env", "")).strip()
        has_llm_key = bool(os.getenv(api_key_env)) if api_key_env else bool(llm_cfg.get("api_key"))
        if has_llm_key:
            llm_client = LLMIntentClient(llm_cfg, logger)
            if llm_client.enabled:
                summary = llm_client.get_usage_summary()
                logger.info(
                    "LLM 已启用: "
                    f"model={summary['configured_model']}, endpoint={summary['endpoint_host']}, "
                    f"credential={summary['credential_source']}"
                )
                if summary["credential_source"] == "config.api_key":
                    logger.warning(
                        "LLM API Key 仍从 config.yaml 读取；请改用 api_key_env 指定的环境变量"
                    )
                llm_client.preflight()
            else:
                logger.warning("LLM 配置异常, 将仅使用规则引擎")
        else:
            logger.warning(
                f"未配置 LLM 环境变量 {api_key_env or '或 config.api_key'}，仅使用规则引擎"
            )

        if self._stop:
            return None

        # M1: 邮件读取
        self._stage1_phase = "邮件下载与附件解析"
        from modules.mail_reader import MailReader
        # 每次运行从邮箱重新读取，不能由旧配置开启原始邮件缓存。
        reader = MailReader({**self.config["email"], "cache_enabled": False}, logger)
        self_email = self.config["email"]["address"]
        mails = reader.fetch_mails(
            self.date_from, self.date_to, self_email,
            progress_callback=lambda c, t: self.progress_signal.emit(c, t),
            cancel_requested=self.is_stop_requested,
        )
        self.imap_total = int(getattr(reader, "last_search_total", 0) or 0)
        self.imap_loaded_total = len(mails)
        if not self.imap_total:
            # 兼容旧版/测试替身 MailReader：没有 SEARCH 元数据时，
            # 至少保留本次解析到的邮件数量，避免统计卡显示空值。
            self.imap_total = len(mails)
        logger.info(
            f"邮件拉取完成: IMAP搜索 {self.imap_total} 封，成功解析 {self.imap_loaded_total} 封"
        )

        if self._stop:
            return None

        # M2: 邮件过滤
        if getattr(reader, "last_fetch_complete", True) is False:
            raise RuntimeError("本次邮件未完整读取，已保留之前结果；请重试所选日期")
        self._stage1_phase = "邮件规则过滤"
        from modules.mail_filter import MailFilter
        audit_addresses = self.config.get("email", {}).get("audit_addresses", [])
        if isinstance(audit_addresses, str):
            audit_addresses = [audit_addresses]
        filt = MailFilter(
            internal_emails,
            logger,
            audit_mailbox=self_email,
            audit_mailboxes=set(audit_addresses or []),
        )
        valid_mails, filtered_mails = filt.filter_mails(mails)
        logger.info(f"规则过滤完成: 有效 {len(valid_mails)} 封, 过滤 {len(filtered_mails)} 封")

        if self._stop:
            return None

        # M2.5: 意图 Agent — 全量识别。
        # 规则引擎只负责预分类和提供证据，不再把“公司改名/地址变更/证书通知/
        # 续费”等复杂情况直接当成最终结论。每封从 IMAP 拉取的邮件都交给同一个
        # 意图 Agent 判断，避免规则词表覆盖不全导致漏单或误收。
        self._stage1_phase = "邮件意图识别"
        intent_candidates = list(valid_mails) + list(filtered_mails)
        if llm_client and llm_client.enabled and intent_candidates:
            logger.info(
                f"LLM 全量意图识别开始: {len(intent_candidates)} 封邮件全部进入意图 Agent..."
            )
            recovered, still_filtered = filt.llm_second_pass(
                intent_candidates, llm_client,
                progress_callback=lambda c, t: self.progress_signal.emit(c, t)
            )
            if recovered:
                logger.info(f"LLM 确认目标或转人工保留 {len(recovered)} 封邮件")
            # 全量模式下，valid_mails 不再绕过模型；最终结果完全由模型结果
            # 加上失败转人工策略组成。
            valid_mails = recovered
            filtered_mails = still_filtered
            logger.info(f"LLM 全量意图识别完成: 保留 {len(valid_mails)} 封, 过滤 {len(filtered_mails)} 封")
        else:
            logger.info("跳过 LLM 全量意图识别 (未配置或本次没有邮件)")
            logger.info(
                "[Agent状态] 意图识别Agent 跳过: "
                f"原因={'未配置/未启用' if not llm_client or not llm_client.enabled else '本次没有邮件'}"
            )

        if self._stop:
            return None

        # M3: 字段提取 (规则 + LLM 补充 + OCR 图片兜底)
        self._stage1_phase = "字段提取"
        from modules.field_extractor import FieldExtractor
        ocr_fallback = self.config.get("ocr", {}).get("fallback", True)
        extractor = FieldExtractor(agent_map, project_table, logger, llm_client,
                                   ocr_fallback=ocr_fallback)
        all_rows = []
        mail_blocks = []   # 每封邮件的提取行，用于后续「同封邮件内」去重与语义校验分组
        for idx, mail in enumerate(valid_mails, 1):
            if self._stop:
                return None
            rows = extractor.extract_fields(mail)
            if not rows:
                # FieldExtractor normally guarantees one review row. This second guard
                # keeps a future parser regression from aborting an entire mailbox run.
                logger.error(
                    f"字段提取未返回行，已保留该邮件供排查: 第{idx}/{len(valid_mails)}封"
                )
                logger.warning(
                    f"[Agent状态] 字段抽取Agent 未产出可用记录: 第{idx}/{len(valid_mails)}封；"
                    "请检查字段解析或模型返回"
                )
                self.progress_signal.emit(idx, len(valid_mails))
                continue
            all_rows.extend(rows)
            mail_blocks.append(rows)
            first_row = rows[0]
            field_llm_status = str(first_row.get("字段LLM状态", "") or "未调用")
            field_llm_used = bool(first_row.get("llm_used"))
            logger.info(
                f"字段提取: 第{idx}/{len(valid_mails)}封 — "
                f"客户='{first_row.get('客户','')}', 代理='{first_row.get('代理','')}', "
                f"字段抽取Agent={'已调用' if field_llm_used else '未调用'}, "
                f"状态={field_llm_status}"
            )
            self.progress_signal.emit(idx, len(valid_mails))

        if self._stop:
            return None

        # M4: 项目标准化拆分 — 所有非空项目统一走 normalizer（组合拆分+标准化）
        self._stage1_phase = "标准化与语义复检"
        from modules.project_normalizer import ProjectNormalizer, prune_umbrella_epr
        normalizer = ProjectNormalizer(project_table, logger)
        normalized_rows = []
        mail_keys = []     # 与 normalized_rows 并行：每行所属邮件序号
        for block_id, block in enumerate(mail_blocks):
            expanded = []
            for row in block:
                proj_raw = row.get("项目", "")
                if proj_raw:
                    split_results = normalizer.normalize_and_split(
                        proj_raw, row.get("subject", ""), row.get("body_text", "")
                    )
                    if len(split_results) > 1:
                        for sr in split_results:
                            new_row = row.copy()
                            new_row["项目"] = sr["standard_name"]
                            new_row["项目原始值"] = proj_raw
                            expanded.append(new_row)
                        continue
                    if len(split_results) == 1 and split_results[0]["standard_name"]:
                        row["项目"] = split_results[0]["standard_name"]
                expanded.append(row)

            # M4.1: 「国家+EPR」统称去重。代理商口语常拿 "奥地利EPR" 代指该国三项，
            # 而项目名称表里只有 奥地利WEEE/电池法/包装法。若同一封邮件已给出该国
            # 具体项目，这个统称行就是冗余；若整封只写了统称，则保留并标注待确认。
            names = [str(r.get("项目", "") or "") for r in expanded]
            drop_idx, flag_idx = prune_umbrella_epr(names)
            for i in sorted(drop_idx):
                country = names[i].replace("EPR", "").replace("epr", "").strip()
                covered = "、".join(n for n in names if n and country and country in n and n != names[i])
                logger.info(
                    f"丢弃同国统称冗余行: {names[i]}"
                    + (f"（同封邮件已有: {covered}）" if covered else "")
                )
            for i in sorted(flag_idx):
                row = expanded[i]
                hint = "项目为“国家+EPR”统称，未写明具体业务类型，请人工确认是 WEEE/电池法/包装法 中哪几项"
                existing = str(row.get("人工复核提示", "") or "")
                row["人工复核提示"] = "；".join(x for x in [existing, hint] if x)
                row["置信度"] = "需人工确认"

            kept = [r for i, r in enumerate(expanded) if i not in drop_idx]
            normalized_rows.extend(kept)
            mail_keys.extend([block_id] * len(kept))

        logger.info(f"项目标准化完成: {len(normalized_rows)} 行")
        # M4.2: 确定性业务校验先拦截“一家公司”、说明句、文件名、空项目等
        # 明显异常。它只打标，不改写模型/规则提取值。
        from modules.business_validator import apply_row_guard
        guarded_count = 0
        for row in normalized_rows:
            if apply_row_guard(row):
                guarded_count += 1
        if guarded_count:
            logger.warning(f"确定性业务校验标记 {guarded_count} 行进入人工复核")

        # M4.5: 语义质量校验。仅提供建议，不覆盖程序抽取值；不确定项会
        # 自动进入人工复核工作台，供商务同事查看正文证据后确认。
        if llm_client and llm_client.enabled and normalized_rows:
            # 语义 Agent 不再无差别复检所有明细。普通的单公司、单项目、
            # 规则已通过且没有 LLM 补全的行，沿用确定性规则即可；只有多实体
            # 邮件或存在不确定信号的行才交给 Agent。这样既保留复杂邮件的
            # 交叉判断能力，也避免每次运行为大量简单行重复付费。
            review_indices, route_reasons = self._semantic_review_indices(
                normalized_rows, mail_keys
            )
            review_rows = [normalized_rows[i] for i in review_indices]
            review_keys = [
                mail_keys[i] if mail_keys is not None and i < len(mail_keys) else i
                for i in review_indices
            ]
            logger.info(
                "LLM 语义字段校验风险路由: "
                f"{len(review_rows)}/{len(normalized_rows)} 行进入 Agent；"
                f"原因={route_reasons}"
            )
            if review_rows:
                reviewed_rows = self._semantic_validate_rows(
                    review_rows, llm_client, logger, mail_keys=review_keys
                )
                for row_index, reviewed in zip(review_indices, reviewed_rows):
                    normalized_rows[row_index] = reviewed

            # 语义 Agent 已确认“意图低置信”但字段完整且业务规则通过的记录，
            # 不再因为过滤器的保守标记继续进入人工复核。确定性业务风险、
            # Agent 返回 invalid/uncertain 或字段缺失的记录仍保留人工状态。
            promoted = 0
            for row_index in review_indices:
                row = normalized_rows[row_index]
                complete = all(str(row.get(field, "") or "").strip()
                               for field in ("代理", "客户", "项目", "需求"))
                if (
                    complete
                    and row.get("语义校验状态") == "valid"
                    and row.get("业务规则校验状态") in (None, "", "通过")
                    and row.get("置信度") not in {"需人工确认", "low", "medium", "低", "中"}
                ):
                    row["置信度"] = "high"
                    # 这是过滤器的低置信提示，不是仍需人工修正的字段问题。
                    row["人工复核提示"] = ""
                    promoted += 1
            if promoted:
                logger.info(f"语义 Agent 已确认 {promoted} 条完整记录，转为可确认")

            reviewed_set = set(review_indices)
            for row_index, row in enumerate(normalized_rows):
                if row_index in reviewed_set:
                    continue
                missing_fields = [
                    field for field in ("代理", "客户", "项目", "需求")
                    if not str(row.get(field, "") or "").strip()
                ]
                if missing_fields or row.get("置信度") in {"low", "medium", "低", "中", "需人工确认"}:
                    row["语义校验状态"] = "人工复核（未调用语义Agent）"
                    row["语义校验建议"] = "先补全缺失字段，再由人工核对原邮件"
                    row["语义校验原因"] = (
                        "字段缺失/置信度不足，需人工补全；本行未命中多公司多项目关联判断"
                    )
                else:
                    row["语义校验状态"] = "规则通过（未调用）"
                    row["语义校验建议"] = "未命中复杂场景，沿用规则提取结果"
                    row["语义校验原因"] = "单公司单项目且确定性规则通过"
                row["语义问题字段"] = ""
                row["语义问题编号"] = ""
                row["语义问题证据"] = ""
                row["语义当前值"] = ""
                row["语义建议值"] = ""
                row["语义校验模型"] = ""
        else:
            logger.info(
                "[Agent状态] 语义复检Agent 跳过: "
                f"原因={'未配置/未启用' if not llm_client or not llm_client.enabled else '没有可复检的标准化记录'}"
            )
        if llm_client:
            summary = llm_client.get_usage_summary()
            logger.info(
                "LLM 本次调用汇总: "
                f"发起={summary['requests_started']}, 成功={summary['requests_succeeded']}, "
                f"失败={summary['requests_failed']}, tokens={summary['total_tokens']}, "
                f"结构失败={summary.get('output_validation_failures', 0)}, "
                f"修复={summary.get('repair_successes', 0)}/{summary.get('repair_attempts', 0)}, "
                f"审计={summary['audit_log_path']}"
            )

        # 同一封邮件可能同时产生“1家公司”和真实公司两套候选。若同一邮件、
        # 同一项目已有唯一合法公司名，丢弃占位公司行，避免工作台展示重复脏数据。
        normalized_rows, mail_keys = self._dedupe_placeholder_rows(
            normalized_rows, mail_keys, logger
        )
        # M4 的项目标准化可能把一个组合项目拆成多行，也可能删除同国统称。
        # 因此附件 Excel 的“18 条”校验必须在最终输出前再计算一次，不能只
        # 使用字段抽取阶段的中间行数。
        self._refresh_attachment_record_counts(normalized_rows, mail_keys, logger)
        return normalized_rows, filtered_mails

    @staticmethod
    def _refresh_attachment_record_counts(rows, mail_keys=None, logger=None):
        """按邮件块复核附件表格记录数与最终输出行数，并留下人工提示。"""
        from collections import defaultdict

        groups = defaultdict(list)
        for index, row in enumerate(rows):
            expected = row.get("附件表格记录数")
            try:
                expected = int(expected or 0)
            except (TypeError, ValueError):
                expected = 0
            if expected <= 0:
                continue
            key = mail_keys[index] if mail_keys is not None and index < len(mail_keys) else (
                row.get("sender_email", ""), row.get("date", ""), row.get("subject", "")
            )
            groups[key].append(index)

        for key, indices in groups.items():
            expected_values = {
                int(rows[index].get("附件表格记录数") or 0)
                for index in indices
            }
            # 同一邮件只能有一条数量基线；若出现多个值，保守按最大值并转人工。
            expected = max(expected_values) if expected_values else 0
            actual = len(indices)
            count_ok = len(expected_values) == 1 and actual == expected
            message = (
                f"附件结构化表 {expected} 条，最终输出 {actual} 条"
                + ("，数量一致" if count_ok else "，数量不一致")
            )
            for index in indices:
                row = rows[index]
                row["附件表格输出数"] = actual
                row["附件表格数量校验"] = "通过" if count_ok else "需人工确认"
                if not count_ok:
                    row["置信度"] = "需人工确认"
                    existing = str(row.get("人工复核提示", "") or "").strip()
                    if message not in existing:
                        row["人工复核提示"] = "；".join(
                            value for value in (existing, message) if value
                        )
            if logger:
                level = "info" if count_ok else "warning"
                getattr(logger, level)(f"附件表格数量复核: {message}")

    @staticmethod
    def _semantic_review_indices(rows, mail_keys=None):
        """返回需要语义 Agent 复检的行号及路由原因计数。

        复杂邮件按邮件块整体送检，避免只把其中一条明细交给模型而丢失
        公司-项目的对应关系。单行如果存在确定性业务校验问题、字段 Agent
        失败或正文明确表示多家公司，也会进入复检。字段缺失但不涉及关联判断
        的行仍会进入人工复核，但不额外消耗语义 Agent 调用。
        """
        from collections import Counter

        group_sizes = Counter(mail_keys or [])
        selected = set()
        reasons = Counter()
        for idx, row in enumerate(rows):
            key = mail_keys[idx] if mail_keys is not None and idx < len(mail_keys) else idx
            reason_set = set()
            try:
                company_candidates = json.loads(str(row.get("客户候选证据") or "[]"))
            except (TypeError, ValueError):
                company_candidates = []
            if isinstance(company_candidates, list) and len(company_candidates) > 1:
                reason_set.add("EPR申请表公司候选待语义仲裁")
            if group_sizes.get(key, 1) > 1:
                reason_set.add("同封邮件多项目/多公司")

            if row.get("业务规则问题") or row.get("业务规则校验状态") not in (None, "", "通过"):
                reason_set.add("业务规则风险")
            if row.get("字段LLM失败原因") or row.get("字段LLM状态") in {"失败", "调用失败", "结构校验失败"}:
                reason_set.add("字段Agent失败")

            # 有些邮件只保留一行，但正文/主题明确写了“2家公司”“多家公司”;
            # 这类漏拆风险无法靠 group_sizes 发现，必须送 Agent 做关联判断。
            text = " ".join(
                str(row.get(field, "") or "")
                for field in ("subject", "body_text")
            )
            if re.search(r"(?:\d+|[两二三四五六七八九十多几]+)\s*家\s*(?:公司|主体)", text, re.I):
                reason_set.add("正文多公司提示")

            # 过滤器对“改名、撤单、证书变更”等非标准措辞会保守标记 uncertain。
            # 字段已经抽取完整时，交给语义 Agent 结合正文判断，不能直接留在
            # “未调用语义Agent”的人工复核队列。
            if str(row.get("filter_status", "") or "").strip().lower() in {
                "uncertain", "不确定", "待确认"
            }:
                reason_set.add("意图低置信度")

            if reason_set:
                selected.add(idx)
                for reason in reason_set:
                    reasons[reason] += 1

        return sorted(selected), dict(reasons)

    @staticmethod
    def _dedupe_placeholder_rows(rows, mail_keys=None, logger=None):
        """删除同邮件同项目下被真实公司名覆盖的占位公司候选。"""
        from collections import defaultdict
        from modules.business_validator import inspect_row

        groups = defaultdict(list)
        for idx, row in enumerate(rows):
            if mail_keys is not None and idx < len(mail_keys):
                mail_key = mail_keys[idx]
            else:
                mail_key = (
                    row.get("sender_email", ""), row.get("date", ""),
                    row.get("subject", "")
                )
            project = str(row.get("项目", "") or "").strip()
            groups[(mail_key, project)].append(idx)

        drop = set()
        for members in groups.values():
            valid_company_rows = []
            for idx in members:
                company_issues = {
                    item.get("field") for item in inspect_row(rows[idx])
                    if item.get("field") == "company"
                }
                if not company_issues and str(rows[idx].get("客户", "") or "").strip():
                    valid_company_rows.append(idx)
            # 只有存在唯一合法公司候选时才清理，避免模型没有真实公司名时
            # 把整封邮件的唯一结果误删。
            valid_companies = {
                str(rows[idx].get("客户", "") or "").strip()
                for idx in valid_company_rows
            }
            if len(valid_companies) != 1:
                continue
            for idx in members:
                if idx in valid_company_rows:
                    continue
                issues = inspect_row(rows[idx])
                if any(item.get("field") == "company" for item in issues):
                    drop.add(idx)

        if drop and logger:
            logger.warning(f"同邮件同项目清理占位公司候选: {len(drop)} 行")
        kept_rows = [row for idx, row in enumerate(rows) if idx not in drop]
        kept_keys = (
            [key for idx, key in enumerate(mail_keys) if idx not in drop]
            if mail_keys is not None else mail_keys
        )
        return kept_rows, kept_keys

    def _semantic_validate_rows(self, rows, llm_client, logger, mail_keys=None):
        """用 LLM 批量审核公司/代理/国家/项目/需求，保留原值并写入建议字段。

        ``mail_keys`` 与 ``rows`` 并行，标记每行来自哪封邮件。一条明细只承载一个
        项目，而正文往往同时写了多个，LLM 拿到单行会误判成"漏拆"。把同封邮件的
        其余项目一并告诉它，这类误报才不会把置信度无谓地压成"需人工确认"。
        """
        batch_size = 6
        try:
            batch_size = max(2, min(10, int(self.config.get("llm", {}).get("semantic_batch_size", 6))))
        except (TypeError, ValueError):
            pass

        # 同一封邮件的全部项目，供 LLM 判断"本行是否只是拆分结果的一部分"
        siblings = {}
        if mail_keys is not None and len(mail_keys) == len(rows):
            for key, row in zip(mail_keys, rows):
                program = str(row.get("项目", "") or "").strip()
                if program:
                    siblings.setdefault(key, []).append(program)

        items = []
        arbitration_owner = {}
        arbitration_key_by_idx = {}
        for idx, row in enumerate(rows):
            program = str(row.get("项目", "") or "").strip()
            # 语义复检不能只看附件文件名：公司法定名称经常只出现在 xlsx
            # 的结构化行/单元格里。阶段一已保存紧凑的附件证据快照，这里把
            # 当前行对应的快照传给 Agent，避免它在缺少附件内容时凭标题猜公司。
            attachment_evidence = str(row.get("附件证据", "") or "").strip()
            if len(attachment_evidence) > 9000:
                attachment_evidence = attachment_evidence[:9000]
            body_original = str(row.get("body_original", "") or "").strip()
            if len(body_original) > 6000:
                body_original = body_original[:6000]
            try:
                company_candidates = json.loads(str(row.get("客户候选证据") or "[]"))
            except (TypeError, ValueError):
                company_candidates = []
            if not isinstance(company_candidates, list):
                company_candidates = []
            company_candidates = [
                candidate for candidate in company_candidates
                if isinstance(candidate, dict) and candidate.get("id") and candidate.get("value")
            ]
            candidate_key = None
            if len(company_candidates) > 1:
                mail_key = mail_keys[idx] if mail_keys is not None and idx < len(mail_keys) else idx
                candidate_key = (
                    str(mail_key),
                    tuple((
                        str(candidate["id"]),
                        str(candidate["value"]),
                        str(candidate.get("source") or ""),
                    ) for candidate in company_candidates),
                )
                arbitration_key_by_idx[idx] = candidate_key
            item = {
                "id": str(idx),
                "sender_email": row.get("sender_email", ""),
                "subject": row.get("subject", ""),
                "body": row.get("body_text", ""),
                "body_original": body_original,
                "attachments": [x for x in str(row.get("附件名称", "") or "").split("；") if x],
                "attachment_evidence": attachment_evidence,
                "fields": {
                    "agent": row.get("代理", ""),
                    "company": row.get("客户", ""),
                    "country": country_of_project(program),
                    "program": program,
                    "request": row.get("需求", ""),
                },
                "rule_risks": row.get("业务规则问题", ""),
            }
            # 多项目同一主体只向模型提交一次公司候选仲裁；这次调用同时完成
            # 本行语义复检，后续行只复用候选结论，不重复做第二轮公司复检。
            if candidate_key is not None and candidate_key not in arbitration_owner:
                arbitration_owner[candidate_key] = idx
                item["company_candidates"] = company_candidates
            if mail_keys is not None and len(mail_keys) == len(rows) and program:
                # 同封邮件拆分出的全部项目（含本行）。LLM 拿到单行会误判"漏拆"，
                # 给它完整清单才能正确判断本行是不是拆分结果的一部分。
                item["same_mail_programs"] = list(siblings.get(mail_keys[idx], []))
            items.append(item)
        logger.info(f"LLM 语义字段校验开始: {len(rows)} 行，批大小={batch_size}")
        results = {}
        for start in range(0, len(items), batch_size):
            if self._stop:
                break
            chunk = items[start:start + batch_size]
            valid_ids = {str(x["id"]) for x in chunk}
            chunk_results = llm_client.validate_extracted_fields_batch(chunk)
            # JSON 修复重试偶尔会多吐几条不属于本批的记录，只认本批 id，
            # 免得把没校验过的行也标上结论。
            chunk_results = {k: v for k, v in chunk_results.items() if str(k) in valid_ids}
            results.update(chunk_results)
            logger.info(
                f"LLM 语义字段校验: {min(start + len(chunk), len(items))}/{len(items)} "
                f"（返回 {len(chunk_results)} 条）"
            )
        for idx, row in enumerate(rows):
            result = results.get(str(idx))
            if result is None:
                row["语义校验状态"] = "LLM调用失败/未完成"
                row["语义问题字段"] = "全部字段"
                row["语义问题编号"] = "AGENT_FAILED"
                row["语义问题证据"] = "模型未返回有效复检结果"
                row["语义当前值"] = ""
                row["语义建议值"] = ""
                row["语义校验建议"] = "保留程序原值，交人工核对"
                row["语义校验原因"] = "LLM未返回该条记录的校验结果"
                row["语义校验模型"] = llm_client.model
                row["置信度"] = "需人工确认"
                candidate_key = arbitration_key_by_idx.get(idx)
                if candidate_key is not None and arbitration_owner.get(candidate_key) == idx:
                    row["公司候选复核状态"] = "LLM未返回候选结论，需人工确认"
                    row["公司候选复核ID"] = ""
                continue
            issues = result.get("issues") or []
            row["语义校验状态"] = result.get("status", "uncertain")
            row["语义问题字段"] = "、".join(dict.fromkeys(str(x.get("field", "")) for x in issues if x.get("field")))
            row["语义问题编号"] = "、".join(dict.fromkeys(
                str(x.get("code", "OTHER") or "OTHER") for x in issues
            ))
            row["语义问题证据"] = "；".join(
                str(x.get("evidence", "")) for x in issues if x.get("evidence")
            )[:1000]
            row["语义当前值"] = "；".join(
                f"{x.get('field')}: {x.get('current_value')}"
                for x in issues if x.get("current_value")
            )[:800]
            row["语义建议值"] = "；".join(
                f"{x.get('field')}: {x.get('suggested_value')}"
                for x in issues if x.get("suggested_value")
            )[:800]
            # 主列表只显示短建议，完整证据保留在单独字段中。
            short_suggestions = [
                str(x.get("suggestion", "")).strip() for x in issues
                if str(x.get("suggestion", "")).strip()
            ]
            row["语义校验建议"] = "；".join(short_suggestions)[:500]
            row["语义校验原因"] = result.get("reason", "") or "；".join(
                str(x.get("reason", "")) for x in issues if x.get("reason")
            )
            row["语义校验模型"] = llm_client.model
            candidate_key = arbitration_key_by_idx.get(idx)
            if candidate_key is not None and arbitration_owner.get(candidate_key) == idx:
                candidates = []
                try:
                    candidates = json.loads(str(row.get("客户候选证据") or "[]"))
                except (TypeError, ValueError):
                    pass
                by_id = {
                    str(candidate.get("id")): candidate
                    for candidate in candidates if isinstance(candidate, dict) and candidate.get("id")
                } if isinstance(candidates, list) else {}
                selected_id = str(result.get("selected_company_candidate_id") or "")
                selected = by_id.get(selected_id)
                if selected and str(selected.get("value") or "").strip():
                    row["客户"] = str(selected["value"]).strip()
                    row["客户提取来源"] = "附件EPR申请表格式化字段（LLM候选仲裁）"
                    row["公司候选复核状态"] = "已仲裁"
                    row["公司候选复核ID"] = selected_id
                    row["公司候选复核证据"] = str(selected.get("evidence") or "")[:500]
                else:
                    row["公司候选复核状态"] = "未能唯一判定，需人工确认"
                    row["公司候选复核ID"] = ""
                    row["置信度"] = "需人工确认"
                    row["语义校验状态"] = "uncertain"
                    hint = "EPR申请表存在多个公司名候选，LLM未能从给定候选中唯一选择"
                    row["人工复核提示"] = "；".join(filter(None, [str(row.get("人工复核提示") or ""), hint]))
            if result.get("status") in {"uncertain", "invalid"}:
                row["置信度"] = "需人工确认"
                hint = "LLM语义校验"
                if row.get("语义问题字段"):
                    hint += f"（{row['语义问题字段']}）"
                if row.get("语义校验原因"):
                    hint += f": {row['语义校验原因']}"
                existing = str(row.get("人工复核提示", "") or "")
                row["人工复核提示"] = "；".join(x for x in [existing, hint] if x)
        # 同一申请人可能因多个项目拆成多行；候选仲裁只执行一次，其结果
        # 稳定复制到同主体的兄弟项目行，不再对每个项目重复选择。
        decisions = {}
        for key, owner_idx in arbitration_owner.items():
            if owner_idx < len(rows):
                owner = rows[owner_idx]
                decisions[key] = (
                    owner.get("客户"), owner.get("客户提取来源"),
                    owner.get("公司候选复核状态"), owner.get("公司候选复核ID"),
                    owner.get("公司候选复核证据"), owner.get("置信度"),
                    owner.get("人工复核提示"), owner.get("语义校验状态"),
                )
        for idx, row in enumerate(rows):
            key = arbitration_key_by_idx.get(idx)
            if key is None or arbitration_owner.get(key) == idx or key not in decisions:
                continue
            (company, source, status, candidate_id, evidence, confidence, hint, semantic_status) = decisions[key]
            if status == "已仲裁" and company:
                row["客户"] = company
                row["客户提取来源"] = source
                row["公司候选复核状态"] = status
                row["公司候选复核ID"] = candidate_id
                row["公司候选复核证据"] = evidence
            else:
                row["公司候选复核状态"] = status or "未能唯一判定，需人工确认"
                row["置信度"] = "需人工确认"
                row["语义校验状态"] = "uncertain"
                row["人工复核提示"] = "；".join(filter(None, [str(row.get("人工复核提示") or ""), str(hint or "EPR公司候选未唯一判定")]))
        return rows

    def _on_finished_stage2(self, primary_file, secondary_file):
        """阶段二完成信号占位（实际 finished_signal 在 run() 内 emit）"""
        pass

    def _log(self, msg):
        self.log_signal.emit(msg)

    def _load_internal_emails(self, path) -> Set[str]:
        """加载附件二: 内部邮箱表"""
        import logging
        logger = logging.getLogger("mail_audit")
        abs_path = str(resolve_runtime_path(path, "data/internal_email_cache.xlsx"))
        if not os.path.exists(abs_path):
            logger.warning(f"内部邮箱表文件不存在: {abs_path}")
            return set()
        from openpyxl import load_workbook
        wb = load_workbook(abs_path, read_only=True, data_only=True)
        ws = wb.active
        emails = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            for cell in row:
                if cell and "@" in str(cell):
                    emails.add(str(cell).strip().lower())
        wb.close()
        return emails

    def _load_agent_emails(self, path) -> Dict[str, dict]:
        """加载附件三: 代理邮箱表"""
        import logging
        logger = logging.getLogger("mail_audit")
        abs_path = str(resolve_runtime_path(path, "data/agent_emails.xlsx"))
        if not os.path.exists(abs_path):
            logger.warning(f"代理邮箱表文件不存在: {abs_path}")
            return {}
        from openpyxl import load_workbook
        wb = load_workbook(abs_path, read_only=True, data_only=True)
        ws = wb.active
        result = {}
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            agent = str(row[0]).strip() if row[0] else ""
            short = str(row[1]).strip() if len(row) > 1 and row[1] else ""
            email = str(row[2]).strip() if len(row) > 2 and row[2] else ""
            if email and "@" in email:
                # 一个代理可能有多个邮箱(分号分隔)
                for em in email.replace("；", ";").split(";"):
                    em = em.strip().lower()
                    if em:
                        result[em] = {"代理": agent, "代理简称": short, "收件人邮箱": email}
        wb.close()
        return result

    def _load_project_names(self, path) -> List[dict]:
        """加载附件四: 项目名称表"""
        import logging
        logger = logging.getLogger("mail_audit")
        abs_path = str(resolve_runtime_path(path, "data/project_names.xlsx"))
        if not os.path.exists(abs_path):
            logger.warning(f"项目名称表文件不存在: {abs_path}")
            return []
        from openpyxl import load_workbook
        wb = load_workbook(abs_path, read_only=True, data_only=True)
        ws = wb.active
        result = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            entry = {}
            keys = ["项目编号", "项目名称", "国家", "业务类型"]
            for idx, key in enumerate(keys):
                if idx < len(row) and row[idx]:
                    entry[key] = str(row[idx]).strip()
            if entry.get("项目名称"):
                result.append(entry)
        wb.close()
        if not result:
            logger.warning(f"项目名称表为空(0 条有效记录): {abs_path}")
        return result


class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.config = None
        self.worker = None
        self.init_ui()
        self.load_config()
        self.load_session()

    def init_ui(self):
        self.setWindowTitle("邮件&系统漏单审核自动化工具 v1.0")
        self.setMinimumSize(1120, 760)
        self.resize(1280, 860)
        self.setStyleSheet("""
            QWidget { font-family: 'Microsoft YaHei UI', 'Microsoft YaHei', sans-serif; font-size: 10pt; color: #243447; }
            QMainWindow, QWidget#centralWidget { background: #f4f7fb; }
            QGroupBox { background: #ffffff; border: 1px solid #dbe3ed; border-radius: 10px; margin-top: 12px; padding: 14px 12px 12px; font-weight: 600; color: #1f3448; }
            QGroupBox::title { subcontrol-origin: margin; left: 14px; padding: 0 6px; background: #f4f7fb; }
            QLabel#pageTitle { font-size: 18pt; font-weight: 700; color: #16324a; }
            QLabel#pageSubtitle { color: #6b7d8f; font-size: 9pt; }
            QLabel#sourceStatus { color: #6b7d8f; }
            QLineEdit, QDateEdit, QComboBox { background: #fbfcfe; border: 1px solid #cbd7e4; border-radius: 6px; padding: 6px 8px; min-height: 20px; }
            QLineEdit:focus, QDateEdit:focus, QComboBox:focus { border: 1px solid #2f80ed; }
            QPushButton { background: #ffffff; border: 1px solid #cbd7e4; border-radius: 6px; padding: 7px 14px; min-height: 22px; color: #243447; }
            QPushButton:hover { background: #eef5fc; border-color: #2f80ed; }
            QPushButton:disabled { background: #eef1f5; color: #9aa8b5; border-color: #dce3ea; }
            QPushButton#primaryButton { background: #1677d2; color: #ffffff; border: none; font-weight: 700; }
            QPushButton#primaryButton:hover { background: #0f65b7; }
            QPushButton#dangerButton { background: #d9534f; color: #ffffff; border: none; font-weight: 600; }
            QPushButton#dangerButton:hover { background: #bd3f3b; }
            QPushButton#secondaryButton { background: #eef5fc; color: #1769aa; border-color: #b9d3eb; }
            QPushButton#secondaryButton:hover { background: #dcecfb; }
            QRadioButton, QCheckBox { spacing: 6px; color: #40566a; }
            QProgressBar { background: #e8eef5; border: none; border-radius: 5px; height: 12px; text-align: center; color: #234; }
            QProgressBar::chunk { background: #2f80ed; border-radius: 5px; }
            QTextEdit { background: #172432; color: #d9e7f2; border: 1px solid #243b50; border-radius: 8px; padding: 8px; }
        """)

        central = QWidget()
        central.setObjectName("centralWidget")
        self.setCentralWidget(central)
        layout = QVBoxLayout(central)
        layout.setSpacing(10)
        layout.setContentsMargins(15, 15, 15, 15)

        # === 顶部控制栏：按“运行模式 → 数据源 → 执行”分组，避免所有控件挤在一行 ===
        header = QHBoxLayout()
        header.setContentsMargins(4, 0, 4, 0)
        title_box = QVBoxLayout()
        page_title = QLabel("邮件审核工作台")
        page_title.setObjectName("pageTitle")
        page_subtitle = QLabel("阶段一邮件解析 · 人工复核 · 阶段二工单比对")
        page_subtitle.setObjectName("pageSubtitle")
        title_box.addWidget(page_title)
        title_box.addWidget(page_subtitle)
        header.addLayout(title_box)
        header.addStretch()

        mode_group = QGroupBox("运行模式")
        mode_layout = QHBoxLayout(mode_group)
        mode_layout.setContentsMargins(8, 4, 8, 4)
        self.radio_all = QRadioButton("一站式")
        self.radio_stage1 = QRadioButton("阶段一：邮件解析")
        self.radio_stage2 = QRadioButton("阶段二：工单核对")
        self.radio_all.setToolTip("邮件解析完成后继续执行工单比对")
        self.radio_stage1.setToolTip("只拉取和解析邮件，不访问工单系统")
        self.radio_stage2.setToolTip("读取阶段一工单待查 Excel，不重新读取邮箱")
        self.radio_all.setChecked(True)
        mode_btn_group = QButtonGroup(self)
        mode_btn_group.addButton(self.radio_all, 0)
        mode_btn_group.addButton(self.radio_stage1, 1)
        mode_btn_group.addButton(self.radio_stage2, 2)
        mode_layout.addWidget(self.radio_all)
        mode_layout.addWidget(self.radio_stage1)
        mode_layout.addWidget(self.radio_stage2)
        header.addWidget(mode_group)
        layout.addLayout(header)

        ctrl_group = QGroupBox("运行参数")
        ctrl_layout = QVBoxLayout(ctrl_group)
        ctrl_layout.setSpacing(10)

        self.date_row_widget = QWidget()
        date_row = QHBoxLayout(self.date_row_widget)
        date_row.setContentsMargins(0, 0, 0, 0)
        self.date_range_label = QLabel("邮件时间范围")
        date_row.addWidget(self.date_range_label)
        self.date_from = QDateEdit()
        self.date_from.setDate(QDate.currentDate().addDays(-30))
        self.date_from.setDisplayFormat("yyyy-MM-dd")
        self.date_from.setCalendarPopup(True)
        date_row.addWidget(self.date_from)
        self.date_to_label = QLabel("至")
        date_row.addWidget(self.date_to_label)
        self.date_to = QDateEdit()
        self.date_to.setDate(QDate.currentDate())
        self.date_to.setDisplayFormat("yyyy-MM-dd")
        self.date_to.setCalendarPopup(True)
        date_row.addWidget(self.date_to)
        date_row.addStretch()
        ctrl_layout.addWidget(self.date_row_widget)

        self.sources_group = QGroupBox("阶段一数据源")
        sources_layout = QHBoxLayout(self.sources_group)
        sources_layout.setContentsMargins(8, 4, 8, 4)

        self.btn_import_agent = QPushButton("导入代理邮箱表")
        self.btn_import_agent.clicked.connect(self.import_agent_table)
        self.agent_table_label = QLabel("未导入")
        self.agent_table_label.setObjectName("sourceStatus")
        agent_box = QVBoxLayout()
        agent_box.addWidget(self.btn_import_agent)
        agent_box.addWidget(self.agent_table_label)
        sources_layout.addLayout(agent_box, 1)

        self.btn_import_internal = QPushButton("导入内部邮箱表")
        self.btn_import_internal.clicked.connect(self.import_internal_table)
        self.internal_table_label = QLabel("未导入")
        self.internal_table_label.setObjectName("sourceStatus")
        internal_box = QVBoxLayout()
        internal_box.addWidget(self.btn_import_internal)
        internal_box.addWidget(self.internal_table_label)
        sources_layout.addLayout(internal_box, 1)

        self.btn_import_project = QPushButton("导入项目名称表")
        self.btn_import_project.clicked.connect(self.import_project_table)
        self.project_table_label = QLabel("未导入")
        self.project_table_label.setObjectName("sourceStatus")
        project_box = QVBoxLayout()
        project_box.addWidget(self.btn_import_project)
        project_box.addWidget(self.project_table_label)
        sources_layout.addLayout(project_box, 1)
        ctrl_layout.addWidget(self.sources_group)

        # === 阶段二输入文件选择 (仅 stage2 模式可见) ===
        self.stage2_panel = QGroupBox("阶段二输入")
        stage2_layout = QHBoxLayout(self.stage2_panel)
        stage2_layout.setContentsMargins(8, 4, 8, 4)
        self.stage2_input_label = QLabel("未选择阶段一工单待查文件")
        self.stage2_input_label.setObjectName("sourceStatus")
        stage2_layout.addWidget(self.stage2_input_label)
        self.stage2_input_combo = QComboBox()
        self.stage2_input_combo.setMinimumWidth(420)
        self.stage2_input_combo.setToolTip(
            "选择阶段一生成的‘工单待查’文件；列表按文件更新时间排列"
        )
        self.stage2_input_combo.currentIndexChanged.connect(
            self._on_stage2_input_combo_changed
        )
        stage2_layout.addWidget(self.stage2_input_combo, 1)
        self.btn_refresh_stage2_inputs = QPushButton("刷新")
        self.btn_refresh_stage2_inputs.clicked.connect(self.refresh_stage2_input_options)
        stage2_layout.addWidget(self.btn_refresh_stage2_inputs)
        self.btn_choose_stage2_input = QPushButton("浏览文件")
        self.btn_choose_stage2_input.clicked.connect(self.choose_stage2_input)
        stage2_layout.addWidget(self.btn_choose_stage2_input)
        self.stage2_input_combo.setVisible(False)
        self.btn_refresh_stage2_inputs.setVisible(False)
        self.btn_choose_stage2_input.setVisible(False)
        self.stage2_input_label.setVisible(False)
        self.chk_force_live = QCheckBox("强制实时查询（忽略缓存）")
        self.chk_force_live.setToolTip("仍使用阶段一待查名单，但每条都会实际填写并点击网页查询")
        stage2_layout.addWidget(self.chk_force_live)
        ctrl_layout.addWidget(self.stage2_panel)
        self.stage2_panel.setVisible(False)
        # 模式切换时显示/隐藏
        self.radio_all.toggled.connect(self._on_mode_changed)
        self.radio_stage1.toggled.connect(self._on_mode_changed)
        self.radio_stage2.toggled.connect(self._on_mode_changed)

        action_row = QHBoxLayout()
        action_row.addStretch()
        self.btn_run = QPushButton("开始运行")
        self.btn_run.setObjectName("primaryButton")
        self.btn_run.setMinimumWidth(130)
        self.btn_run.clicked.connect(self.start_run)
        action_row.addWidget(self.btn_run)

        self.btn_stop = QPushButton("停止")
        self.btn_stop.setObjectName("dangerButton")
        self.btn_stop.setMinimumWidth(100)
        self.btn_stop.setEnabled(False)
        self.btn_stop.clicked.connect(self.stop_run)
        action_row.addWidget(self.btn_stop)
        ctrl_layout.addLayout(action_row)

        layout.addWidget(ctrl_group)

        # === 凭据配置区 ===
        cred_group = QGroupBox("账号配置 (人员变动时可直接修改并保存)")
        cred_layout = QFormLayout(cred_group)

        # 阿里邮箱
        self.input_email_addr = QLineEdit()
        self.input_email_pwd = QLineEdit()
        self.input_email_pwd.setEchoMode(QLineEdit.Password)
        cred_layout.addRow("阿里邮箱账号:", self.input_email_addr)
        cred_layout.addRow("阿里邮箱密码:", self.input_email_pwd)

        # 工单系统
        self.input_wo_user = QLineEdit()
        self.input_wo_pwd = QLineEdit()
        self.input_wo_pwd.setEchoMode(QLineEdit.Password)
        cred_layout.addRow("工单系统账号:", self.input_wo_user)
        cred_layout.addRow("工单系统密码:", self.input_wo_pwd)

        # 保存按钮
        btn_save_cred = QPushButton("保存配置到 config.yaml")
        btn_save_cred.setObjectName("secondaryButton")
        btn_save_cred.clicked.connect(self.save_credentials)
        cred_layout.addRow(btn_save_cred)

        layout.addWidget(cred_group)

        # === 进度条 ===
        self.progress = QProgressBar()
        self.progress.setValue(0)
        self.progress.setFormat("准备中...")
        layout.addWidget(self.progress)

        # === 日志面板 ===
        log_group = QGroupBox("运行日志")
        log_layout = QVBoxLayout(log_group)
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        self.log_text.setFont(QFont("Consolas", 9))
        self.log_text.setStyleSheet(
            "QTextEdit { background-color: #1e1e1e; color: #d4d4d4; }"
        )
        log_layout.addWidget(self.log_text)
        layout.addWidget(log_group, stretch=1)

        # === 结果与人工处理 ===
        # 将结果入口和工作台入口放进独立分组，避免像一排孤立按钮。
        results_group = QGroupBox("结果与人工处理")
        bottom_layout = QHBoxLayout(results_group)
        bottom_layout.setContentsMargins(8, 4, 8, 4)
        bottom_layout.setSpacing(8)

        self.btn_open_missing = QPushButton("查看待查名单")
        self.btn_open_missing.setEnabled(False)
        self.btn_open_missing.clicked.connect(lambda: self.open_excel(self.missing_file))
        bottom_layout.addWidget(self.btn_open_missing)

        self.btn_open_filtered = QPushButton("查看人工复核名单")
        self.btn_open_filtered.setEnabled(False)
        self.btn_open_filtered.clicked.connect(lambda: self.open_excel(self.filtered_file))
        bottom_layout.addWidget(self.btn_open_filtered)

        self.btn_open_stage2_result = QPushButton("查看核对结果")
        self.btn_open_stage2_result.setEnabled(False)
        self.btn_open_stage2_result.clicked.connect(lambda: self.open_excel(self.stage2_result_file))
        bottom_layout.addWidget(self.btn_open_stage2_result)

        self.btn_load_stage1_cache = QPushButton("读取阶段一缓存")
        self.btn_load_stage1_cache.setToolTip(
            "按邮件日期选择已经保存的阶段一结果，不重新连接邮箱"
        )
        self.btn_load_stage1_cache.setEnabled(False)
        self.btn_load_stage1_cache.clicked.connect(self.load_stage1_cache)
        bottom_layout.addWidget(self.btn_load_stage1_cache)

        # 可视化人工复核工作台：始终使用正式持久历史，不向操作人员暴露测试会话切换。
        self.btn_open_workbench = QPushButton("打开人工复核工作台")
        self.btn_open_workbench.setToolTip(
            "读取阶段一待查/人工补全结果，在浏览器中进行证据核对和人工确认"
        )
        self.btn_open_workbench.setEnabled(False)
        self.btn_open_workbench.clicked.connect(self.open_workbench)
        bottom_layout.addWidget(self.btn_open_workbench)

        bottom_layout.addStretch()
        btn_exit = QPushButton("退出")
        btn_exit.clicked.connect(self.close)
        bottom_layout.addWidget(btn_exit)
        layout.addWidget(results_group)

        self.missing_file = ""
        self.filtered_file = ""
        self.stage2_result_file = ""
        self.stage2_input_path = None
        self.stage2_input_user_selected = False
        self.stage1_cache_selection = None
        self.workbench_server = None
        # 允许打开上一次阶段一结果；若不存在，运行阶段一后自动启用。
        output_root = os.path.join(APP_DIR, "output")
        categorized_primary = os.path.join(
            output_root, "stage1_email", "to_workorder_list.xlsx"
        )
        legacy_primary = os.path.join(output_root, "to_workorder_list.xlsx")
        self.btn_open_workbench.setEnabled(
            os.path.exists(categorized_primary) or os.path.exists(legacy_primary)
        )
        self._refresh_stage1_cache_button()
        # 统一初始化控件可见性。
        self._on_mode_changed()

    def load_config(self):
        config_path = os.path.join(APP_DIR, "config.yaml")
        if os.path.exists(config_path):
            with open(config_path, "r", encoding="utf-8") as f:
                self.config = yaml.safe_load(f)
        else:
            QMessageBox.critical(self, "错误", "找不到 config.yaml 配置文件!")
            self.config = {}
        self._normalize_config_paths()
        self.load_credentials()

    def _normalize_config_paths(self):
        """启动时迁移旧版绝对路径，避免换电脑后仍引用开发机目录。"""
        if not isinstance(self.config, dict):
            self.config = {}
        reference = self.config.setdefault("reference_tables", {})
        defaults = {
            "internal_emails": "data/internal_emails.xlsx",
            "agent_emails": "data/agent_emails.xlsx",
            "project_names": "data/project_names.xlsx",
        }
        for key, fallback in defaults.items():
            reference[key] = str(resolve_runtime_path(reference.get(key), fallback))

        output = self.config.setdefault("output", {})
        output["dir"] = str(resolve_runtime_path(output.get("dir"), "output"))
        logging_cfg = self.config.setdefault("logging", {})
        logging_cfg["dir"] = str(resolve_runtime_path(logging_cfg.get("dir"), "logs"))

        email_cfg = self.config.setdefault("email", {})
        email_cfg["cache_dir"] = str(resolve_runtime_path(email_cfg.get("cache_dir"), "cache/mails"))
        workorder = self.config.setdefault("workorder", {})
        workorder["query_cache_path"] = str(
            resolve_runtime_path(workorder.get("query_cache_path"), "storage/query_cache.json")
        )
        workorder["login_state_path"] = str(
            resolve_runtime_path(workorder.get("login_state_path"), "storage/login_state.json")
        )
        llm_cfg = self.config.setdefault("llm", {})
        llm_cfg["audit_log_path"] = str(
            resolve_runtime_path(llm_cfg.get("audit_log_path"), "logs/llm_calls.jsonl")
        )

    def _refresh_stage1_cache_button(self):
        """有日期缓存或旧版稳定输出时，允许操作人员读取阶段一结果。"""
        has_cache = bool(_stage1_cache_options())
        output_root = Path(APP_DIR) / "output"
        has_legacy = any(
            (output_root / relative).is_file()
            for relative in (
                "stage1_email/to_workorder_list.xlsx",
                "stage1_email/to_review_list.xlsx",
                "to_workorder_list.xlsx",
            )
        )
        self.btn_load_stage1_cache.setEnabled(has_cache or has_legacy)

    @staticmethod
    def _set_date_widget(widget, value):
        """把缓存清单中的日期安全地恢复到日期控件。"""
        text = str(value or "")[:10]
        if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", text):
            return
        try:
            year, month, day = (int(part) for part in text.split("-"))
            widget.setDate(QDate(year, month, day))
        except (TypeError, ValueError):
            return

    def load_session(self):
        """加载上次保存的时间范围、参考表路径和阶段二输入选择。"""
        if not os.path.exists(SESSION_FILE):
            return
        try:
            with open(SESSION_FILE, "r", encoding="utf-8") as f:
                state = json.load(f)
            # 恢复时间范围
            date_from_str = state.get("date_from")
            date_to_str = state.get("date_to")
            if date_from_str:
                parts = date_from_str.split("-")
                self.date_from.setDate(QDate(int(parts[0]), int(parts[1]), int(parts[2])))
            if date_to_str:
                parts = date_to_str.split("-")
                self.date_to.setDate(QDate(int(parts[0]), int(parts[1]), int(parts[2])))
            # 恢复上次选择的阶段一日期缓存；文件不存在时保留当前默认结果，
            # 不把旧电脑上的绝对路径继续带入新环境。
            cached_primary = resolve_runtime_path(state.get("stage1_cache_primary"), None)
            cached_review = resolve_runtime_path(state.get("stage1_cache_review"), None)
            if cached_primary.is_file() or cached_review.is_file():
                self.missing_file = str(cached_primary) if cached_primary.is_file() else ""
                self.filtered_file = str(cached_review) if cached_review.is_file() else ""
                self.btn_open_missing.setEnabled(bool(self.missing_file))
                self.btn_open_filtered.setEnabled(bool(self.filtered_file))
                self.btn_open_workbench.setEnabled(bool(self.missing_file or self.filtered_file))
                self.stage1_cache_selection = {
                    "primary": self.missing_file,
                    "review": self.filtered_file,
                }
            # 恢复邮箱汇总表路径
            agent_path = state.get("agent_table_path")
            agent_resolved = resolve_runtime_path(agent_path, "data/agent_emails.xlsx")
            if agent_resolved.is_file():
                # 旧会话可能指向开发机桌面；将仍可访问的外部参考表复制进当前包。
                if APP_ROOT not in agent_resolved.parents:
                    agent_resolved = copy_reference_file(agent_resolved, "agent_emails.xlsx")
                self.agent_email_path = str(agent_resolved)
                fname = agent_resolved.name
                self.agent_table_label.setText(f"已导入: {fname}")
                self.agent_table_label.setStyleSheet("color: #4CAF50;")
                self.log(f"已恢复上次邮箱汇总表: {fname}")
            # 恢复内部邮箱表路径
            internal_path = state.get("internal_table_path")
            internal_resolved = None
            # 缓存优先：即使旧会话保存的是原始文件路径，也展示并使用本地规范化副本。
            if os.path.exists(INTERNAL_EMAIL_CACHE_FILE):
                self.internal_email_path = INTERNAL_EMAIL_CACHE_FILE
                internal_resolved = Path(INTERNAL_EMAIL_CACHE_FILE).resolve()
                fname = os.path.basename(INTERNAL_EMAIL_CACHE_FILE)
                self.internal_table_label.setText(f"已缓存: {fname}")
                self.internal_table_label.setStyleSheet("color: #4CAF50;")
                self.log(f"已恢复本地内部邮箱缓存: {fname}")
            elif internal_path:
                internal_resolved = resolve_runtime_path(internal_path, "data/internal_emails.xlsx")
                if internal_resolved.is_file():
                    self.internal_email_path = str(internal_resolved)
                    fname = internal_resolved.name
                else:
                    fname = ""
                if not fname:
                    internal_resolved = None
            else:
                internal_resolved = None
            if internal_path and internal_resolved and not os.path.exists(INTERNAL_EMAIL_CACHE_FILE):
                self.internal_table_label.setText(f"已导入: {fname}")
                self.internal_table_label.setStyleSheet("color: #4CAF50;")
                self.log(f"已恢复上次内部邮箱表: {fname}")
            # 恢复项目表路径
            project_path = state.get("project_table_path")
            project_resolved = resolve_runtime_path(project_path, "data/project_names.xlsx")
            if project_resolved.is_file():
                if APP_ROOT not in project_resolved.parents:
                    project_resolved = copy_reference_file(project_resolved, "project_names.xlsx")
                self.project_table_path = str(project_resolved)
                fname = project_resolved.name
                self.project_table_label.setText(f"已导入: {fname}")
                self.project_table_label.setStyleSheet("color: #4CAF50;")
                self.log(f"已恢复上次项目表: {fname}")
            # 阶段二输入不再无条件锁定上次的时序副本。启动时先扫描
            # 当前 output 中最新的阶段一“工单待查”文件；只有用户曾经
            # 明确选择过某个文件时，才恢复该选择。
            stage2_in = state.get("stage2_input_path")
            self.stage2_input_user_selected = bool(
                state.get("stage2_input_user_selected", False)
            )
            self.refresh_stage2_input_options(
                preferred_path=(stage2_in if self.stage2_input_user_selected else None)
            )
            self.chk_force_live.setChecked(bool(state.get("force_live_query", False)))
            # 恢复运行模式
            mode = state.get("mode", "all")
            if mode == "stage1":
                self.radio_stage1.setChecked(True)
            elif mode == "stage2":
                self.radio_stage2.setChecked(True)
            else:
                self.radio_all.setChecked(True)
            self._on_mode_changed()
        except Exception as e:
            self.log(f"[WARNING] 恢复会话状态失败: {e}")

    def save_session(self):
        """保存当前时间范围、参考表路径和阶段二输入选择"""
        if self.radio_stage2.isChecked():
            mode = "stage2"
        elif self.radio_stage1.isChecked():
            mode = "stage1"
        else:
            mode = "all"
        state = {
            "date_from": self.date_from.date().toString("yyyy-MM-dd"),
            "date_to": self.date_to.date().toString("yyyy-MM-dd"),
            "agent_table_path": app_relative_path(
                getattr(self, "agent_email_path", None), "data/agent_emails.xlsx"
            ),
            "internal_table_path": app_relative_path(
                getattr(self, "internal_email_path", None), "data/internal_email_cache.xlsx"
            ),
            "project_table_path": app_relative_path(
                getattr(self, "project_table_path", None), "data/project_names.xlsx"
            ),
            "stage2_input_path": app_relative_path(
                getattr(self, "stage2_input_path", None), None
            ) if getattr(self, "stage2_input_path", None) else None,
            "stage2_input_user_selected": bool(
                getattr(self, "stage2_input_user_selected", False)
            ),
            "force_live_query": self.chk_force_live.isChecked(),
            "stage1_cache_primary": (
                app_relative_path(self.missing_file, None)
                if getattr(self, "missing_file", "") and os.path.exists(self.missing_file)
                else None
            ),
            "stage1_cache_review": (
                app_relative_path(self.filtered_file, None)
                if getattr(self, "filtered_file", "") and os.path.exists(self.filtered_file)
                else None
            ),
            "mode": mode,
        }
        try:
            temporary = SESSION_FILE + ".tmp"
            with open(temporary, "w", encoding="utf-8") as f:
                json.dump(state, f, ensure_ascii=False, indent=2)
            os.replace(temporary, SESSION_FILE)
            return True
        except Exception as exc:
            self.log(f"[WARNING] 保存会话失败: {exc}")
            return False

    def load_credentials(self):
        """从 config 加载凭据到界面输入框"""
        if not self.config:
            return
        email_cfg = self.config.get("email", {})
        self.input_email_addr.setText(email_cfg.get("address", ""))
        self.input_email_pwd.setText(email_cfg.get("password", ""))

        wo_cfg = self.config.get("workorder", {})
        self.input_wo_user.setText(wo_cfg.get("username", ""))
        self.input_wo_pwd.setText(wo_cfg.get("password", ""))

    def save_credentials(self):
        """将界面输入的凭据保存到 config.yaml 并更新内存配置"""
        if not self.config:
            QMessageBox.critical(self, "错误", "配置文件未加载!")
            return

        self.config.setdefault("email", {})
        self.config["email"]["address"] = self.input_email_addr.text().strip()
        self.config["email"]["password"] = self.input_email_pwd.text().strip()

        self.config.setdefault("workorder", {})
        self.config["workorder"]["username"] = self.input_wo_user.text().strip()
        self.config["workorder"]["password"] = self.input_wo_pwd.text().strip()

        # 配置文件是可迁移文件，不能把当前电脑的绝对路径写回去。
        reference = self.config.setdefault("reference_tables", {})
        reference["internal_emails"] = app_relative_path(
            reference.get("internal_emails"), "data/internal_emails.xlsx"
        )
        reference["agent_emails"] = app_relative_path(
            reference.get("agent_emails"), "data/agent_emails.xlsx"
        )
        reference["project_names"] = app_relative_path(
            reference.get("project_names"), "data/project_names.xlsx"
        )
        self.config.setdefault("output", {})["dir"] = app_relative_path(
            self.config["output"].get("dir"), "output"
        )
        self.config.setdefault("logging", {})["dir"] = app_relative_path(
            self.config["logging"].get("dir"), "logs"
        )
        self.config.setdefault("email", {})["cache_dir"] = app_relative_path(
            self.config["email"].get("cache_dir"), "cache/mails"
        )
        self.config.setdefault("workorder", {})["query_cache_path"] = app_relative_path(
            self.config["workorder"].get("query_cache_path"), "storage/query_cache.json"
        )
        self.config["workorder"]["login_state_path"] = app_relative_path(
            self.config["workorder"].get("login_state_path"), "storage/login_state.json"
        )
        self.config.setdefault("llm", {})["audit_log_path"] = app_relative_path(
            self.config["llm"].get("audit_log_path"), "logs/llm_calls.jsonl"
        )

        config_path = os.path.join(APP_DIR, "config.yaml")
        try:
            with open(config_path, "w", encoding="utf-8") as f:
                yaml.dump(self.config, f, allow_unicode=True, default_flow_style=False, sort_keys=False)
            self.log("凭据已保存到 config.yaml")
            QMessageBox.information(self, "保存成功", "账号配置已保存到 config.yaml")
        except Exception as e:
            self.log(f"[ERROR] 保存配置失败: {e}")
            QMessageBox.critical(self, "保存失败", str(e))

    def _is_stage2_workorder_file(self, path):
        """只把阶段一或人工复核后仍保留“工单待查”表的工作簿放入候选列表。"""
        if not path or not os.path.isfile(path):
            return False
        try:
            from openpyxl import load_workbook

            wb = load_workbook(path, read_only=True, data_only=True)
            if "工单待查" not in wb.sheetnames:
                wb.close()
                return False
            ws = wb["工单待查"]
            headers = {
                str(value).strip()
                for value in next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
                if value is not None
            }
            data_row = next(
                ws.iter_rows(min_row=2, max_row=2, values_only=True), None
            )
            wb.close()
            if not data_row or not any(
                value is not None and str(value).strip() for value in data_row
            ):
                return False
            return {"代理", "客户公司名称", "标准化项目名称"}.issubset(headers)
        except Exception:
            return False

    def _stage2_candidate_paths(self):
        """返回当前项目中可供阶段二选择的阶段一输出文件。"""
        output_dir = (self.config or {}).get("output", {}).get("dir", "output")
        if not os.path.isabs(output_dir):
            output_dir = os.path.join(APP_DIR, output_dir)
        if not os.path.isdir(output_dir):
            return []

        candidates = []
        # 阶段一结果现在位于 output/stage1_email；同时递归兼容旧版 output 根目录。
        for root, _dirs, names in os.walk(output_dir):
            for name in names:
                lower_name = name.lower()
                # 原始阶段一输出和人工确认后的工作台输出都可以作为阶段二输入。
                if not (
                    lower_name.startswith("to_workorder_list")
                    or lower_name.startswith("workbench_reviewed")
                ) or not lower_name.endswith((".xlsx", ".xls")):
                    continue
                path = os.path.abspath(os.path.join(root, name))
                if self._is_stage2_workorder_file(path):
                    candidates.append(path)
        # 稳定名是阶段一的 canonical 输出；同一批次时间相同的时候优先它，
        # 时序副本仍保留在列表中供人工选择历史批次。
        candidates.sort(
            key=lambda p: (
                -os.path.getmtime(p),
                os.path.basename(p).lower() != "to_workorder_list.xlsx",
            )
        )
        return candidates

    def _set_stage2_input_path(self, path, user_selected=False, save=False):
        self.stage2_input_path = path if path and os.path.exists(path) else None
        if self.stage2_input_path:
            fname = os.path.basename(self.stage2_input_path)
            self.stage2_input_label.setText(f"阶段二输入: {fname}")
            self.stage2_input_label.setStyleSheet("color: #4CAF50;")
            self.stage2_input_combo.setToolTip(self.stage2_input_path)
        else:
            self.stage2_input_label.setText("阶段二输入: 未选择")
            self.stage2_input_label.setStyleSheet("color: #999;")
            self.stage2_input_combo.setToolTip(
                "请选择阶段一生成的‘工单待查’文件"
            )
        if user_selected:
            self.stage2_input_user_selected = True
            self.log(
                f"阶段二输入文件已选: "
                f"{os.path.basename(self.stage2_input_path) if self.stage2_input_path else '未选择'}"
            )
        if save:
            self.save_session()

    def refresh_stage2_input_options(self, preferred_path=None):
        """刷新阶段一输出候选，并默认选择最新文件而不是历史旧副本。"""
        candidates = self._stage2_candidate_paths()
        preferred = os.path.abspath(preferred_path) if preferred_path else None
        if preferred and os.path.exists(preferred) and preferred not in candidates:
            if self._is_stage2_workorder_file(preferred):
                candidates.append(preferred)

        self.stage2_input_combo.blockSignals(True)
        self.stage2_input_combo.clear()
        if not candidates:
            self.stage2_input_combo.addItem("暂无可用的工单待查文件", "")
            self.stage2_input_combo.setCurrentIndex(0)
            self.stage2_input_combo.blockSignals(False)
            self._set_stage2_input_path(None)
            self.log("未找到可用的阶段一工单待查文件，请先运行阶段一或手动选择文件", "warning")
            return

        for path in candidates:
            stamp = datetime.fromtimestamp(os.path.getmtime(path)).strftime("%Y-%m-%d %H:%M")
            self.stage2_input_combo.addItem(
                f"{os.path.basename(path)}  ({stamp})", path
            )
        selected_path = preferred if preferred in candidates else candidates[0]
        selected_index = candidates.index(selected_path)
        self.stage2_input_combo.setCurrentIndex(selected_index)
        self.stage2_input_combo.blockSignals(False)
        self._set_stage2_input_path(selected_path, user_selected=bool(preferred))
        self.log(
            f"阶段二输入候选已刷新，当前选择: {os.path.basename(selected_path)}"
        )

    def _on_stage2_input_combo_changed(self, index):
        if index < 0:
            return
        path = self.stage2_input_combo.itemData(index)
        if not path:
            self._set_stage2_input_path(None)
            return
        self._set_stage2_input_path(path, user_selected=True, save=True)

    def choose_stage2_input(self):
        path, _ = QFileDialog.getOpenFileName(
            self, "选择阶段一产出的工单待查文件", "", "Excel Files (*.xlsx *.xls)"
        )
        if path:
            if not self._is_stage2_workorder_file(path):
                QMessageBox.warning(
                    self,
                    "文件不符合要求",
                    "请选择阶段一输出的 Excel，并且工作表必须为“工单待查”。",
                )
                return
            path = os.path.abspath(path)
            try:
                if APP_ROOT not in Path(path).resolve().parents:
                    path = str(copy_imported_file(path))
            except OSError:
                pass
            existing = [
                self.stage2_input_combo.itemData(i)
                for i in range(self.stage2_input_combo.count())
            ]
            if path not in existing:
                self.stage2_input_combo.addItem(os.path.basename(path), path)
            self.stage2_input_combo.setCurrentIndex(
                [self.stage2_input_combo.itemData(i) for i in range(self.stage2_input_combo.count())].index(path)
            )
            self._set_stage2_input_path(path, user_selected=True, save=True)

    def _on_mode_changed(self):
        is_stage2 = self.radio_stage2.isChecked()
        self.stage2_panel.setVisible(is_stage2)
        self.date_row_widget.setVisible(not is_stage2)
        self.sources_group.setVisible(not is_stage2)
        self.btn_choose_stage2_input.setVisible(is_stage2)
        self.stage2_input_combo.setVisible(is_stage2)
        self.btn_refresh_stage2_inputs.setVisible(is_stage2)
        self.stage2_input_label.setVisible(is_stage2)
        self.chk_force_live.setVisible(is_stage2)
        self.btn_import_agent.setVisible(not is_stage2)
        self.agent_table_label.setVisible(not is_stage2)
        self.btn_import_internal.setVisible(not is_stage2)
        self.internal_table_label.setVisible(not is_stage2)
        self.btn_import_project.setVisible(not is_stage2)
        self.project_table_label.setVisible(not is_stage2)
        self.btn_run.setText("开始运行")
        if is_stage2:
            self.btn_run.setText("开始工单核对")

    def import_agent_table(self):
        """导入邮箱汇总表"""
        path, _ = QFileDialog.getOpenFileName(
            self, "选择邮箱汇总表 Excel 文件", "", "Excel Files (*.xlsx *.xls)"
        )
        if path:
            local_path = copy_reference_file(path, "agent_emails.xlsx")
            self.agent_email_path = str(local_path)
            self.config.setdefault("reference_tables", {})["agent_emails"] = "data/agent_emails.xlsx"
            fname = local_path.name
            self.agent_table_label.setText(f"已导入: {fname}")
            self.agent_table_label.setStyleSheet("color: #4CAF50;")
            self.log(f"邮箱汇总表已导入: {fname}")
            self.save_session()

    def import_internal_table(self):
        """导入附件二: 内部邮箱表"""
        path, _ = QFileDialog.getOpenFileName(
            self, "选择内部邮箱表 Excel 文件", "", "Excel Files (*.xlsx *.xls)"
        )
        if path:
            try:
                cache_path = _cache_internal_email_table(path)
            except Exception as exc:
                QMessageBox.warning(self, "内部邮箱表导入失败", f"未能生成本地缓存，原有规则未改变：\n{exc}")
                self.log(f"[WARNING] 内部邮箱表缓存失败: {exc}")
                return
            self.internal_email_path = cache_path
            self.config.setdefault("reference_tables", {})["internal_emails"] = "data/internal_email_cache.xlsx"
            fname = os.path.basename(path)
            self.internal_table_label.setText(f"已缓存: {os.path.basename(cache_path)}（来源：{fname}）")
            self.internal_table_label.setStyleSheet("color: #4CAF50;")
            self.log(f"内部邮箱表已更新并缓存: {cache_path}")
            self.save_session()

    def import_project_table(self):
        """导入附件四: 项目名称表"""
        path, _ = QFileDialog.getOpenFileName(
            self, "选择项目名称表 Excel 文件", "", "Excel Files (*.xlsx *.xls)"
        )
        if path:
            local_path = copy_reference_file(path, "project_names.xlsx")
            self.project_table_path = str(local_path)
            self.config.setdefault("reference_tables", {})["project_names"] = "data/project_names.xlsx"
            fname = local_path.name
            self.project_table_label.setText(f"已导入: {fname}")
            self.project_table_label.setStyleSheet("color: #4CAF50;")
            self.log(f"项目名称表已导入: {fname}")
            self.save_session()

    def start_run(self):
        if not self.config:
            QMessageBox.critical(self, "错误", "配置文件未加载!")
            return

        # 模式选择
        if self.radio_stage2.isChecked():
            mode = "stage2"
        elif self.radio_stage1.isChecked():
            mode = "stage1"
        else:
            mode = "all"

        # 阶段二模式必须选输入文件
        if mode == "stage2" and not getattr(self, "stage2_input_path", None):
            QMessageBox.critical(self, "错误", "阶段二模式必须先选择阶段一输出的工单待查文件")
            return

        # 使用界面当前输入的凭据覆盖内存配置
        self.config.setdefault("email", {})
        self.config["email"]["address"] = self.input_email_addr.text().strip()
        self.config["email"]["password"] = self.input_email_pwd.text().strip()
        self.config.setdefault("workorder", {})
        self.config["workorder"]["username"] = self.input_wo_user.text().strip()
        self.config["workorder"]["password"] = self.input_wo_pwd.text().strip()
        self.config["workorder"]["force_live_query"] = self.chk_force_live.isChecked()

        # 仅 stage1/all 需要时间范围
        if mode != "stage2":
            date_from_q = self.date_from.date()
            date_to_q = self.date_to.date()
            if date_from_q.daysTo(date_to_q) < 0:
                QMessageBox.warning(self, "日期范围错误", "开始日期不能晚于结束日期")
                return
            date_from = datetime(
                date_from_q.year(), date_from_q.month(), date_from_q.day()
            )
            # MailReader 负责把 UI 的闭区间结束日转换成 IMAP 的 BEFORE 上界；
            # 这里不能再加一天，否则会多处理一天邮件。
            date_to = datetime(
                date_to_q.year(), date_to_q.month(), date_to_q.day()
            )
            self.log(
                f"邮件查询范围: {date_from:%Y-%m-%d} ~ {date_to:%Y-%m-%d}"
            )
        else:
            date_from = date_to = None

        agent_path = getattr(self, "agent_email_path", None)
        internal_path = getattr(self, "internal_email_path", None)
        project_path = getattr(self, "project_table_path", None)
        stage2_in = getattr(self, "stage2_input_path", None)

        self.worker = WorkerThread(
            self.config, mode=mode,
            date_from=date_from, date_to=date_to,
            agent_email_path=agent_path,
            internal_email_path=internal_path,
            project_table_path=project_path,
            stage2_input_path=stage2_in,
        )
        self.save_session()
        self.worker.log_signal.connect(self.log)
        self.worker.progress_signal.connect(self.on_progress)
        self.worker.finished_signal.connect(self.on_finished)
        self.worker.cancelled_signal.connect(self.on_cancelled)
        self.worker.error_signal.connect(self.on_error)
        self.worker.start()

        self.btn_run.setEnabled(False)
        self.btn_stop.setEnabled(True)
        self.btn_open_missing.setEnabled(False)
        self.btn_open_filtered.setEnabled(False)
        self.btn_open_stage2_result.setEnabled(False)
        self.btn_open_workbench.setEnabled(False)
        self.progress.setValue(0)

    def stop_run(self):
        if self.worker and self.worker.isRunning():
            self.worker.stop("桌面停止按钮")
            self.log("用户请求停止运行，正在取消当前操作...")
            self.btn_stop.setEnabled(False)

    def on_progress(self, current, total):
        if total <= 0:
            return
        self.progress.setMaximum(total)
        self.progress.setValue(current)
        self.progress.setFormat(f"{current} / {total}")

    def on_finished(self, primary_file, secondary_file):
        # 根据当前模式分别放置文件名 + 启用按钮
        if self.radio_stage2.isChecked():
            self.stage2_result_file = primary_file
            self.btn_open_stage2_result.setEnabled(True)
            self.log(f"工单核对结果: {primary_file}")
        elif self.radio_stage1.isChecked():
            self.missing_file = primary_file
            self.filtered_file = secondary_file
            self.stage1_cache_selection = {
                "primary": primary_file,
                "review": secondary_file,
            }
            self.btn_open_missing.setEnabled(True)
            self.btn_open_filtered.setEnabled(True)
            self.btn_open_workbench.setEnabled(True)
            self._refresh_stage1_cache_button()
            self.log(f"阶段一待查清单: {primary_file}")
            self.log(f"阶段一漏单复查: {secondary_file}")
        else:
            # 一站式: primary=漏单清单, secondary=过滤清单
            self.missing_file = primary_file
            self.filtered_file = secondary_file
            self.btn_open_missing.setEnabled(True)
            self.btn_open_filtered.setEnabled(True)
            # 一站式模式已经直接进入阶段二，不能把“漏单清单/过滤清单”误当成
            # 阶段一“工单待查/漏单复查”输入；请切换为仅阶段一后再人工确认。
            self.btn_open_workbench.setEnabled(False)
            self.log(f"漏单清单: {primary_file}")
            self.log(f"过滤清单: {secondary_file}")
            self.log("提示：人工复核工作台请使用‘仅阶段一（邮件解析）’运行，确认后再进入阶段二。")
        self.btn_run.setEnabled(True)
        self.btn_stop.setEnabled(False)
        self.progress.setFormat("完成")
        self.log("=" * 50)
        self.log("运行完成!")
        if self.radio_stage2.isChecked():
            self.log(f"  工单核对结果: {self.stage2_result_file}")
        else:
            self.log(f"  主输出: {self.missing_file}")
            self.log(f"  辅助输出: {self.filtered_file}")

    def on_error(self, error_msg):
        self.btn_run.setEnabled(True)
        self.btn_stop.setEnabled(False)
        self.log(f"[ERROR] {error_msg}")
        QMessageBox.critical(self, "运行错误", error_msg)

    def on_cancelled(self, message):
        """停止后恢复界面；不把未完成的 RPA 查询伪装成“运行完成”。"""
        self.btn_run.setEnabled(True)
        self.btn_stop.setEnabled(False)
        self.progress.setFormat("已停止")
        self.log("=" * 50)
        self.log(message)

    def log(self, msg):
        self.log_text.append(msg)
        self.log_text.verticalScrollBar().setValue(
            self.log_text.verticalScrollBar().maximum()
        )

    def open_excel(self, filepath):
        if filepath and os.path.exists(filepath):
            os.startfile(filepath)
        else:
            QMessageBox.warning(self, "提示", "文件不存在")

    def load_stage1_cache(self):
        """按日期选择阶段一缓存，并将其作为当前工作台数据源。"""
        options = _stage1_cache_options()
        if not options:
            QMessageBox.information(
                self,
                "没有可读取的缓存",
                "当前还没有按日期保存的阶段一结果，请先运行一次“阶段一：邮件解析”。",
            )
            return
        labels = [item["label"] for item in options]
        label, ok = QInputDialog.getItem(
            self,
            "读取阶段一缓存",
            "选择要读取的邮件日期范围：",
            labels,
            0,
            False,
        )
        if not ok or not label:
            return
        selected = next((item for item in options if item["label"] == label), None)
        if not selected:
            return
        self.missing_file = selected.get("primary", "")
        self.filtered_file = selected.get("review", "")
        self.stage1_cache_selection = selected
        self.btn_open_missing.setEnabled(bool(self.missing_file))
        self.btn_open_filtered.setEnabled(bool(self.filtered_file))
        self.btn_open_workbench.setEnabled(bool(self.missing_file or self.filtered_file))
        self._set_date_widget(self.date_from, selected.get("date_from"))
        self._set_date_widget(self.date_to, selected.get("date_to"))
        self.save_session()
        self.log(
            "已读取阶段一日期缓存: "
            f"{selected.get('date_from') or '未知'} ~ {selected.get('date_to') or '未知'}；"
            f"待查={self.missing_file or '无'}，复查={self.filtered_file or '无'}"
        )

    def open_workbench(self):
        """启动本机人工复核工作台；原 PyQt GUI 和阶段流程保持不变。"""
        try:
            import glob
            from workbench_server import WorkbenchServer
            primary = self.missing_file if self.missing_file and os.path.exists(self.missing_file) else None
            review = self.filtered_file if self.filtered_file and os.path.exists(self.filtered_file) else None

            # GUI 重启后没有内存中的文件路径时，直接从阶段一输出缓存恢复最近一批
            # 结果；这里读取的是 xlsx 产物，不会重新连接 IMAP，也不会读取历史人工状态。
            stage1_dir = os.path.join(APP_DIR, "output", "stage1_email")
            if not primary:
                stable = os.path.join(stage1_dir, "to_workorder_list.xlsx")
                candidates = [p for p in glob.glob(
                    os.path.join(stage1_dir, "to_workorder_list_*.xlsx")
                ) if os.path.isfile(p)]
                primary = stable if os.path.exists(stable) else (
                    max(candidates, key=os.path.getmtime) if candidates else None
                )
            if not review:
                stable = os.path.join(stage1_dir, "to_review_list.xlsx")
                candidates = [p for p in glob.glob(
                    os.path.join(stage1_dir, "to_review_list_*.xlsx")
                ) if os.path.isfile(p)]
                review = stable if os.path.exists(stable) else (
                    max(candidates, key=os.path.getmtime) if candidates else None
                )
            filtered = None
            if primary:
                candidate = os.path.join(os.path.dirname(primary), "filtered_mail_record.xlsx")
                if os.path.exists(candidate):
                    filtered = candidate

            # 工作台始终读取并延续正式人工复核历史。
            # 先停掉旧 HTTP 服务，避免浏览器仍连接到旧状态实例。
            if self.workbench_server:
                self.workbench_server.stop()
            self.workbench_server = WorkbenchServer(
                primary_path=primary,
                review_path=review,
                filtered_path=filtered,
                test_mode=False,
            )
            url = self.workbench_server.start(open_browser=True)
            self.log(f"人工复核工作台已启动: {url}")
            self.log(f"工作台数据源: 阶段一缓存文件（待查={primary or '无'}，复查={review or '无'}）")
            self.log("工作台正式模式：读取并延续历史人工复核状态")
        except Exception as exc:
            self.log(f"[ERROR] 人工复核工作台启动失败: {exc}")
            QMessageBox.critical(self, "工作台启动失败", str(exc))

    def closeEvent(self, event):
        """协作停止任务后保存会话；不能销毁仍在运行的 QThread。"""
        worker = getattr(self, "worker", None)
        if worker and worker.isRunning():
            event.ignore()
            if not getattr(self, "_close_requested", False):
                answer = QMessageBox.question(
                    self, "停止任务后退出",
                    "任务仍在运行。是否请求停止，并等待当前处理安全结束后退出？",
                    QMessageBox.Yes | QMessageBox.No, QMessageBox.No,
                )
                if answer != QMessageBox.Yes:
                    return
                self._close_requested = True
                worker.stop("关闭桌面窗口")
                self.log("正在等待任务安全停止，请勿强制结束程序。")
            QTimer.singleShot(250, self.close)
            return
        try:
            if self.save_session() is False:
                raise RuntimeError("会话保存失败，请检查文件权限或磁盘空间后再次退出。")
            if self.workbench_server:
                self.workbench_server.stop()
        except Exception as exc:
            self._close_requested = False
            event.ignore()
            QMessageBox.warning(self, "暂未退出", str(exc))
            return
        event.accept()


def _find_running_instances():
    """只读发现尚未采用实例锁的旧进程；永不结束其他进程。"""
    if os.name != "nt":
        return []
    my_pid = os.getpid()

    main_path = os.path.normcase(os.path.abspath(os.path.join(APP_DIR, "main.py")))
    escaped_main_path = main_path.replace("'", "''")
    ps_cmd = f"""
$ErrorActionPreference='Stop'
$appPid={my_pid}
$appPath='{escaped_main_path}'
$pattern='(?i)(?:^|[\\s\"])'+[regex]::Escape($appPath)+'(?=[\\s\"]|$)'
Get-CimInstance Win32_Process -Filter "Name='pythonw.exe' OR Name='python.exe'" |
 Where-Object {{ $_.ProcessId -ne $appPid -and $_.CommandLine -and
     ($_.CommandLine.Replace('/','\\') -match $pattern) }} |
 ForEach-Object {{ $_.ProcessId }}
"""
    try:
        out = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", ps_cmd],
            capture_output=True, text=True, timeout=10, creationflags=0x08000000,
        )
        if out.returncode:
            raise RuntimeError("实例查询未成功")
        return [int(line.strip()) for line in (out.stdout or "").splitlines() if line.strip()]
    except Exception as exc:
        raise RuntimeError("无法检查旧程序是否仍在运行；未自动关闭任何程序，请稍后重新启动。") from exc


def main():
    # ---- 防御层: 让 pythonw.exe 启动的 GUI 也能留下崩溃日志 ----
    import sys as _sys
    import traceback as _tb
    log_dir = os.path.join(APP_DIR, "logs")
    os.makedirs(log_dir, exist_ok=True)
    crash_log = os.path.join(log_dir, "crash.log")
    stderr_log = os.path.join(log_dir, "gui_stderr.log")

    # 1) pythonw.exe 默认丢弃 stderr, 这里重定向到文件, 任何 C 扩展 / Qt 内部错误才能看到
    try:
        _stderr_fh = open(stderr_log, "a", encoding="utf-8", buffering=1)
        _sys.stderr = _stderr_fh
    except Exception:
        pass

    # 2) 全局未捕获异常钩子: 写 crash.log + 同步写 stderr (PyQt5 槽函数异常 QApplication 会静默)
    def _excepthook(exc_type, exc_value, exc_tb):
        msg = "".join(_tb.format_exception(exc_type, exc_value, exc_tb))
        try:
            with open(crash_log, "a", encoding="utf-8") as f:
                f.write(f"\n[{_sys.executable}] uncaught:\n{msg}\n")
        except Exception:
            pass
        try:
            _sys.__stderr__.write(msg)
        except Exception:
            pass
    _sys.excepthook = _excepthook

    # 3) PyQt5 槽函数未捕获异常钩子 (PyQt5 默认会 qFatal 直接退出 QApplication)
    try:
        from PyQt5.QtCore import qInstallMessageHandler, QtMsgType
        def _qt_msg_handler(mode, ctx, msg):
            msg_str = str(msg)
            # 过滤无害警告: 字体目录缺失 (PyQt5 5.15+ Windows 已知问题, 不影响功能)
            if "QFontDatabase" in msg_str or "fonts" in msg_str.lower() and "Cannot find" in msg_str:
                return
            tag = {QtMsgType.QtDebugMsg: "DEBUG",
                   QtMsgType.QtInfoMsg: "INFO",
                   QtMsgType.QtWarningMsg: "WARN",
                   QtMsgType.QtCriticalMsg: "CRIT",
                   QtMsgType.QtFatalMsg: "FATAL"}.get(mode, "MSG")
            line = f"[Qt{tag}] {msg_str}\n"
            try:
                with open(crash_log, "a", encoding="utf-8") as f:
                    f.write(line)
            except Exception:
                pass
            if mode in (QtMsgType.QtFatalMsg, QtMsgType.QtCriticalMsg):
                try:
                    _sys.stderr.write(line)
                except Exception:
                    pass
        qInstallMessageHandler(_qt_msg_handler)
    except Exception:
        pass

    # 实例锁防止两个新版本同时启动；旧版进程另做只读检测。
    app = QApplication(_sys.argv)
    lock_dir = Path(APP_DIR) / "storage"
    lock_dir.mkdir(parents=True, exist_ok=True)
    instance_lock = QLockFile(str(lock_dir / "gui-instance.lock"))
    instance_lock.setStaleLockTime(0)
    if not instance_lock.tryLock(100):
        QMessageBox.information(None, "程序已运行", "请先保存并正常退出已有主程序，再启动新版。")
        return
    try:
        if _find_running_instances():
            QMessageBox.information(None, "旧程序仍在运行", "请先保存并正常退出旧主程序，再启动新版。不会强制结束旧任务。")
            return
        app.setStyle("Fusion")
        window = MainWindow()
        window.show()
        app.exec_()
    except RuntimeError as exc:
        QMessageBox.warning(None, "暂未启动", str(exc))
    finally:
        instance_lock.unlock()


if __name__ == "__main__":
    main()
