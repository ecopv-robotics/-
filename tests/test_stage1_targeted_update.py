from copy import copy

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from utils.stage1_targeted_update import mail_key, replace_selected_rows


HEADERS = ["发件人邮箱", "发件日期", "邮件主题", "客户公司名称", "标准化项目名称", "需求"]


def row(company, subject="target"):
    return dict(zip(HEADERS, ["test@example.invalid", "2026-08-01 12:00:00", subject,
                              company, "法国包装法", "注册"]))


def book():
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    for item in [row("Alpha LLC"), row("Fixture 01974750 LLC", "unrelated"), row("Wrong label")]:
        ws.append(list(item.values()))
    ws["D3"].fill = PatternFill("solid", fgColor="AABBCC")
    ws.freeze_panes = "B2"
    wb.create_sheet("audit")["A1"] = "keep audit"
    return wb


def test_keep_unrelated_row_positions_styles_and_matching_identity(tmp_path):
    wb = book()
    style = copy(wb.active["D3"]._style)
    key = mail_key(row("Alpha LLC"))
    report = replace_selected_rows(wb, {key: [row("Beta LLC"), row("Alpha LLC"), row("Gamma LLC")]})
    assert report[key]["after"] == 3
    assert wb.active["D2"].value == "Alpha LLC"
    assert wb.active["D3"].value == "Fixture 01974750 LLC"
    assert wb.active["D4"].value == "Beta LLC"
    assert wb.active["D5"].value == "Gamma LLC"
    assert wb.active["D3"]._style == style
    target = tmp_path / "check.xlsx"
    wb.save(target)
    reopened = load_workbook(target)
    assert reopened.active.freeze_panes == "B2"
    assert reopened["audit"]["A1"].value == "keep audit"
    assert reopened.active["D3"]._style == style
    reopened.close()


def test_clear_target_without_shifting_unrelated_rows():
    wb = book()
    replace_selected_rows(wb, {mail_key(row("Alpha LLC")): []})
    assert wb.active["D2"].value is None
    assert wb.active["D3"].value == "Fixture 01974750 LLC"
    assert wb.active["D4"].value is None


@pytest.mark.parametrize("bad", [
    [row("Alpha LLC"), row("Alpha LLC")],
    [row("Alpha LLC", "different mail")],
    [{**row("Alpha LLC"), "unknown": "bad"}],
    [row("x" * 32768)],
])
def test_reject_invalid_batch_before_any_changes(bad):
    wb = book()
    before = list(wb.active.values)
    with pytest.raises(ValueError):
        replace_selected_rows(wb, {mail_key(row("Alpha LLC")): bad})
    assert list(wb.active.values) == before


def test_mail_text_starting_equals_is_not_a_formula(tmp_path):
    wb = book()
    replace_selected_rows(wb, {mail_key(row("Alpha LLC")): [row("=formula is text")]})
    target = tmp_path / "literal.xlsx"
    wb.save(target)
    reopened = load_workbook(target)
    assert reopened.active["D2"].data_type == "s"
    reopened.close()
