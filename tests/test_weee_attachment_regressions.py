import io
import json
import shutil
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

from modules.field_extractor import FieldExtractor
from modules.weee_category_audit import extract_weee_items
from utils.attachment_parser import _xlsx_structured_records, parse_attachment


def _malformed_weee_xlsx(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "EPR申请表"
    sheet.cell(39, 2).value = "WEEE"
    sheet.cell(39, 3).value = "*WEEE类别（必填）"
    sheet.cell(39, 5).value = "*对应的品牌（多个品牌，请用逗号隔开）"
    sheet.cell(40, 3).value = "小型非光伏设备"
    sheet.cell(40, 4).value = "https://example.test/product"
    sheet.cell(40, 5).value = "DemoBrand"
    # 后续电池区块也含有“类别/品牌”字样，不能被当成 WEEE 项目。
    sheet.cell(43, 2).value = "电池"
    sheet.cell(43, 3).value = "*电池类型（必填）"
    sheet.cell(43, 5).value = "*对应的品牌（多个品牌，请用逗号隔开）"
    workbook.save(source)

    malformed = tmp_path / "malformed.xlsx"
    with zipfile.ZipFile(source) as source_zip, zipfile.ZipFile(malformed, "w") as output_zip:
        for info in source_zip.infolist():
            payload = source_zip.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                payload = payload.replace(b'<dimension ref="A1:W43"', b'<dimension ref="A1"')
            output_zip.writestr(info, payload)
    return malformed


def test_malformed_dimension_keeps_weee_brand_and_category(tmp_path):
    path = _malformed_weee_xlsx(tmp_path)
    attachment = parse_attachment(str(path), path.name)

    records = attachment["structured_records"]
    assert [(record["brand"], record["category"]) for record in records] == [
        ("DemoBrand", "小型非光伏设备")
    ]
    result = extract_weee_items(
        subject="Demo customer-德国WEEE",
        body="",
        attachments=[attachment],
        project="德国WEEE",
    )
    assert [(item["brand"], item["category"], item["category_class"]) for item in result["items"]] == [
        ("DemoBrand", "小型非光伏设备", "5")
    ]


def test_weee_form_reads_selected_dropdown_category_with_company_and_cell_evidence(tmp_path):
    path = tmp_path / "德国WEEE申请表.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "德国WEEE申请表"
    sheet["A2"] = "公司中文名称"
    sheet["B2"] = "深圳示例科技有限公司"
    sheet["A8"] = "WEEE产品信息"
    sheet["B9"] = "*类别"
    sheet["B10"] = "小型设备"
    sheet["C9"] = "公司类型"
    sheet["C10"] = "有限公司"
    category_validation = DataValidation(type="list", formula1='"大型设备,小型设备"')
    sheet.add_data_validation(category_validation)
    category_validation.add(sheet["B10"])
    company_type_validation = DataValidation(type="list", formula1='"有限公司,股份有限公司"')
    sheet.add_data_validation(company_type_validation)
    company_type_validation.add(sheet["C10"])
    workbook.save(path)

    attachment = parse_attachment(str(path), path.name)
    selected = [
        record for record in attachment["structured_records"]
        if record.get("weee_category_source") == "已选下拉值"
    ]

    assert len(selected) == 1
    assert selected[0]["category"] == "小型设备"
    assert selected[0]["customer"] == "深圳示例科技有限公司"
    assert selected[0]["row_number"] == 10
    assert "B10" in selected[0]["raw_text"]
    assert all(record.get("category") != "有限公司" for record in selected)

    result = extract_weee_items(
        subject="深圳示例科技有限公司-德国WEEE",
        body="",
        attachments=[attachment],
        project="德国WEEE",
        company="深圳示例科技有限公司",
    )
    assert [(item["brand"], item["category"], item["category_class"]) for item in result["items"]] == [
        ("", "小型设备", "5")
    ]
    assert "B10" in result["items"][0]["evidence"]


def test_weee_form_does_not_read_dropdowns_from_battery_section(tmp_path):
    path = tmp_path / "mixed-services.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "WEEE产品信息"
    sheet["A2"] = "WEEE类别"
    sheet["A3"] = "小型设备"
    sheet["A8"] = "电池产品信息"
    sheet["A9"] = "类别"
    sheet["A10"] = "锂电池"
    validation = DataValidation(type="list", formula1='"锂电池,铅酸电池"')
    sheet.add_data_validation(validation)
    validation.add(sheet["A10"])
    workbook.save(path)

    attachment = parse_attachment(str(path), path.name)

    assert not any(
        record.get("weee_category_source") == "已选下拉值"
        and record.get("category") == "锂电池"
        for record in attachment["structured_records"]
    )


def test_weee_workbook_prefers_new_brand_column_over_status_columns():
    rows = [
        [
            "编号", "公司中文名", "公司英文名", "原品牌品牌", "德文类别", "类别",
            "WEEE号", "联系邮箱", "日期", "RV号码", "备注", "", "法人", "需新增品牌", "品牌",
        ],
        [
            "1", "示例fe00185c有限公司", "entity-a", "old-brand", "Kleingeräte", "第五类",
            "17256311", "a@example.test", "2026-08-14", "RV-1", "重命名成功", "", "法人A", "DemoBrandA", "授权",
        ],
    ]

    records = _xlsx_structured_records(rows, "Sheet1")

    assert len(records) == 1
    assert records[0]["customer"] == "示例fe00185c有限公司"
    assert records[0]["brand"] == "DemoBrandA"
    assert records[0]["category"] == "Kleingeräte 第五类"


def test_weee_extraction_does_not_duplicate_structured_and_preview_rows():
    rows = [
        [
            "编号", "公司中文名", "公司英文名", "原品牌品牌", "德文类别", "类别",
            "WEEE号", "联系邮箱", "日期", "RV号码", "备注", "", "法人", "需新增品牌", "品牌",
        ],
        [
            "1", "示例fe00185c有限公司", "entity-a", "old-brand", "Kleingeräte", "第五类",
            "17256311", "a@example.test", "2026-08-14", "RV-1", "重命名成功", "", "法人A", "DemoBrandA", "授权",
        ],
    ]
    records = _xlsx_structured_records(rows, "Sheet1")
    attachment = {
        "filename": "德国weee新增品牌.xlsx",
        "structured_records": records,
        "sheets": [{
            "sheet_name": "Sheet1",
            "preview_rows": [
                {"row_number": 1, "cells": rows[0]},
                {"row_number": 2, "cells": rows[1]},
            ],
        }],
    }

    result = extract_weee_items(
        subject="示例fe00185c有限公司-德国WEEE",
        body="",
        attachments=[attachment],
        project="德国WEEE",
        company="示例fe00185c有限公司",
    )

    assert [(item["brand"], item["category_class"]) for item in result["items"]] == [
        ("DemoBrandA", "5")
    ]
    assert result["items"][0]["category"].startswith("Kleingeräte")
    assert "第五类" in result["items"][0]["category"]


def test_weee_five_subject_workbook_emits_five_rows_with_one_item_each():
    header = [
        "编号", "公司中文名", "公司英文名", "原品牌品牌", "德文类别", "类别",
        "WEEE号", "联系邮箱", "日期", "RV号码", "备注", "", "法人", "需新增品牌", "品牌",
    ]
    brands = ["DemoBrandA", "DemoBrandB", "DemoBrandC", "DemoBrandC", "DemoBrandD"]
    rows = [header]
    for index, brand in enumerate(brands, start=1):
        rows.append([
            str(index), f"主体{index}有限公司", f"entity-{index}", "old-brand",
            "Kleingeräte", "第五类", f"WEEE-{index}", f"{index}@example.test",
            "2026-08-14", f"RV-{index}", "重命名成功", "", f"法人{index}", brand, "授权",
        ])
    records = _xlsx_structured_records(rows, "Sheet1")
    attachment = {
        "filename": "德国weee新增品牌.xlsx",
        "structured_records": records,
        "sheets": [{"sheet_name": "Sheet1", "preview_rows": [
            {"row_number": index + 1, "cells": row} for index, row in enumerate(rows)
        ]}],
    }
    extractor = FieldExtractor(
        {}, [{"项目名称": "德国WEEE", "国家": "德国", "业务类型": "WEEE"}],
        ocr_fallback=False,
    )
    output = extractor.extract_fields({
        "subject": "示例代理戊科技——5个主体德国WEEE新增品牌",
        "body_text": "",
        "body_original": "",
        "sender_email": "agent@example.test",
        "date": "2026-08-26 17:25:24",
        "attachments": [attachment],
    })

    assert len(output) == 5
    parsed = [json.loads(row["德国WEEE品类明细"])[0] for row in output]
    assert [item["brand"] for item in parsed] == brands
    assert all(
        item["category"].startswith("Kleingeräte") and "第五类" in item["category"]
        for item in parsed
    )
    assert [len(json.loads(row["德国WEEE品类明细"])) for row in output] == [1] * 5


def test_workbench_recovers_items_from_stale_attachment_index(tmp_path, monkeypatch):
    import workbench_server

    cached = tmp_path / "cache" / "attachments"
    cached.mkdir(parents=True)
    source = _malformed_weee_xlsx(tmp_path)
    shutil.copyfile(source, cached / "demo_form.xlsx")
    monkeypatch.setattr(workbench_server, "APP_ROOT", tmp_path)
    row = {
        "标准化项目名称": "德国WEEE",
        "德国WEEE专项": "是",
        "客户公司名称": "Demo customer",
        "邮件主题": "Demo customer-德国WEEE",
        "附件文件索引": json.dumps([{"filename": "demo_form.xlsx", "token": "demo_form.xlsx"}]),
    }

    recovered = workbench_server._recover_weee_items_from_cached_attachments(row, [])
    assert [(item["brand"], item["category"]) for item in recovered] == [
        ("DemoBrand", "小型非光伏设备")
    ]


def test_workbench_exposes_manual_weee_item_entry():
    html = Path("workbench.html").read_text(encoding="utf-8")
    assert "data-add-weee-item" in html
    assert "新增品牌/品类" in html
    assert "save_weee_draft" in html
    assert "function scheduleWeeeDraftSave(detail,delay=360)" in html
    assert "updateWeeeDraftFromField(field" in html
    assert "function attachmentFilesFor(m,d=null)" in html
    assert "attachmentChip(mail,name,'',false,detail)" in html
    assert "attachment-download" in html


def test_workbench_persists_weee_draft_without_confirming_it(tmp_path):
    import json
    import workbench_server

    state_path = tmp_path / "workbench_review.json"
    store = workbench_server.WorkbenchStore(
        primary_path=str(tmp_path / "primary.xlsx"),
        review_path=str(tmp_path / "review.xlsx"),
        filtered_path=str(tmp_path / "filtered.xlsx"),
        state_path=str(state_path),
    )
    result = store.action({
        "action": "save_weee_draft",
        "record_id": "DETAIL-DRAFT-1",
        "mail_id": "MAIL-DRAFT-1",
        "weee_items": [{
            "item_id": "WEEE-DRAFT-1",
            "brand": "YWNYT",
            "category": "电气与电子设备废料-大型设备(光伏面板除外)",
            "category_class": "4",
            "category_class_name": "大型设备",
            "weee_confirmed": False,
        }],
        "reason": "输入草稿自动保存",
    })
    assert result["ok"] is True
    saved = json.loads(state_path.read_text(encoding="utf-8"))
    entry = saved["records"]["DETAIL-DRAFT-1"]
    assert entry["weee_items"][0]["category"] == "电气与电子设备废料-大型设备(光伏面板除外)"
    assert entry["weee_status"] == "pending"
    assert entry["weee_items"][0]["weee_confirmed"] is False


def test_workbench_reads_raw_imap_total(tmp_path, monkeypatch):
    import workbench_server

    stage1 = tmp_path / "stage1_email"
    stage1.mkdir()
    (stage1 / "imap_summary.json").write_text(
        json.dumps({"imap_read_total": 296, "parsed_total": 268, "date_from": "2026-08-23", "date_to": "2026-09-27"}),
        encoding="utf-8",
    )
    monkeypatch.setattr(workbench_server, "OUTPUT_ROOT", tmp_path)
    assert workbench_server._read_imap_summary()["imap_read_total"] == 296
    assert workbench_server._read_imap_summary()["parsed_total"] == 268
