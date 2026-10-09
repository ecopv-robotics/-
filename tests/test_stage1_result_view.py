import json

from openpyxl import load_workbook

from modules.excel_writer import ExcelWriter, WORKORDER_RESULT_HEADERS


def _weee_row(company="Fixture WEEE GmbH", brand="FixtureBrand", category="第五类"):
    return {
        "sender_email": "sender@example.invalid",
        "recipient": "report@example.invalid",
        "date": "2026-08-26 10:00:00",
        "subject": "Fixture registration",
        "代理": "华之星",
        "代理匹配方式": "邮箱精确匹配",
        "国家": "德国",
        "项目": "德国WEEE",
        "需求": "注册",
        "客户": company,
        "置信度": "high",
        "德国WEEE专项": "是",
        "德国WEEE品类明细": json.dumps(
            [{"brand": brand, "category": "Small equipment", "category_class_name": category}],
            ensure_ascii=False,
        ),
        "德国WEEE品类核对": json.dumps({"status": "matched"}, ensure_ascii=False),
    }


def test_stage1_outputs_share_stage2_schema_and_preserve_hidden_inputs(tmp_path):
    output = tmp_path / "output"
    writer = ExcelWriter(str(output), categorized=True)
    battery = {
        "代理": "示例代理",
        "代理匹配方式": "邮箱精确匹配",
        "国家": "德国",
        "项目": "德国电池法",
        "需求": "新增",
        "客户": "Fixture Battery GmbH",
        "置信度": "high",
        "德国电池法品类明细": json.dumps(
            [{"brand": "BatteryBrand", "category": "设备电池"}], ensure_ascii=False
        ),
    }
    needs_review = {
        "项目": "德国WEEE",
        "国家": "德国",
        "需求": "注册",
        "客户": "Fixture Review GmbH",
        "置信度": "low",
    }
    filtered = [{
        "filter_object": "德国一次性塑料税",
        "filter_country": "德国",
        "filter_agent": "示例代理",
        "filter_company": "Fixture Filter GmbH",
        "filter_request": "注册",
        "filter_reason": "明确非目标项目",
    }]

    primary, review = writer.write_stage1_outputs(
        [_weee_row(), battery, needs_review], filtered
    )
    filtered_path = output / "stage1_email" / "filtered_mail_record.xlsx"

    for path, raw_sheet in (
        (primary, "工单待查"),
        (review, "漏单复查"),
        (filtered_path, "过滤日志"),
    ):
        workbook = load_workbook(path, read_only=False, data_only=True)
        try:
            assert workbook.active.title == "阶段一结果"
            assert [cell.value for cell in workbook.active[1]] == WORKORDER_RESULT_HEADERS
            assert workbook[raw_sheet].sheet_state == "hidden"
        finally:
            workbook.close()

    workbook = load_workbook(primary, read_only=True, data_only=True)
    try:
        rows = list(workbook["阶段一结果"].iter_rows(min_row=2, values_only=True))
        assert rows[0][:7] == (
            "德国WEEE", "华之星", "德国", "德国WEEE - 注册",
            "Fixture WEEE GmbH", "FixtureBrand", "第五类",
        )
        assert rows[0][7] == "否"
        assert "尚未执行阶段二" in rows[0][8]
        assert rows[1][:7] == (
            "德国电池法", "示例代理", "德国", "德国电池法 - 新增",
            "Fixture Battery GmbH", "BatteryBrand", "设备电池",
        )
        raw = workbook["工单待查"]
        headers = {cell.value for cell in next(raw.iter_rows(max_row=1))}
        assert {"发件日期", "代理", "客户公司名称", "标准化项目名称"}.issubset(headers)
    finally:
        workbook.close()

    workbook = load_workbook(review, read_only=True, data_only=True)
    try:
        row = next(workbook["阶段一结果"].iter_rows(min_row=2, values_only=True))
        assert row[7] == "否"
        assert "待人工复核" in row[8]
    finally:
        workbook.close()

    workbook = load_workbook(filtered_path, read_only=True, data_only=True)
    try:
        row = next(workbook["阶段一结果"].iter_rows(min_row=2, values_only=True))
        assert row[0] == "其他项目"
        assert row[5:7] == ("无", "无")
        assert "阶段一已过滤" in row[8]
    finally:
        workbook.close()


def test_stage2_actual_result_still_uses_same_schema_and_success_gate(tmp_path):
    writer = ExcelWriter(str(tmp_path / "output"), categorized=True)
    success = _weee_row()
    success.update({
        "是否已录单": "是",
        "匹配状态": "唯一匹配",
        "RPA查询状态": "查询完成",
    })
    workbook = writer._do_write_workorder_check([success], [success], [])
    try:
        sheet = workbook.active
        assert [cell.value for cell in sheet[1]] == WORKORDER_RESULT_HEADERS
        row = next(sheet.iter_rows(min_row=2, values_only=True))
        assert row[7] == "是"
        assert row[8] == ""
    finally:
        workbook.close()

    missing_values = writer._workorder_export_values(
        {"国家": "德国", "项目": "德国WEEE", "需求": "注册", "客户": "Fixture Missing GmbH"},
        "stage1_pending",
    )[0]
    assert missing_values[5:7] == ["无", "无"]
