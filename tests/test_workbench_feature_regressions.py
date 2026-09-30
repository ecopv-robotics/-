import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from modules.field_extractor import FieldExtractor
from workbench_server import (
    WorkbenchStore,
    _detail_confirmation_ready,
    _mail_in_date_range,
    _repair_attachment_fields,
)
from gui import _filter_workbench_confirmed_rows


def test_epr_application_company_has_priority_over_body_and_subject():
    extractor = FieldExtractor({}, [{"项目名称": "德国包装法"}])
    rows = extractor.extract_fields({
        "subject": "代理-示例2753d332有限公司-德国包装法",
        "body_text": "公司名称：示例39ecdfaf有限公司\n申请注册",
        "sender_email": "agent@example.test",
        "attachments": [{
            "filename": "泛欧EPR申请表.xlsx",
            "text_content": "公司名称：示例2aa087c1有限公司\nApplication for EPR",
        }],
    })

    assert rows
    assert rows[0]["客户"] == "示例2aa087c1有限公司"
    assert rows[0]["客户提取来源"] == "附件EPR申请表"


def test_stage2_workbench_input_only_keeps_confirmed_details():
    headers = ["客户公司名称", "标准化项目名称", "工作台状态"]
    rows = [
        {"客户公司名称": "A有限公司", "工作台状态": "confirmed"},
        {"客户公司名称": "B有限公司", "工作台状态": "review"},
    ]

    selected, is_workbench, has_status = _filter_workbench_confirmed_rows(
        headers, rows, "workbench_reviewed_20260923.xlsx"
    )
    assert is_workbench is True
    assert has_status is True
    assert [row["客户公司名称"] for row in selected] == ["A有限公司"]

    legacy, is_legacy, has_legacy_status = _filter_workbench_confirmed_rows(
        ["客户公司名称"], rows, "to_workorder_list.xlsx"
    )
    assert is_legacy is False
    assert has_legacy_status is False
    assert legacy == rows


def test_export_date_range_uses_mail_date_only():
    assert _mail_in_date_range({"date": "2026-08-25 09:00:00"}, "2026-08-25", "2026-08-25")
    assert not _mail_in_date_range({"date": "2026-08-26 09:00:00"}, "2026-08-25", "2026-08-25")
    assert not _mail_in_date_range({"date": ""}, "2026-08-25", "2026-08-25")
    # 空范围保持历史“导出全部”行为。
    assert _mail_in_date_range({"date": ""}, "", "")


def test_confirming_one_detail_does_not_finish_other_details(tmp_path):
    primary = tmp_path / "primary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单待查"
    sheet.append([
        "邮件编号", "明细编号", "发件人邮箱", "发件日期", "邮件主题", "代理",
        "客户公司名称", "标准化项目名称", "需求", "置信度",
    ])
    for index, company in enumerate(("A公司", "B公司"), start=1):
        sheet.append([
            "MAIL-1", f"DETAIL-{index}", "agent@example.test", "2026-08-25 09:00:00",
            "双明细测试", "代理A", company, "德国WEEE", "注册", "high",
        ])
    workbook.save(primary)

    store = WorkbenchStore(
        primary_path=str(primary),
        review_path=str(tmp_path / "review.xlsx"),
        filtered_path=str(tmp_path / "filtered.xlsx"),
        state_path=str(tmp_path / "state.json"),
        test_mode=True,
    )
    before = store.snapshot()
    details = before["mails"][0]["details"]
    first_id = details[0]["id"]
    store.action({
        "action": "confirm_detail",
        "record_id": first_id,
        "mail_number": "MAIL-1",
        "detail_number": "DETAIL-1",
        "reason": "确认当前明细",
    })
    after = store.snapshot()

    assert [(detail["fields"]["company"], detail["status"]) for detail in after["mails"][0]["details"]] == [
        ("A公司", "confirmed"),
        ("B公司", "ready"),
    ]
    assert after["mails"][0]["status"] == "partial"


def test_partial_confirmation_exposes_confirmed_detail_and_keeps_pending_detail(tmp_path):
    primary = tmp_path / "primary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单待查"
    sheet.append([
        "邮件编号", "明细编号", "发件人邮箱", "发件日期", "邮件主题", "代理",
        "客户公司名称", "标准化项目名称", "需求", "置信度",
    ])
    for index, company in enumerate(("A公司", "B公司"), start=1):
        sheet.append([
            "MAIL-PARTIAL", f"DETAIL-{index}", "agent@example.test", "2026-08-25 09:00:00",
            "部分确认测试", "代理A", company, "德国包装法", "注册", "high",
        ])
    workbook.save(primary)

    store = WorkbenchStore(
        primary_path=str(primary), review_path=str(tmp_path / "review.xlsx"),
        filtered_path=str(tmp_path / "filtered.xlsx"), state_path=str(tmp_path / "state.json"),
        test_mode=True,
    )
    first, second = store.snapshot()["mails"][0]["details"]
    store.action({"action": "confirm_detail", "record_id": first["id"], "mail_number": "MAIL-PARTIAL", "detail_number": "DETAIL-1", "reason": "确认当前明细"})
    mail = store.snapshot()["mails"][0]
    assert mail["status"] == "partial"
    assert mail["confirmed_detail_count"] == 1
    assert mail["pending_detail_count"] == 1
    assert _detail_confirmation_ready(mail["details"][0]) is True
    assert _detail_confirmation_ready(mail["details"][1]) is False


def test_weee_category_confirmation_is_stage2_eligible_without_second_business_confirm(tmp_path):
    primary = tmp_path / "primary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单待查"
    sheet.append([
        "邮件编号", "明细编号", "发件人邮箱", "发件日期", "邮件主题", "代理",
        "客户公司名称", "标准化项目名称", "需求", "置信度", "德国WEEE专项", "德国WEEE品类明细",
    ])
    sheet.append([
        "MAIL-WEEE-CONFIRM", "DETAIL-WEEE-1", "agent@example.test", "2026-08-25 09:00:00",
        "WEEE确认测试", "代理A", "WEEE示例ddc0cb87有限公司", "德国WEEE", "注册", "high", "是",
        json.dumps([{"brand": "品牌A", "category": "小型设备", "category_class": "5"}], ensure_ascii=False),
    ])
    workbook.save(primary)
    store = WorkbenchStore(
        primary_path=str(primary), review_path=str(tmp_path / "review.xlsx"),
        filtered_path=str(tmp_path / "filtered.xlsx"), state_path=str(tmp_path / "state.json"),
        test_mode=True,
    )
    detail = store.snapshot()["mails"][0]["details"][0]
    items = [{"brand": "品牌A", "category": "小型设备", "category_class": "5"}]
    store.action({
        "action": "confirm_weee", "record_id": detail["id"], "mail_id": store.snapshot()["mails"][0]["id"],
        "mail_number": "MAIL-WEEE-CONFIRM", "detail_number": "DETAIL-WEEE-1", "weee_items": items,
        "reason": "确认德国 WEEE 品牌与品类",
    })
    after = store.snapshot()
    confirmed = after["mails"][0]["details"][0]
    assert confirmed["weee"]["confirmed"] is True
    assert confirmed["confirmation_ready"] is True
    assert after["mails"][0]["status"] == "confirmed"


def test_export_contains_only_current_confirmation_result(tmp_path):
    primary = tmp_path / "primary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单待查"
    sheet.append([
        "邮件编号", "明细编号", "发件人邮箱", "发件日期", "邮件主题", "代理",
        "客户公司名称", "标准化项目名称", "需求", "置信度",
    ])
    for index, company in enumerate(("已确认公司", "未确认公司"), start=1):
        sheet.append([
            "MAIL-EXPORT", f"DETAIL-EXPORT-{index}", "agent@example.test", "2026-08-25 09:00:00",
            "导出确认测试", "代理A", company, "德国包装法", "注册", "high",
        ])
    workbook.save(primary)
    store = WorkbenchStore(
        primary_path=str(primary), review_path=str(tmp_path / "review.xlsx"),
        filtered_path=str(tmp_path / "filtered.xlsx"), state_path=str(tmp_path / "state.json"),
        test_mode=True,
    )
    details = store.snapshot()["mails"][0]["details"]
    store.action({"action": "confirm_detail", "record_id": details[0]["id"], "mail_number": "MAIL-EXPORT", "detail_number": "DETAIL-EXPORT-1", "reason": "确认"})
    out = store.export(date_start="2026-08-25", date_end="2026-08-25")
    exported = load_workbook(out, read_only=True, data_only=True)
    ws = exported["工单待查"]
    rows = list(ws.iter_rows(values_only=True))
    headers = list(rows[0])
    body = [dict(zip(headers, row)) for row in rows[1:] if any(value not in (None, "") for value in row)]
    assert len(body) == 1
    assert body[0]["客户公司名称"] == "已确认公司"
    assert body[0]["工作台状态"] == "confirmed"
    assert body[0]["工作台确认类型"] == "普通业务"
    exported.close()


def test_delete_mail_removes_all_details_and_keeps_deleted_audit(tmp_path):
    primary = tmp_path / "primary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单待查"
    sheet.append(["邮件编号", "明细编号", "发件人邮箱", "发件日期", "邮件主题", "代理", "客户公司名称", "标准化项目名称", "需求", "置信度"])
    for index in (1, 2):
        sheet.append(["MAIL-DELETE", f"DETAIL-DELETE-{index}", "agent@example.test", "2026-08-25 09:00:00", "整封删除测试", "代理A", f"公司{index}", "德国包装法", "注册", "high"])
    workbook.save(primary)
    store = WorkbenchStore(
        primary_path=str(primary), review_path=str(tmp_path / "review.xlsx"),
        filtered_path=str(tmp_path / "filtered.xlsx"), state_path=str(tmp_path / "state.json"),
        test_mode=True,
    )
    mail = store.snapshot()["mails"][0]
    result = store.action({"action": "delete_mail", "mail_id": mail["id"], "mail_number": "MAIL-DELETE", "reason": "整封删除"})
    assert result["deleted_details"] == 2
    assert store.snapshot()["mails"] == []
    out = store.export()
    exported = load_workbook(out, read_only=True, data_only=True)
    deleted = exported["已删除明细"]
    assert deleted.max_row == 3
    exported.close()


def test_legacy_unconfirmed_weee_snapshot_yields_to_latest_stage1_row(tmp_path):
    primary = tmp_path / "primary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单待查"
    sheet.append([
        "邮件编号", "明细编号", "发件人邮箱", "发件日期", "邮件主题",
        "客户公司名称", "标准化项目名称", "德国WEEE专项", "德国WEEE品类明细",
        "代理", "需求", "置信度",
    ])
    current_company = "保定白沟新城轩琥商贸商行（个人独资）"
    current_item = {
        "brand": "DemoBrandA",
        "category": "Kleingeräte 第五类",
        "evidence": f"1 | {current_company} | DemoBrandA",
        "extraction_method": "规则提取·LLM复检通过",
    }
    sheet.append([
        "MAIL-WEEE-LEGACY", "DETAIL-WEEE-1", "agent@example.test", "2026-08-26 17:25:24",
        "5个主体德国WEEE新增品牌", current_company, "德国WEEE", "是",
        json.dumps([current_item], ensure_ascii=False), "示例代理戊", "新增", "high",
    ])
    workbook.save(primary)

    store = WorkbenchStore(
        primary_path=str(primary),
        review_path=str(tmp_path / "review.xlsx"),
        filtered_path=str(tmp_path / "filtered.xlsx"),
        state_path=str(tmp_path / "state.json"),
        test_mode=True,
    )
    rid = store._raw_records()[0]["_id"]
    store.state_path.write_text(json.dumps({
        "records": {
            rid: {
                "weee_status": "pending",
                "weee_items": [
                    {
                        "brand": "DemoBrandA",
                        "category": "Kleingeräte 第五类",
                        "evidence": f"1 | {current_company} | DemoBrandA",
                        "extraction_method": "LLM重新提取",
                        "llm_review_status": "invalid",
                    },
                    {
                        "brand": "DemoBrandB",
                        "category": "第五类",
                        "evidence": "2 | 示例2cbabe64有限公司 | DemoBrandB",
                        "extraction_method": "LLM重新提取",
                        "llm_review_status": "invalid",
                    },
                ],
                "events": [],
            }
        }
    }, ensure_ascii=False), encoding="utf-8")

    detail = store.snapshot()["mails"][0]["details"][0]
    assert len(detail["weee"]["items"]) == 1
    assert detail["weee"]["items"][0]["brand"] == "DemoBrandA"


def test_company_name_recovered_from_labeled_body_and_attachment_filename():
    body = (
        "你好，麻烦安排以下公司德国一次性塑料，如有问题请随时联系我们，谢谢\n\n"
        "公司：\n示例a2f907fc有限公司\n\n项目：德国一次性塑料\n生效时间：2026年"
    )
    row = _repair_attachment_fields({
        "客户公司名称": "",
        "客户": "",
        "邮件主题": "示例代理丙-示例a2f907fc有限公司-2026年德国一次性塑料注册",
        "邮件正文原文": body,
        "附件名称": "德国一次性塑料法-示例a2f907fc有限公司.zip；营业执照.jpg",
        "人工复核提示": "核对客户；程序客户公司字段缺失，需补全；company: 客户公司名称为空",
        "附件证据": "[]",
    })

    assert row["客户公司名称"] == "示例a2f907fc有限公司"
    assert row["客户"] == row["客户公司名称"]
    assert row["客户提取来源"] == "正文确定性回填"
    assert row["人工复核提示"] == "核对客户"

    # 正文缺失时才使用附件文件名；已有合法主体不能被证据回填覆盖。
    filename_only = _repair_attachment_fields({
        "客户公司名称": "",
        "客户": "",
        "附件名称": "德国一次性塑料法-示例a2f907fc有限公司.zip",
        "附件证据": "[]",
    })
    assert filename_only["客户公司名称"] == "示例a2f907fc有限公司"
    assert filename_only["客户提取来源"] == "附件文件名确定性回填"

    existing = _repair_attachment_fields({
        "客户公司名称": "示例2a9acf99有限公司",
        "客户": "示例2a9acf99有限公司",
        "邮件正文原文": body,
        "附件名称": "德国一次性塑料法-示例a2f907fc有限公司.zip",
        "附件证据": "[]",
    })
    assert existing["客户公司名称"] == "示例2a9acf99有限公司"


def test_field_extractor_prefers_company_label_over_introduction_sentence():
    extractor = FieldExtractor({}, [{"项目名称": "德国一次性塑料"}], ocr_fallback=False)
    rows = extractor.extract_fields({
        "subject": "示例代理丙-示例a2f907fc有限公司-2026年德国一次性塑料注册",
        "body_text": (
            "你好，麻烦安排以下公司德国一次性塑料，如有问题请随时联系我们，谢谢\n\n"
            "公司：\n示例a2f907fc有限公司\n\n项目：德国一次性塑料\n生效时间：2026年"
        ),
        "sender_email": "fixture-6610afc6@external.example.invalid",
        "attachments": [{"filename": "德国一次性塑料法-示例a2f907fc有限公司.zip"}],
    })

    assert rows
    assert rows[0]["客户"] == "示例a2f907fc有限公司"
    assert rows[0]["客户提取来源"] == "正文"


def test_legacy_blank_state_does_not_hide_repaired_company(tmp_path):
    primary = tmp_path / "primary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "工单待查"
    sheet.append([
        "邮件编号", "明细编号", "发件人邮箱", "发件日期", "邮件主题", "邮件正文原文",
        "附件名称", "代理", "客户公司名称", "标准化项目名称", "需求", "置信度",
    ])
    sheet.append([
        "MAIL-CAC2", "DETAIL-1", "fixture-6610afc6@external.example.invalid", "2026-08-26 15:56:22",
        "示例代理丙-示例a2f907fc有限公司-2026年德国一次性塑料注册",
        "公司：\n示例a2f907fc有限公司",
        "德国一次性塑料法-示例a2f907fc有限公司.zip", "示例代理丁", "",
        "德国一次性塑料", "注册", "需人工确认",
    ])
    workbook.save(primary)

    store = WorkbenchStore(
        primary_path=str(primary),
        review_path=str(tmp_path / "review.xlsx"),
        filtered_path=str(tmp_path / "filtered.xlsx"),
        state_path=str(tmp_path / "state.json"),
        test_mode=True,
    )
    raw = store._raw_records()
    store.state_path.write_text(
        json.dumps({"records": {raw[0]["_id"]: {"fields": {"company": ""}}}}, ensure_ascii=False),
        encoding="utf-8",
    )

    detail = store.snapshot()["mails"][0]["details"][0]
    assert detail["fields"]["company"] == "示例a2f907fc有限公司"


def test_confirmations_apply_locally_and_preserve_queue_and_detail_position():
    html = Path("workbench.html").read_text(encoding="utf-8")
    assert "function captureNavigation()" in html
    assert "function restoreNavigationSelection(snapshot)" in html
    assert "function restoreNavigationScroll(snapshot)" in html
    do_action = html.split("async function doAction(", 1)[1].split("async function addProject(", 1)[0]
    assert "await load()" not in do_action
    assert "restoreNavigationSelection(navigation);render();restoreNavigationScroll(navigation)" in do_action
    weee_confirmation = html.split("async function confirmWeeeItem(", 1)[1].split("async function deleteWeeeItem(", 1)[0]
    assert "await load()" not in weee_confirmation
    assert "restoreNavigationScroll(navigation)" in weee_confirmation
    assert "保存中…" in html
