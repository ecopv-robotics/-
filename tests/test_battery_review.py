import json

import pytest
from openpyxl import Workbook, load_workbook

from modules.battery_review import (battery_review, classify_battery_category,
                                    compare_battery_categories, extract_battery_items,
                                    normalize_battery_items)
from workbench_server import WorkbenchStore, _detail_confirmation_ready


def item(item_id="B1", brand="Example"):
    return {"item_id": item_id, "brand": brand, "category": "Original battery type",
            "category_class": "Manual battery class"}


@pytest.mark.parametrize("filename,expected", [
    ("GRS_Anlage_FZG 电动汽车电池.pdf", "汽车电池"),
    ("LV 轻型交通工具电池合同.pdf", "轻型交通工具电池"),
    ("STB 汽车启动电池合同.pdf", "启动电池"),
    ("NB_BESS 工业电池合同.pdf", "工业电池"),
    ("设备电池申请表.xlsx", "设备电池"),
])
def test_battery_five_classes_from_specific_evidence(filename, expected):
    assert classify_battery_category(filename) == expected


def test_generic_ofh_contract_only_gives_unconfirmed_equipment_candidate():
    result = extract_battery_items(project="德国电池法", attachments=[{
        "filename": "Vertrag_OfH合同.pdf",
        "text_content": "Gerätebatterien Starterbatterien LV-Batterien Elektrofahrzeugbatterien Industriebatterien",
    }])
    assert len(result["items"]) == 1
    assert result["items"][0]["category_class"] == "设备电池"
    assert result["items"][0]["confidence"] == "low"
    assert not result["items"][0]["battery_confirmed"]


def test_archive_member_contracts_are_kept_as_separate_battery_classes():
    result = extract_battery_items(project="德国电池法", attachments=[{
        "filename": "资料包.zip",
        "text_content": "[contracts/STB 汽车启动电池合同.pdf] Anlage STB\n"
                        "[contracts/NB_BESS 工业电池合同.pdf] Anlage INB BESS",
    }])
    assert [item["category_class"] for item in result["items"]] == ["启动电池", "工业电池"]


def test_mixed_weee_and_battery_stay_separate_by_project_and_company():
    attachments = [{"filename": "mixed.xlsx", "structured_records": [
        {"customer": "甲甲科技有限公司", "sheet_name": "德国WEEE", "brand": "WEEEBrand", "category": "小型设备"},
        {"customer": "甲甲科技有限公司", "sheet_name": "德国电池法", "brand": "BatteryBrand", "category": "设备电池"},
        {"customer": "乙乙科技有限公司", "sheet_name": "德国电池法", "brand": "OtherBrand", "category": "工业电池"},
    ]}]
    result = extract_battery_items(project="德国电池法", company="甲甲科技有限公司", attachments=attachments)
    assert [(item["brand"], item["category_class"]) for item in result["items"]] == [("BatteryBrand", "设备电池")]


def test_battery_form_section_reads_selected_value_but_not_weee_section():
    attachment = {"filename": "甲甲科技有限公司-EPR申请表.xlsx", "sheets": [{"sheet_name": "申请表", "preview_rows": [
        {"row_number": 1, "cells": ["公司中文名称", "甲甲科技有限公司"]},
        {"row_number": 2, "cells": ["WEEE产品信息"]},
        {"row_number": 3, "cells": ["类别", "品牌"]},
        {"row_number": 4, "cells": ["小型设备", "WEEEBrand"]},
        {"row_number": 5, "cells": ["电池产品信息"]},
        {"row_number": 6, "cells": ["类别", "品牌"]},
        {"row_number": 7, "cells": ["工业电池", "BatteryBrand"]},
    ]}]}
    result = extract_battery_items(project="德国电池法", company="甲甲科技有限公司", attachments=[attachment])
    assert [(item["brand"], item["category"]) for item in result["items"]] == [("BatteryBrand", "工业电池")]


def test_manual_confirmation_keeps_original_evidence():
    original = extract_battery_items(project="德国电池法", attachments=[{"filename": "STB 汽车启动电池合同.pdf"}])["items"][0]
    confirmed = normalize_battery_items([{**original, "brand": "DemoBrand"}], confirm_id=original["item_id"])
    assert confirmed[0]["battery_confirmed"]
    assert confirmed[0]["source"] == "附件合同名称"
    assert "STB" in confirmed[0]["evidence"]


def test_stage1_row_contains_battery_items_from_contract():
    from modules.field_extractor import FieldExtractor
    rows = FieldExtractor({}, [], ocr_fallback=False).extract_fields({
        "subject": "甲甲科技有限公司德国电池法注册",
        "body_text": "请办理德国电池法注册",
        "attachments": [{"filename": "甲甲科技有限公司-STB 汽车启动电池合同.pdf"}],
    })
    battery_rows = [row for row in rows if row.get("项目") == "德国电池法"]
    assert len(battery_rows) == 1
    items = json.loads(battery_rows[0]["德国电池法品类明细"])
    assert len(items) == 1
    assert items[0]["category_class"] == "启动电池"
    assert not items[0]["battery_confirmed"]


def test_battery_workorder_category_is_compared_independently():
    items = [{"brand": "DemoBrand", "category": "汽车启动电池", "category_class": "启动电池"}]
    matched = compare_battery_categories(items, [{"品类明细": [{"电池类别": "Starterbatterie", "品牌": "DemoBrand"}]}])
    assert matched["status"] == "matched"
    missing = compare_battery_categories(items, [{"品类明细": [{"电池类别": "工业电池", "品牌": "DemoBrand"}]}])
    assert missing["status"] == "missing"
    pending = compare_battery_categories(items, [{"工单编号": "WO-1"}])
    assert pending["status"] == "pending"


def test_battery_brand_and_type_cannot_be_cross_paired_between_rows():
    items = [{"brand": "DemoBrand", "category": "汽车启动电池", "category_class": "启动电池"}]
    result = compare_battery_categories(items, [{"品类明细": [
        {"品牌": "DemoBrand", "电池类别": "工业电池"},
        {"品牌": "OtherBrand", "电池类别": "Starterbatterie"},
    ]}])
    assert result["status"] == "missing"


def test_battery_brand_prefix_is_not_treated_as_exact_match():
    items = [{"brand": "DemoBrand", "category": "汽车启动电池", "category_class": "启动电池"}]
    result = compare_battery_categories(items, [{"品类明细": [
        {"品牌": "DemoBrand Plus", "电池类别": "Starterbatterie"},
    ]}])
    assert result["status"] == "missing"


def test_battery_extra_types_are_returned_separately():
    items = [{"brand": "DemoBrand", "category": "汽车启动电池", "category_class": "启动电池"}]
    result = compare_battery_categories(items, [{"品类明细": [
        {"品牌": "DemoBrand", "电池类别": "Starterbatterie"},
        {"品牌": "ExtraBrand", "电池类别": "Industriebatterie"},
    ]}])
    assert result["status"] == "matched"
    assert result["extra_items"] == [{"brand": "ExtraBrand", "category": "Industriebatterie",
                                     "category_class": "工业电池"}]


def make_store(tmp_path):
    path = tmp_path / "primary.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["邮件编号", "明细编号", "发件人邮箱", "发件日期", "邮件主题",
                  "代理", "客户公司名称", "标准化项目名称", "需求", "置信度",
                  "德国WEEE专项", "德国WEEE品类明细"])
    for index, project in enumerate(["德国WEEE", "德国电池法"]):
        sheet.append(["MIXED", f"D{index}", "fixture@example.test", "2026-08-26 09:00:00",
                      "Anonymous mixed projects", "Example Agent", "Example LLC", project,
                      "注册", "high", "是" if index == 0 else "",
                      json.dumps([item("W1")]) if index == 0 else ""])
    workbook.save(path)
    workbook.close()
    return WorkbenchStore(primary_path=str(path), review_path=str(tmp_path / "review.xlsx"),
                          filtered_path=str(tmp_path / "filtered.xlsx"),
                          state_path=str(tmp_path / "state.json"), test_mode=True)


def battery_detail(store):
    return next(d for m in store.snapshot()["mails"] for d in m["details"]
                if d["fields"]["program"] == "德国电池法")


def test_battery_draft_cannot_forge_confirmation():
    assert not normalize_battery_items([{**item(), "battery_confirmed": True}])[0]["battery_confirmed"]


def test_battery_edit_invalidates_confirmation_and_empty_list_is_pending():
    saved = normalize_battery_items([item()], confirm_id="B1")
    assert normalize_battery_items([item()], saved)[0]["battery_confirmed"]
    assert not normalize_battery_items([item(brand="Changed")], saved)[0]["battery_confirmed"]
    assert not battery_review("德国电池法", {"battery_items": []}, [], "confirmed")["confirmed"]


def test_battery_preserves_legacy_confirmation_but_return_blocks_it():
    assert battery_review("德国电池法", {}, [], "confirmed")["confirmed"]
    saved = {"battery_items": normalize_battery_items([item()], confirm_id="B1")}
    assert not battery_review("德国电池法", saved, [], "returned")["confirmed"]


@pytest.mark.parametrize("items,kwargs", [([item(), item()], {}), ([item()], {"confirm_id": "missing"}),
                                          (None, {}), (["invalid"], {})])
def test_invalid_battery_payload_rejected(items, kwargs):
    with pytest.raises(ValueError):
        normalize_battery_items(items, **kwargs)


def test_mixed_mail_battery_actions_do_not_change_weee(tmp_path):
    store = make_store(tmp_path)
    before = store.snapshot()["mails"][0]
    battery = battery_detail(store)
    weee_before = next(d["weee"] for d in before["details"] if d["weee"]["enabled"])
    payload = {"record_id": battery["id"], "mail_id": before["id"],
               "battery_items": [item(), item("B2", "Second")], "battery_item_id": "B1"}
    store.action({**payload, "action": "save_battery_draft"})
    assert not _detail_confirmation_ready(battery_detail(store))
    store.action({**payload, "action": "confirm_battery_item"})
    first = battery_detail(store)
    assert first["battery"]["items"][0]["battery_confirmed"]
    assert not _detail_confirmation_ready(first)
    store.action({**payload, "action": "confirm_battery_item", "battery_item_id": "B2"})
    assert _detail_confirmation_ready(battery_detail(store))
    weee_after = next(d["weee"] for d in store.snapshot()["mails"][0]["details"] if d["weee"]["enabled"])
    assert weee_after == weee_before
    reopened = WorkbenchStore(primary_path=store.primary_path, review_path=store.review_path,
                               filtered_path=store.filtered_path, state_path=store.state_path, test_mode=True)
    assert _detail_confirmation_ready(battery_detail(reopened))
    store.action({**payload, "action": "delete_battery_item", "battery_item_id": "B2"})
    current = battery_detail(store)
    assert len(current["battery"]["items"]) == 1
    store.action({**payload, "action": "save_battery_draft", "battery_items": [item(brand="Edited")]})
    assert not _detail_confirmation_ready(battery_detail(store))


def test_battery_action_rejects_wrong_project(tmp_path):
    store = make_store(tmp_path)
    weee = next(d for d in store.snapshot()["mails"][0]["details"] if d["weee"]["enabled"])
    with pytest.raises(ValueError, match="德国电池法"):
        store.action({"action": "confirm_battery_item", "record_id": weee["id"],
                      "battery_items": [item()], "battery_item_id": "B1"})


def test_battery_last_item_delete_is_not_completion(tmp_path):
    store = make_store(tmp_path)
    detail = battery_detail(store)
    payload = {"record_id": detail["id"], "battery_items": [item()], "battery_item_id": "B1"}
    store.action({**payload, "action": "confirm_battery_item"})
    assert _detail_confirmation_ready(battery_detail(store))
    store.action({**payload, "action": "delete_battery_item"})
    assert battery_detail(store)["battery"]["items"] == []
    assert not _detail_confirmation_ready(battery_detail(store))


def test_battery_draft_does_not_create_confirmation_event(tmp_path):
    store = make_store(tmp_path)
    detail = battery_detail(store)
    store.action({"action": "save_battery_draft", "record_id": detail["id"], "battery_items": [item()]})
    assert not battery_detail(store)["events"]


def test_battery_export_retains_independent_manual_fields(tmp_path, monkeypatch):
    import workbench_server
    monkeypatch.setattr(workbench_server, "MANUAL_OUTPUT", tmp_path / "export")
    store = make_store(tmp_path)
    detail = battery_detail(store)
    store.action({"action": "confirm_battery_item", "record_id": detail["id"],
                  "battery_items": [item()], "battery_item_id": "B1"})
    workbook = load_workbook(store.export(), read_only=True)
    rows = list(workbook.active.values)
    assert len(rows) == 2  # Unconfirmed WEEE business must not be exported.
    row = dict(zip(rows[0], rows[1]))
    assert row["标准化项目名称"] == "德国电池法"
    # 业务要求：导出表不带品类列（品类可留空），但业务字段必须保留。
    assert "德国电池法品类明细" not in rows[0]
    assert "品牌" not in rows[0] and "分类" not in rows[0]
    assert row["客户公司名称"] and row["代理"] is not None
    assert row.get("具体业务", "").startswith("德国电池法")
    workbook.close()


def test_battery_brand_candidate_strips_amazon_link():
    """R10：品牌候选混入亚马逊链接时，剥离链接、保留品牌部分。"""
    result = extract_battery_items(
        subject="示例-德国电池法注册",
        body="",
        attachments=[{
            "filename": "battery_list.xlsx",
            "structured_records": [{
                "record_type": "battery_section",
                "brand": "SexYln https://www.amazon.de/dp/B0G2BYTMKN",
                "category": "设备电池",
            }],
        }],
        project="德国电池法",
    )
    assert result["enabled"] is True
    brands = [item["brand"] for item in result["items"] if item["brand"]]
    assert "SexYln" in brands
    assert not any("http" in brand or "amazon" in brand for brand in brands)


def test_battery_brand_candidate_pure_link_is_dropped():
    """品牌候选只有链接时直接作废，不产生空壳品牌项。"""
    result = extract_battery_items(
        subject="示例-德国电池法注册",
        body="",
        attachments=[{
            "filename": "battery_list.xlsx",
            "structured_records": [{
                "record_type": "battery_section",
                "brand": "https://www.amazon.de/dp/B0G2BYTMKN",
                "category": "设备电池",
            }],
        }],
        project="德国电池法",
    )
    assert result["enabled"] is True
    assert all("amazon" not in (item["brand"] or "") for item in result["items"])


def test_battery_brand_candidate_strips_amazon_link():
    """R10：品牌候选混入亚马逊链接时，剥离链接、保留品牌部分。"""
    result = extract_battery_items(
        subject="示例-德国电池法注册",
        body="",
        attachments=[{
            "filename": "battery_list.xlsx",
            "structured_records": [{
                "record_type": "battery_section",
                "brand": "SexYln https://www.amazon.de/dp/B0G2BYTMKN",
                "category": "设备电池",
            }],
        }],
        project="德国电池法",
    )
    assert result["enabled"] is True
    brands = [item["brand"] for item in result["items"] if item["brand"]]
    assert "SexYln" in brands
    assert not any("http" in brand or "amazon" in brand for brand in brands)


def test_battery_brand_candidate_pure_link_is_dropped():
    """品牌候选只有链接时直接作废，不产生空壳品牌项。"""
    result = extract_battery_items(
        subject="示例-德国电池法注册",
        body="",
        attachments=[{
            "filename": "battery_list.xlsx",
            "structured_records": [{
                "record_type": "battery_section",
                "brand": "https://www.amazon.de/dp/B0G2BYTMKN",
                "category": "设备电池",
            }],
        }],
        project="德国电池法",
    )
    assert result["enabled"] is True
    assert all("amazon" not in (item["brand"] or "") for item in result["items"])


def test_battery_body_table_cells_on_separate_lines_are_parsed():
    """无附件邮件：正文表格（表头列 + 数据行）也要读出品牌与电池类型。"""
    from modules.mail_reader import _html_to_text
    html = ("<table><tr><td>公司名称</td><td>电池类型</td><td>品牌</td></tr>"
            "<tr><td>武汉示例网络科技有限公司</td><td>设备电池</td><td>Hethlif</td></tr></table>"
            "<p>示例代理-武汉示例网络科技有限公司 德国电池法 新增注册</p>")
    result = extract_battery_items(
        subject="示例代理-1家公司 德国电池法 新增注册",
        body=_html_to_text(html),
        attachments=[],
        project="德国电池法",
        company="武汉示例网络科技有限公司",
    )
    brands = [item["brand"] for item in result["items"] if item["brand"]]
    classes = [item["category_class"] for item in result["items"] if item["category_class"]]
    assert "Hethlif" in brands
    assert "设备电池" in classes
    assert result["status"] == "pending"


def test_battery_body_table_cells_on_separate_lines_are_parsed():
    """无附件邮件：正文表格（表头列 + 数据行）也要读出品牌与电池类型。"""
    from modules.mail_reader import _html_to_text
    html = ("<table><tr><td>公司名称</td><td>电池类型</td><td>品牌</td></tr>"
            "<tr><td>武汉示例网络科技有限公司</td><td>设备电池</td><td>Hethlif</td></tr></table>"
            "<p>示例代理-武汉示例网络科技有限公司 德国电池法 新增注册</p>")
    result = extract_battery_items(
        subject="示例代理-1家公司 德国电池法 新增注册",
        body=_html_to_text(html),
        attachments=[],
        project="德国电池法",
        company="武汉示例网络科技有限公司",
    )
    brands = [item["brand"] for item in result["items"] if item["brand"]]
    classes = [item["category_class"] for item in result["items"] if item["category_class"]]
    assert "Hethlif" in brands
    assert "设备电池" in classes
    assert result["status"] == "pending"


def test_application_form_brand_field_row_is_used_as_brand_evidence():
    """申请表「品牌」字段行（带必填标记与填写说明）要能抽出品牌。

    真实反馈：德国电池法 重庆振丰维，Sheet1 第10行写着
    `品牌 *（多个品牌，请用逗号隔开） KDTEND`，界面品牌却是空的。
    """
    attachment = {
        "filename": "重庆振丰维商贸有限公司＋电池法申请表.xlsx",
        "sheets": [{
            "sheet_name": "Sheet1",
            "text": "营业执照地址 * 重庆市丰都县三合街道龙腾大道\n"
                    "品牌 *（多个品牌，请用逗号隔开） KDTEND\n"
                    "法人姓名-拼音 * Yuanyuan Liu",
        }],
        "structured_records": [],
    }
    result = extract_battery_items(
        subject="美鸥跨境——2026新注册 德国电池法 BG1943\t重庆振丰维商贸有限公司",
        body="", attachments=[attachment], project="德国电池法",
        company="重庆振丰维商贸有限公司",
    )
    assert [entry["brand"] for entry in result["items"]] == ["KDTEND"]
    # 品牌补上了，类别仍留空由人工五选一，不趁机猜类别。
    assert result["items"][0]["category"] == ""


def test_application_form_brand_is_not_taken_from_other_company_sheet():
    attachment = {
        "filename": "共用申请表.xlsx",
        "sheets": [{"sheet_name": "Sheet1", "text": "品牌 KDTEND"}],
        "structured_records": [],
    }
    result = extract_battery_items(
        subject="德国电池法", body="", attachments=[attachment],
        project="德国电池法", company="重庆振丰维商贸有限公司",
    )
    assert [entry["brand"] for entry in result["items"]] == [""]


def test_blank_battery_snapshot_does_not_hide_recomputed_brand():
    """快照里的空壳不能盖掉阶段一后来补出的品牌，但人工确认结果要保住。"""
    shell = {"battery_items": [{"brand": "", "category": "", "category_class": "设备电池",
                                "source": "未见专用类别合同", "evidence": "x",
                                "item_id": "BAT-7027b55e833ddeed", "battery_confirmed": False}]}
    computed = [{"brand": "KDTEND", "category": "", "category_class": "设备电池",
                 "source": "申请表品牌字段", "evidence": "品牌 KDTEND"}]
    assert battery_review("德国电池法", shell, computed, "ready")["items"][0]["brand"] == "KDTEND"

    confirmed = {"battery_items": [{"brand": "", "category": "", "category_class": "设备电池",
                                    "battery_confirmed": True}]}
    assert battery_review("德国电池法", confirmed, computed, "confirmed")["items"][0]["brand"] == ""

    manual = {"battery_items": [{"brand": "手填品牌", "category": "工业电池",
                                "battery_confirmed": False}]}
    assert battery_review("德国电池法", manual, computed, "ready")["items"][0]["brand"] == "手填品牌"


def _sheet_text_values(text, filename="重庆振丰维商贸有限公司＋电池法申请表.xlsx"):
    from modules.battery_review import _sheet_brand_field_values
    attachment = {"filename": filename,
                  "sheets": [{"sheet_name": "Sheet1", "text": text}]}
    return [value for value, _ in _sheet_brand_field_values(attachment, "重庆振丰维商贸有限公司")]


def test_application_form_brand_row_rejects_header_and_instruction_rows():
    """「品牌」是长说明表里到处都是的词，只能认「行首标签 + 短品牌值」。

    实测 183 个申请表一次抽出 80 种取值，绝大多数是表头/说明文
    （logo（必填）…／德文／商标注册编号／t in Ireland?…）。
    """
    assert _sheet_text_values("品牌 logo（必填） 产品链接 预估销售量") == []
    assert _sheet_text_values("Brand register Nr.") == []
    assert _sheet_text_values("品牌 德文") == []
    assert _sheet_text_values("单个电池重量 品牌 品牌logo（必填）") == []
    # 说明文正文里出现 brand 不算字段行（标签必须在行首）。
    assert _sheet_text_values("Q2 您是否以自己的 brand 在该国销售") == []


def test_application_form_brand_row_keeps_real_brand_values():
    assert _sheet_text_values("品牌 *（多个品牌，请用逗号隔开） KDTEND") == ["KDTEND"]
    assert _sheet_text_values("品牌：Sandsddy") == ["Sandsddy"]
    assert _sheet_text_values("品牌 TianTian Co., Ltd.") == ["TianTian Co., Ltd."]


def _sheet_text_values(text, filename="重庆振丰维商贸有限公司＋电池法申请表.xlsx"):
    from modules.battery_review import _sheet_brand_field_values
    attachment = {"filename": filename,
                  "sheets": [{"sheet_name": "Sheet1", "text": text}]}
    return [value for value, _ in _sheet_brand_field_values(attachment, "重庆振丰维商贸有限公司")]


def test_application_form_brand_row_rejects_header_and_instruction_rows():
    """「品牌」是长说明表里到处都是的词，只能认「行首标签 + 短品牌值」。

    实测 183 个申请表一次抽出 80 种取值，绝大多数是表头/说明文
    （logo（必填）…／德文／商标注册编号／t in Ireland?…）。
    """
    assert _sheet_text_values("品牌 logo（必填） 产品链接 预估销售量") == []
    assert _sheet_text_values("Brand register Nr.") == []
    assert _sheet_text_values("品牌 德文") == []
    assert _sheet_text_values("单个电池重量 品牌 品牌logo（必填）") == []
    # 说明文正文里出现 brand 不算字段行（标签必须在行首）。
    assert _sheet_text_values("Q2 您是否以自己的 brand 在该国销售") == []


def test_application_form_brand_row_keeps_real_brand_values():
    assert _sheet_text_values("品牌 *（多个品牌，请用逗号隔开） KDTEND") == ["KDTEND"]
    assert _sheet_text_values("品牌：Sandsddy") == ["Sandsddy"]
    assert _sheet_text_values("品牌 TianTian Co., Ltd.") == ["TianTian Co., Ltd."]


def test_multi_brand_cell_is_split_but_company_suffix_is_not():
    """申请表写明「多个品牌，请用逗号隔开」，一格两品牌要拆成两条待确认。"""
    assert _sheet_text_values("品牌 mesterth,Naxidf") == ["mesterth", "Naxidf"]
    # 公司后缀的逗号不是多品牌分隔符。
    assert _sheet_text_values("品牌 TianTian Co., Ltd.") == ["TianTian Co., Ltd."]


@pytest.mark.parametrize("text", [
    "品牌 logo（必填） 单边最大尺寸（必填） 产品链接",
    "品牌 单个电池重量 Gewicht pro srück",
    "品牌 对应的品牌logo（必填）",
    "品牌 32.9g",
    "品牌 品牌1",
])
def test_body_brand_line_rejects_form_header_fragments(text):
    """正文里「品牌」后面跟表头/说明文时不能当品牌（实测这几种都会漏进来）。"""
    result = extract_battery_items(
        subject="德国电池法 注册 重庆振丰维商贸有限公司",
        body=f"重庆振丰维商贸有限公司 德国电池法\n{text}",
        attachments=[], project="德国电池法", company="重庆振丰维商贸有限公司",
    )
    assert [entry["brand"] for entry in result["items"] if entry["brand"]] == []


def test_saved_snapshot_drops_junk_brand_items():
    """早期误存进快照的「logo（必填）…」条目要自动消失，人工品牌不受影响。"""
    saved = {"battery_items": [
        {"brand": "logo（必填） 单边最大尺寸（必填） 产品链接", "category": "", "category_class": "设备电池"},
        {"brand": "单个电池重量", "category": "", "category_class": "设备电池"},
    ]}
    assert battery_review("德国电池法", saved, [], "ready")["items"] == []

    manual = {"battery_items": [{"brand": "Coolita", "category": "屏幕设备", "category_class": "工业电池"}]}
    assert battery_review("德国电池法", manual, [], "ready")["items"][0]["brand"] == "Coolita"


@pytest.mark.parametrize("text", [
    "品牌 logo（必填） 单边最大尺寸（必填） 产品链接",
    "品牌 单个电池重量 Gewicht pro srück",
    "品牌 对应的品牌logo（必填）",
    "品牌 32.9g",
    "品牌 品牌1",
])
def test_body_brand_line_rejects_form_header_fragments(text):
    """正文里「品牌」后面跟表头/说明文时不能当品牌（实测这几种都会漏进来）。"""
    result = extract_battery_items(
        subject="德国电池法 注册 重庆振丰维商贸有限公司",
        body=f"重庆振丰维商贸有限公司 德国电池法\n{text}",
        attachments=[], project="德国电池法", company="重庆振丰维商贸有限公司",
    )
    assert [entry["brand"] for entry in result["items"] if entry["brand"]] == []


def test_saved_snapshot_drops_junk_brand_items():
    """早期误存进快照的「logo（必填）…」条目要自动消失，人工品牌不受影响。"""
    saved = {"battery_items": [
        {"brand": "logo（必填） 单边最大尺寸（必填） 产品链接", "category": "", "category_class": "设备电池"},
        {"brand": "单个电池重量", "category": "", "category_class": "设备电池"},
    ]}
    assert battery_review("德国电池法", saved, [], "ready")["items"] == []

    manual = {"battery_items": [{"brand": "Coolita", "category": "屏幕设备", "category_class": "工业电池"}]}
    assert battery_review("德国电池法", manual, [], "ready")["items"][0]["brand"] == "Coolita"


@pytest.mark.parametrize("subject,body,company,expected", [
    ("德国电池新增品牌-2026.8.24",
     "您好，EPR_DE_025030602902\n深圳市远德景科技有限公司\n"
     "shenzhenshiyuandejingkejiyouxiangongsi\n新增品牌：Fyrove\n\n请安排注册，谢谢",
     "深圳市远德景科技有限公司", "Fyrove"),
    ("SY003-首阳跨境咨询-德国电池法+意大利包装法-Goffii",
     "主体：斯洛伐克\n公司名称：ID corp s. r. o\n"
     "安排项目：德国电池法+意大利包装法\n品牌名称：Goffii",
     "ID corp s. r. o", "Goffii"),
])
def test_brand_on_its_own_line_is_extracted(subject, body, company, expected):
    """品牌常单独成行，该行既无公司名也无「电池」字样，不能被上下文过滤整行丢掉。

    真实反馈：德国电池新增品牌（Fyrove）、首阳跨境（Goffii）卡片品牌一直是空的
    —— 界面上看到的「品牌」其实是输入框 placeholder。
    """
    result = extract_battery_items(subject=subject, body=body, attachments=[],
                                   project="德国电池法", company=company)
    assert [entry["brand"] for entry in result["items"] if entry["brand"]] == [expected]


def test_brand_line_fallback_skips_weee_only_lines():
    result = extract_battery_items(
        subject="德国电池法", body="深圳某某科技有限公司\nWEEE 品牌：Coolita 品类：屏幕设备",
        attachments=[], project="德国电池法", company="深圳某某科技有限公司",
    )
    assert [entry["brand"] for entry in result["items"] if entry["brand"]] == []


@pytest.mark.parametrize("subject,body,company,expected", [
    ("德国电池新增品牌-2026.8.24",
     "您好，EPR_DE_025030602902\n深圳市远德景科技有限公司\n"
     "shenzhenshiyuandejingkejiyouxiangongsi\n新增品牌：Fyrove\n\n请安排注册，谢谢",
     "深圳市远德景科技有限公司", "Fyrove"),
    ("SY003-首阳跨境咨询-德国电池法+意大利包装法-Goffii",
     "主体：斯洛伐克\n公司名称：ID corp s. r. o\n"
     "安排项目：德国电池法+意大利包装法\n品牌名称：Goffii",
     "ID corp s. r. o", "Goffii"),
])
def test_brand_on_its_own_line_is_extracted(subject, body, company, expected):
    """品牌常单独成行，该行既无公司名也无「电池」字样，不能被上下文过滤整行丢掉。

    真实反馈：德国电池新增品牌（Fyrove）、首阳跨境（Goffii）卡片品牌一直是空的
    —— 界面上看到的「品牌」其实是输入框 placeholder。
    """
    result = extract_battery_items(subject=subject, body=body, attachments=[],
                                   project="德国电池法", company=company)
    assert [entry["brand"] for entry in result["items"] if entry["brand"]] == [expected]


def test_brand_line_fallback_skips_weee_only_lines():
    result = extract_battery_items(
        subject="德国电池法", body="深圳某某科技有限公司\nWEEE 品牌：Coolita 品类：屏幕设备",
        attachments=[], project="德国电池法", company="深圳某某科技有限公司",
    )
    assert [entry["brand"] for entry in result["items"] if entry["brand"]] == []
