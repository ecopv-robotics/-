from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import Workbook

from utils.attachment_parser import parse_attachment
from utils.workbook_preview import PreviewError, download_archive_member, preview_attachment


def _xlsx_bytes(sheet_name="EPR申请表", rows=3):
    buffer = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for row in range(1, rows + 1):
        sheet.cell(row=row, column=1, value=f"字段{row}")
        sheet.cell(row=row, column=2, value=f"值{row}")
    workbook.save(buffer)
    return buffer.getvalue()


def test_direct_workbook_preview_pages_all_rows(tmp_path):
    workbook_path = tmp_path / "application.xlsx"
    workbook_path.write_bytes(_xlsx_bytes(rows=125))

    first = preview_attachment(workbook_path, workbook_path.name, offset=0, page_size=40)
    second = preview_attachment(workbook_path, workbook_path.name, offset=40, page_size=40)
    final = preview_attachment(workbook_path, workbook_path.name, offset=120, page_size=40)

    assert first["sheets"][0]["row_count"] == 125
    assert [row["row_number"] for row in first["sheets"][0]["rows"]] == list(range(1, 41))
    assert first["sheets"][0]["next_offset"] == 40
    assert second["sheets"][0]["rows"][0]["row_number"] == 41
    assert final["sheets"][0]["rows"][-1]["row_number"] == 125
    assert final["sheets"][0]["next_offset"] is None


def test_direct_workbook_preview_keeps_every_sheet(tmp_path):
    workbook_path = tmp_path / "multi_sheet.xlsx"
    workbook = Workbook()
    workbook.active.title = "EPR申请表"
    workbook.active["A1"] = "第一张表"
    second = workbook.create_sheet("公司信息")
    second["A1"] = "第二张表"
    workbook.save(workbook_path)

    result = preview_attachment(workbook_path, workbook_path.name)

    assert [sheet["sheet_name"] for sheet in result["sheets"]] == ["EPR申请表", "公司信息"]
    assert [sheet["rows"][0]["cells"][0] for sheet in result["sheets"]] == ["第一张表", "第二张表"]


def test_preview_reports_row_and_column_safety_limits(tmp_path, monkeypatch):
    workbook_path = tmp_path / "wide_form.xlsx"
    buffer = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    for column in range(1, 5):
        sheet.cell(row=1, column=column, value=f"字段{column}")
    for row in range(2, 7):
        sheet.cell(row=row, column=1, value=f"值{row}")
    workbook.save(buffer)
    workbook_path.write_bytes(buffer.getvalue())
    monkeypatch.setattr("utils.workbook_preview.MAX_PREVIEW_ROWS", 4)
    monkeypatch.setattr("utils.workbook_preview.MAX_PREVIEW_COLUMNS", 2)

    result = preview_attachment(workbook_path, workbook_path.name, page_size=20)
    sheet_preview = result["sheets"][0]

    assert sheet_preview["row_count"] == 6
    assert sheet_preview["column_count"] == 4
    assert sheet_preview["row_limit_reached"] is True
    assert sheet_preview["column_limit_reached"] is True
    assert sheet_preview["complete"] is False
    assert len(sheet_preview["rows"]) == 4
    assert len(sheet_preview["rows"][0]["cells"]) == 2


def test_zip_manifest_keeps_all_workbooks_and_duplicate_sheet_names(tmp_path):
    archive_path = tmp_path / "forms.zip"
    members = {
        "one/form.xlsx": _xlsx_bytes(rows=4),
        "two/form.xlsx": _xlsx_bytes(rows=5),
        "three/form.xlsx": _xlsx_bytes(rows=6),
    }
    with zipfile.ZipFile(archive_path, "w") as archive:
        for name, payload in members.items():
            archive.writestr(name, payload)

    manifest = preview_attachment(archive_path, archive_path.name)
    assert manifest["workbook_count"] == 3
    assert {item["path"] for item in manifest["members"]} == set(members)
    previews = [preview_attachment(archive_path, archive_path.name, path)
                for path in members]
    assert [item["sheets"][0]["sheet_name"] for item in previews] == ["EPR申请表"] * 3
    assert [item["sheets"][0]["row_count"] for item in previews] == [4, 5, 6]

    payload, filename = download_archive_member(archive_path, archive_path.name, "two/form.xlsx")
    assert filename == "form.xlsx"
    assert payload == members["two/form.xlsx"]


def test_archive_member_path_traversal_is_rejected(tmp_path):
    archive_path = tmp_path / "unsafe.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("../outside.xlsx", _xlsx_bytes())

    with pytest.raises(PreviewError, match="不安全成员路径"):
        preview_attachment(archive_path, archive_path.name)


def test_archive_parser_preserves_full_member_path_for_same_named_sheets(tmp_path):
    archive_path = tmp_path / "forms.zip"
    with zipfile.ZipFile(archive_path, "w") as archive:
        archive.writestr("first/application.xlsx", _xlsx_bytes(rows=2))
        archive.writestr("second/application.xlsx", _xlsx_bytes(rows=3))

    parsed = parse_attachment(str(archive_path), archive_path.name)
    assert {sheet["member_path"] for sheet in parsed["sheets"]} == {
        "first/application.xlsx", "second/application.xlsx",
    }


def test_macro_enabled_workbook_is_parsed_as_a_table(tmp_path):
    workbook_path = tmp_path / "application.xlsm"
    workbook_path.write_bytes(_xlsx_bytes(rows=2))

    parsed = parse_attachment(str(workbook_path), workbook_path.name)

    assert len(parsed["sheets"]) == 1
    assert parsed["sheets"][0]["preview_rows"][0]["cells"] == ["字段1", "值1"]
