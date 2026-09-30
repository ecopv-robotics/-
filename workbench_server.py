"""本机邮件询单人工复核工作台。

工作台是阶段一输出的可视化复核层，不替代现有邮件解析和阶段二 RPA：
原始 Excel 只读，人工修改保存在 storage/workbench_review.json，确认后另存
为 output/manual_review/workbench_reviewed_*.xlsx。
"""

from __future__ import annotations

import hashlib
import io
import json
import mimetypes
import os
import re
import socket
import threading
import unicodedata
import uuid
import webbrowser
from datetime import datetime, timezone
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
from urllib.parse import parse_qs, quote, urlparse

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from modules.project_normalizer import country_of_project
from modules.mail_filter import is_internal_sender_address
from modules.battery_review import battery_review, is_german_battery, normalize_battery_items
from workbench_database import WorkbenchDatabase, WORKBENCH_DATA_LOCK
from utils.runtime_paths import APP_ROOT, resolve_runtime_path


OUTPUT_ROOT = APP_ROOT / "output"
STAGE1_OUTPUT = OUTPUT_ROOT / "stage1_email"
STAGE2_OUTPUT = OUTPUT_ROOT / "stage2_workorder"
MANUAL_OUTPUT = OUTPUT_ROOT / "manual_review"
DEFAULT_PRIMARY = STAGE1_OUTPUT / "to_workorder_list.xlsx"
DEFAULT_REVIEW = STAGE1_OUTPUT / "to_review_list.xlsx"
DEFAULT_FILTERED = STAGE1_OUTPUT / "filtered_mail_record.xlsx"
DEFAULT_WORKORDER_RESULT = STAGE2_OUTPUT / "workorder_check_result.xlsx"
DEFAULT_MISSING_WORKORDER = STAGE2_OUTPUT / "漏单.xlsx"
LEGACY_PRIMARY = OUTPUT_ROOT / "to_workorder_list.xlsx"
LEGACY_REVIEW = OUTPUT_ROOT / "to_review_list.xlsx"
LEGACY_FILTERED = OUTPUT_ROOT / "filtered_mail_record.xlsx"
LEGACY_WORKORDER_RESULT = OUTPUT_ROOT / "workorder_check_result.xlsx"
LEGACY_MISSING_WORKORDER = OUTPUT_ROOT / "漏单.xlsx"
REVIEW_STATE = APP_ROOT / "storage" / "workbench_review.json"
HISTORY_DIR = APP_ROOT / "storage" / "workbench_history"
HISTORY_STATE = HISTORY_DIR / "mail_history.json"
HISTORY_OUTPUT = OUTPUT_ROOT / "workbench_history"
DATABASE_PATH = APP_ROOT / "storage" / "workbench.db"
WORKORDER_RETRY_QUEUE = APP_ROOT / "storage" / "workorder_retry_queue.json"
COMPLETED_HISTORY_XLSX = HISTORY_OUTPUT / "邮件处理完成总表.xlsx"
UNFINISHED_HISTORY_XLSX = HISTORY_OUTPUT / "邮件未完成总表.xlsx"
HTML_PATH = APP_ROOT / "workbench.html"
WORKBENCH_BUILD = "2026.09.30-r11"
SESSION_STATE = APP_ROOT / "session_state.json"
CONFIG_PATH = APP_ROOT / "config.yaml"

# 过滤日志里已经"定论"的过滤原因：证书/下号通知是规则确定的交付通知，
# 非注册业务(收款/报价/续费/信息变更/发票)是词表确定的运营事务，
# LLM确认=... 是二次识别已判定为非注册邮件。它们只参与右上角计数，
# 不再进入复核队列；其余过滤原因（可能误过滤）仍保留给人工判断。
# 这些原因已经由阶段一确定为不应进入询单复核：
# - 证书/下号通知、非注册业务、LLM确认：业务规则或意图模型的定论；
# - ECOPV 内部发件人/收件人：公司内部流转，不是客户询单。
# 它们仍保留在“已过滤邮件”视图供追溯，但不应被重新拼回主询单队列。
SETTLED_FILTER_HINTS = (
    "证书/下号通知",
    "非注册业务",
    "LLM确认",
    "ECOPV内部发件人邮箱",
    "ECOPV系统内部收件人邮箱",
)


def _is_settled_filter(reason: str) -> bool:
    text = _text(reason)
    return any(hint in text for hint in SETTLED_FILTER_HINTS)


def _is_ecopv_internal_sender(value: Any) -> bool:
    """判断发件人地址是否属于 ECOPV 内部域名。

    阶段一新邮件由 ``modules.mail_filter`` 负责过滤；工作台还需要对旧的
    主表/数据库记录执行同一闸门，否则历史“人工补全”记录会绕过邮件过滤，
    继续出现在询单队列。
    """
    return is_internal_sender_address(value)


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _read_imap_summary() -> Dict[str, Any]:
    """读取最近一次阶段一的原始 IMAP SEARCH 统计。"""
    path = OUTPUT_ROOT / "stage1_email" / "imap_summary.json"
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError, TypeError, json.JSONDecodeError):
        return {}
    if not isinstance(payload, dict):
        return {}
    try:
        total = max(0, int(payload.get("imap_read_total") or 0))
    except (TypeError, ValueError):
        total = 0
    try:
        parsed = max(0, int(payload.get("parsed_total") or 0))
    except (TypeError, ValueError):
        parsed = 0
    return {
        "imap_read_total": total,
        "parsed_total": parsed,
        "date_from": _text(payload.get("date_from")),
        "date_to": _text(payload.get("date_to")),
        "created_at": _text(payload.get("created_at")),
    }


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _json_list(value: Any) -> List[dict]:
    """读取阶段一写入的附件证据 JSON；旧文件/手工 Excel 为空时兼容。"""
    if isinstance(value, list):
        return [item for item in value if isinstance(item, dict)]
    text = _text(value)
    if not text:
        return []
    try:
        parsed = json.loads(text)
    except (TypeError, ValueError, json.JSONDecodeError):
        # 旧版为了适配 Excel 单元格上限把 JSON 从中间截断。不能把整列
        # 直接判成“无附件”：从截断前仍完整的 filename/row 对象恢复一份
        # 可预览的结构化快照，原始文件不存在时也能生成可打开的预览表。
        return _recover_truncated_attachment_evidence(text)
    return [item for item in parsed if isinstance(item, dict)] if isinstance(parsed, list) else []


def _weee_item_id(item: Dict[str, Any], index: int = 0) -> str:
    """为品牌/品类项目生成跨刷新稳定的编号。"""
    existing = _text(item.get("item_id"))
    if existing:
        return existing
    seed = "|".join(
        _text(item.get(key))
        for key in ("brand", "category", "category_original", "source", "evidence")
    ) or str(index)
    return f"WEEE-{hashlib.sha1(seed.encode('utf-8', errors='ignore')).hexdigest()[:16].upper()}"


def _with_weee_item_ids(items: Iterable[dict]) -> List[dict]:
    result: List[dict] = []
    for index, item in enumerate(items or []):
        if not isinstance(item, dict):
            continue
        copied = dict(item)
        copied["item_id"] = _weee_item_id(copied, index)
        result.append(copied)
    return result


def _should_refresh_legacy_weee_snapshot(
    saved: Dict[str, Any],
    saved_items: List[dict],
    current_items: List[dict],
    company: str,
) -> bool:
    """判断旧工作台快照是否应让位于最新阶段一结果。

    早期版本的 LLM 复提取会把同一附件中其他公司的品牌写入当前明细，
    并把这份错误结果保存为 ``weee_items``。后续阶段一 Excel 已经按公司
    正确拆成一行一项，但工作台原先始终优先读取旧快照，导致刷新后仍显示
    8 项/12 项。只迁移具有明确旧版标记、尚未确认且没有人工操作记录的
    快照；人工确认、人工删除或普通草稿均不覆盖。
    """
    if not saved_items or not current_items or len(saved_items) <= len(current_items):
        return False
    if _text(saved.get("weee_status")) in {"confirmed", "partial"}:
        return False
    events = saved.get("events")
    if isinstance(events, list) and events:
        return False
    if any(bool(item.get("weee_confirmed")) for item in saved_items):
        return False

    legacy_marked = any(
        _text(item.get("extraction_method")).startswith("LLM重新提取")
        or _text(item.get("llm_review_status")).casefold() in {"invalid", "failed", "error"}
        for item in saved_items
    )
    if not legacy_marked:
        return False

    company_key = _attachment_match_text(company)
    if not company_key:
        return False
    # 只有能证明旧快照混入了其他公司证据时才迁移；同一公司的多品牌草稿
    # 即使数量大于阶段一当前结果，也不在这里强行覆盖。
    foreign_evidence = 0
    for item in saved_items:
        evidence = _attachment_match_text(
            item.get("evidence") or " ".join(item.get("evidences") or [])
        )
        if evidence and company_key not in evidence:
            foreign_evidence += 1
    return foreign_evidence > 0


_WEEE_CATEGORY_NAMES = {
    "1": "热交换设备",
    "2": "屏幕和显示设备",
    "3": "灯具和光源",
    "4": "大型设备",
    "5": "小型设备",
    "6": "小型信息和电信设备",
}


def _normalize_weee_draft_items(raw_items: Any) -> List[Dict[str, Any]]:
    """规范化工作台 WEEE 草稿，不把输入草稿误标为已确认。"""
    if not isinstance(raw_items, list):
        raise ValueError("德国 WEEE 品牌/品类数据格式不正确")
    normalized: List[Dict[str, Any]] = []
    for index, raw_item in enumerate(raw_items):
        if not isinstance(raw_item, dict):
            continue
        item = dict(raw_item)
        item["item_id"] = _weee_item_id(item, index)
        item["brand"] = _text(item.get("brand"))
        item["category"] = _text(item.get("category") or item.get("category_original"))
        item["category_original"] = item["category"]
        category_class = _text(item.get("category_class"))
        if category_class not in _WEEE_CATEGORY_NAMES:
            category_class = ""
        item["category_class"] = category_class
        item["category_class_name"] = _WEEE_CATEGORY_NAMES.get(category_class, "")
        item["category_class_status"] = "matched" if category_class else (_text(item.get("category_class_status")) or "unmatched")
        # 前端在用户修改过的项目上会清除该标记；后端不允许草稿把未完成项目
        # 伪装成已确认，但保留没有发生修改的既有单项确认状态。
        item["weee_confirmed"] = bool(item.get("weee_confirmed"))
        normalized.append(item)
    if not normalized:
        raise ValueError("没有可保存的德国 WEEE 品牌/品类项目")
    return normalized


_WEEE_ATTACHMENT_PARSE_CACHE: Dict[str, Tuple[float, Dict[str, Any]]] = {}


def _attachment_match_text(value: Any) -> str:
    """用于把公司名与附件文件名做保守关联的规范化文本。"""
    return re.sub(r"[^0-9a-z\u4e00-\u9fff]+", "", _text(value).casefold())


def _recover_weee_items_from_cached_attachments(
    row: Dict[str, Any],
    existing_items: List[dict],
) -> List[dict]:
    """为旧阶段一结果回读附件缓存，补上当时漏掉的 WEEE 品牌/品类。

    旧版 Excel 可能已经写入 ``德国WEEE品类明细=[]``，但同一行仍保留了
    ``附件文件索引``。这种情况下仅刷新工作台不会重新跑阶段一，操作人员
    看到的卡片就永远没有可填写的品牌/品类。这里仅读取本行自己的缓存附件，
    优先按客户公司名匹配压缩包，不从产品常识猜字段；解析失败则保持空值。
    """
    if existing_items:
        return existing_items
    project = _text(row.get("标准化项目名称") or row.get("项目"))
    if not re.search(r"WEEE", project, re.I):
        return existing_items

    index_items = _json_list(row.get("附件文件索引"))
    if not index_items:
        return existing_items
    company = _attachment_match_text(row.get("客户公司名称") or row.get("客户"))
    archive_items = []
    for item in index_items:
        filename = _text(item.get("filename"))
        token = os.path.basename(_text(item.get("token")))
        if not filename or not token:
            continue
        if not Path(filename).suffix.casefold() in {".zip", ".rar", ".7z", ".xlsx", ".xlsm", ".xls"}:
            continue
        archive_items.append((filename, token))
    if not archive_items:
        return existing_items

    # 同一封邮件可能包含多个公司的压缩包；每条业务明细只回读自己的包。
    matched = [
        item for item in archive_items
        if company and company in _attachment_match_text(item[0])
    ]
    selected = matched or archive_items
    attachments: List[dict] = []
    cache_dir = (APP_ROOT / "cache" / "attachments").resolve()
    try:
        from utils.attachment_parser import parse_attachment
    except Exception:
        return existing_items
    for filename, token in selected:
        path = (cache_dir / token).resolve()
        try:
            path.relative_to(cache_dir)
        except ValueError:
            continue
        if not path.is_file():
            continue
        try:
            mtime = path.stat().st_mtime
        except OSError:
            continue
        cached = _WEEE_ATTACHMENT_PARSE_CACHE.get(str(path))
        if cached and cached[0] == mtime:
            parsed = cached[1]
        else:
            try:
                parsed = parse_attachment(str(path), filename)
            except Exception:
                continue
            _WEEE_ATTACHMENT_PARSE_CACHE[str(path)] = (mtime, parsed)
        if isinstance(parsed, dict):
            attachments.append(parsed)
    if not attachments:
        return existing_items

    try:
        from modules.weee_category_audit import extract_weee_items
        recovered = extract_weee_items(
            subject=_text(row.get("邮件主题")),
            body=_text(row.get("邮件正文原文") or row.get("邮件正文摘要(最多300字)")),
            attachments=attachments,
            project=project,
            llm_client=None,
        )
    except Exception:
        return existing_items
    items = recovered.get("items") if isinstance(recovered, dict) else []
    return _with_weee_item_ids(items) if items else existing_items


def _json_object(value: Any) -> Dict[str, Any]:
    """读取阶段一/阶段二写入的专项核对 JSON，失败时返回空对象。"""
    if isinstance(value, dict):
        return value
    text = _text(value)
    if not text:
        return {}
    try:
        parsed = json.loads(text)
    except (TypeError, ValueError, json.JSONDecodeError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _recover_truncated_attachment_evidence(text: str) -> List[dict]:
    """尽可能恢复被 Excel 单元格长度截断的附件证据。

    恢复目标是证据展示，不把不完整 JSON 当成完整数据：只接受
    ``json.JSONDecoder.raw_decode`` 能完整解析的附件名、工作表行和文本片段。
    这样旧数据不会再静默变成空附件，新数据仍由正常 json.loads 处理。
    """
    decoder = json.JSONDecoder()
    filename_matches = list(re.finditer(r'"filename"\s*:', text))
    if not filename_matches:
        return []
    recovered: List[dict] = []
    seen_names = set()
    for index, match in enumerate(filename_matches):
        start = match.end()
        while start < len(text) and text[start].isspace():
            start += 1
        try:
            filename, _ = decoder.raw_decode(text, start)
        except (TypeError, ValueError, json.JSONDecodeError):
            continue
        filename = _text(filename)
        if not filename:
            continue
        end = filename_matches[index + 1].start() if index + 1 < len(filename_matches) else len(text)
        segment = text[start:end]
        rows: List[dict] = []
        seen_rows = set()
        for row_match in re.finditer(r'\{"row_number"\s*:', segment):
            try:
                row, _ = decoder.raw_decode(segment, row_match.start())
            except (TypeError, ValueError, json.JSONDecodeError):
                continue
            if not isinstance(row, dict):
                continue
            key = (_text(row.get("row_number")), json.dumps(row.get("cells") or [], ensure_ascii=False))
            if key in seen_rows:
                continue
            seen_rows.add(key)
            rows.append({
                "row_number": row.get("row_number") or "?",
                "cells": [_text(cell)[:160] for cell in (row.get("cells") or [])[:24]],
            })
        sheet_names = []
        for sheet_match in re.finditer(r'"sheet_name"\s*:', segment):
            pos = sheet_match.end()
            while pos < len(segment) and segment[pos].isspace():
                pos += 1
            try:
                sheet_name, _ = decoder.raw_decode(segment, pos)
            except (TypeError, ValueError, json.JSONDecodeError):
                continue
            if _text(sheet_name) and _text(sheet_name) not in sheet_names:
                sheet_names.append(_text(sheet_name))
        text_preview = ""
        text_match = re.search(r'"text"\s*:', segment)
        if text_match:
            pos = text_match.end()
            while pos < len(segment) and segment[pos].isspace():
                pos += 1
            try:
                text_value, _ = decoder.raw_decode(segment, pos)
                text_preview = _text(text_value)[:1200]
            except (TypeError, ValueError, json.JSONDecodeError):
                pass
        item: Dict[str, Any] = {"filename": filename, "records": [], "sheets": [], "text": text_preview}
        if rows:
            item["sheets"] = [{"sheet_name": sheet_names[0] if sheet_names else "工作表", "rows": rows}]
        if filename not in seen_names:
            seen_names.add(filename)
            recovered.append(item)
    return recovered


def _looks_like_company_name(value: Any) -> bool:
    """仅接受带法定后缀的附件公司候选，拒绝城市/提示语。"""
    text = _text(value)
    if not text or _is_non_company_customer_value(text):
        return False
    if re.search(r"(?:有限责任公司|有限公司|股份有限公司|集团公司|合伙企业|公司|（个体工商户）|\(个体工商户\)|经营部|门市部|服务部|商店|商行|工厂|工作室|店)$", text):
        return len(text) >= 4
    return bool(re.search(
        r"(?i)(?:^|\s)(?:limited|ltd\.?|llc|gmbh|s\.?\s*p\.?\s*z\.?\s*o\.?\s*o\.?|"
        r"sarl|sas|sia|b\.?v\.?|a\.?b\.?|s\.?l\.?|oy|a\.?/\.?s\.?|plc|inc\.?|corp\.?|company)"
        r"(?:$|\s|[,.)])",
        text,
    ))


def _company_evidence_candidates(value: Any) -> List[str]:
    """从正文、主题或附件文件名中提取有法定后缀的公司候选。

    这是工作台读取旧阶段一 Excel 时的确定性回填兜底，不替代阶段一的
    完整解析。优先识别“公司：”标签后的值，并按常见文件名/主题分隔符
    拆分，避免把“德国一次性塑料法-示例a2f907fc有限公司.zip”整体
    当成公司名。
    """
    text = _text(value).replace("\r\n", "\n").replace("\r", "\n")
    if not text:
        return []
    suffix = (
        r"有限责任公司|股份有限公司|集团有限公司|有限公司|责任公司|合伙企业|公司|"
        r"经营部|门市部|服务部|商店|商行|工厂|工作室|店|（个体工商户）|\(个体工商户\)"
    )
    boundary_pattern = re.compile(
        rf"(?:^|[-_—–\s:：;；,，])"
        rf"([\u4e00-\u9fffA-Za-z0-9·（）()&.'’ ]{{2,80}}(?:{suffix}))"
        rf"(?=$|[-_—–\s:：;；,，.])",
        re.I,
    )

    def is_noise(candidate: str) -> bool:
        compact = re.sub(r"\s+", "", candidate)
        if compact in {"公司", "客户公司", "公司名称", "申请单位", "待确认", "未知"}:
            return True
        # 说明句可能被后缀正则截成“麻烦安排以下公司”这类伪主体。
        return bool(re.search(r"(?:安排|以下|下列|上述|涉及|通知|告知).{0,16}公司$", compact))

    found: List[str] = []

    def add_fragment(fragment: str) -> None:
        fragment = re.sub(
            r"^\s*(?:公司名称|公司名|客户公司|申请单位|公司)\s*[:：]?\s*",
            "",
            _text(fragment),
            flags=re.I,
        ).strip(" \t\n,，;；:+＋-—()（）[]【】")
        if not fragment:
            return
        pieces = re.split(r"[-_—–]", fragment) if re.search(r"[-_—–]", fragment) else [fragment]
        for piece in pieces:
            piece = piece.strip(" \t\n,，;；:+＋()（）[]【】")
            if not piece:
                continue
            matches = list(boundary_pattern.finditer(piece))
            candidates = [match.group(1).strip(" .,-") for match in matches]
            if _looks_like_company_name(piece) and not is_noise(piece):
                candidates.append(piece)
            for candidate in candidates:
                candidate = candidate.strip(" .,-")
                if _looks_like_company_name(candidate) and not is_noise(candidate):
                    key = re.sub(r"\s+", "", candidate).casefold()
                    if key and key not in {re.sub(r"\s+", "", item).casefold() for item in found}:
                        found.append(candidate)

    lines = text.split("\n")
    # 标签独占一行时，取下一行作为值；标签与值同一行时直接取冒号后的值。
    for index, line in enumerate(lines):
        label = re.match(
            r"^\s*(?:公司名称|公司名|客户公司|申请单位|公司)\s*[:：]?\s*(.*)$",
            line,
            re.I,
        )
        if not label:
            continue
        value_part = label.group(1).strip()
        if not value_part:
            for next_line in lines[index + 1:]:
                value_part = next_line.strip()
                if value_part:
                    break
        add_fragment(value_part)

    # 正文/主题按行，附件名称按分隔符均可复用同一确定性提取。
    for line in lines:
        add_fragment(line)
    if len(lines) > 1:
        add_fragment(text)
    return found


def _merge_saved_fields(base: Dict[str, str], saved: Dict[str, Any]) -> Dict[str, str]:
    """合并人工覆盖层，避免旧缓存的空公司字段遮住新证据回填。

    旧版状态文件可能在首次打开时保存了 ``company: ""``，这不是一次
    有意的清空操作；只有带有字段变更事件的空值才视为人工明确清空。
    """
    result = dict(base)
    saved_fields = saved.get("fields") if isinstance(saved, dict) and isinstance(saved.get("fields"), dict) else {}
    explicitly_cleared = set()
    for event in (saved.get("events") if isinstance(saved, dict) and isinstance(saved.get("events"), list) else []):
        if not isinstance(event, dict):
            continue
        changes = event.get("changed_fields") if isinstance(event.get("changed_fields"), dict) else {}
        for key, change in changes.items():
            if isinstance(change, dict) and not _text(change.get("after")):
                explicitly_cleared.add(key)
    placeholders = {"待确认", "未知", "unknown", "none", "null", "空"}
    for key, value in saved_fields.items():
        if key not in result:
            continue
        text = _text(value)
        if key == "company" and result.get(key) and text.casefold() in placeholders:
            continue
        if text or not result.get(key) or key in explicitly_cleared:
            result[key] = text
    return result


def _attachment_sheet_company_candidates(attachment: Dict[str, Any]) -> List[Tuple[str, str, str]]:
    """读取附件表格嵌套 cells 中的公司中文名/英文名。

    阶段一保存的 xlsx 证据不是传统 ``records.raw_text``，而是
    ``sheets[].rows[].cells``。一个表格行可能同时有中文名和英文名，
    也可能中文单元格为空；这里返回候选、证据类型和行定位，交给上层按
    当前业务明细的正文/来源做最终选择，不从城市或注册号推断公司名。
    """
    candidates: List[Tuple[str, str, str]] = []
    filename = _text(attachment.get("filename"))
    for sheet in attachment.get("sheets") or []:
        if not isinstance(sheet, dict):
            continue
        sheet_name = _text(sheet.get("sheet_name") or "工作表")
        rows = sheet.get("rows") or sheet.get("preview_rows") or []
        for row in rows:
            if not isinstance(row, dict):
                continue
            cells = row.get("cells") or row.get("values") or []
            if not isinstance(cells, list):
                continue
            row_number = _text(row.get("row_number") or "?")
            locator = f"附件表格：{filename} / {sheet_name} 第{row_number}行"
            for index, raw_label in enumerate(cells):
                label = _text(raw_label)
                compact = re.sub(r"\s+", "", label).lower()
                if not (
                    re.search(r"公司中文名称|中文公司名称", compact)
                    or re.search(r"companyname.*(?:chinese|in chinese)", compact)
                    or re.search(r"公司英文名称|英文公司名称", compact)
                    or re.search(r"companyname.*(?:english|in english)", compact)
                ):
                    continue
                english = bool(re.search(r"英文|english", compact))
                for raw_value in cells[index + 1:]:
                    value = _text(raw_value)
                    next_compact = re.sub(r"\s+", "", value).lower()
                    if not value:
                        continue
                    if (
                        re.search(r"公司中文名称|中文公司名称|公司英文名称|英文公司名称", next_compact)
                        or re.search(r"companyname", next_compact)
                    ):
                        break
                    if _looks_like_company_name(value):
                        kind = "附件表格英文公司名" if english else "附件表格中文公司名"
                        candidates.append((value, kind, locator))
                    break
    return candidates


def _subject_loose_company(subject: Any) -> str:
    """从旧阶段一记录的主题中恢复无公司后缀的英文商号。

    只处理“案件编号 + 英文商号 + 国家/项目”这一种明确格式，避免把
    中文正文中的说明句重新猜成公司。该兜底用于读取旧 Excel，新的阶段一
    结果仍以 field_extractor 的完整证据链为准。
    """
    text = _text(subject)
    code = re.search(
        r"(?<![A-Za-z0-9])(?:[A-Za-z]{1,3}\s*[-_]\s*)?"
        r"[A-Za-z]{1,8}\s*[-_]?\s*\d{2,8}(?![A-Za-z0-9])",
        text,
        re.I,
    )
    if not code:
        return ""
    tail = text[code.end():].strip(" \t-—_:：+＋")
    boundary = re.search(
        r"(?:德国|比利时|比利時|法国|法國|意大利|義大利|西班牙|荷兰|荷蘭|"
        r"波兰|波蘭|瑞典|爱尔兰|愛爾蘭|葡萄牙|奥地利|奧地利|芬兰|芬蘭|"
        r"WEEE|EEE|EPR|电池法|電池法|包装法|包裝法|注册|註冊|新注册|新註冊|"
        r"申报|申報|注销|註銷|修改|变更|變更)",
        tail,
        re.I,
    )
    candidate = (tail[:boundary.start()] if boundary else tail).strip(" \t-—_:：+＋,，;；")
    english = re.fullmatch(r"[A-Za-z][A-Za-z0-9 .,&'’()\-]{3,80}", candidate)
    chinese = re.fullmatch(
        r"[\u4e00-\u9fffA-Za-z0-9·（）()]{4,60}"
        r"(?:经营部|门市部|服务部|商店|商行|工厂|工作室|店|有限公司|公司|（个体工商户）|\(个体工商户\))",
        candidate,
    )
    if not english and not chinese:
        return ""
    if english and len(candidate.split()) < 2:
        return ""
    if _is_non_company_customer_value(candidate):
        return ""
    return candidate


def _repair_attachment_fields(row: Dict[str, Any]) -> Dict[str, Any]:
    """从附件证据行修复旧版“城市冒充公司”的阶段一结果。

    早期导出曾把注册表的 City 列写进客户公司字段。附件证据保留了
    ``代理 | 编号 | 公司中文名`` 的原始顺序，因此无需重新下载邮件即可
    修复已有工作台数据；人工状态仍由后续 state 覆盖。
    """
    if not isinstance(row, dict):
        return row
    current_company = _text(row.get("客户公司名称") or row.get("客户"))
    # 即使旧行没有附件证据，也能安全拆分“编号 + 法定英文名称”。
    code_prefix = re.match(
        r"^\s*(?:[A-Za-z]{1,3}\s*[-_]\s*)?[A-Za-z]{1,8}"
        r"\s*[-_]?\s*\d{2,8}\s+(.+?)\s*$",
        current_company,
    )
    if code_prefix and _looks_like_company_name(code_prefix.group(1)):
        cleaned_company = _text(code_prefix.group(1)).strip(" .,-")
        row["客户公司名称"] = cleaned_company
        row["客户"] = cleaned_company
        row["客户提取来源"] = "公司编号与法定名称拆分"
        current_company = cleaned_company

    # 代理名有时直接粘在客户公司前面（例如 TBA示例b691064e有限公司）。
    # 只使用当前明细已确认的代理值，并要求剩余文本仍是合法公司候选，避免
    # 把真实以同样字母开头的公司名误删。
    agent_aliases = []
    for raw_agent in re.split(r"[、,，;；/／|]", _text(row.get("代理") or row.get("代理(可能空)"))):
        alias = raw_agent.strip()
        if len(alias) >= 2:
            agent_aliases.append(alias)
    for alias in sorted(set(agent_aliases), key=len, reverse=True):
        if len(current_company) <= len(alias) or current_company[:len(alias)].casefold() != alias.casefold():
            continue
        remainder = current_company[len(alias):].lstrip(" \t-—_:：+＋/／")
        if remainder and _looks_like_company_name(remainder):
            row["客户公司名称"] = remainder.strip(" .,-")
            row["客户"] = row["客户公司名称"]
            row["客户提取来源"] = "代理前缀清洗"
            current_company = row["客户公司名称"]
            break
    # 旧记录中“不能与其它公司”“国公司”等确定性说明值可能遮住了主题中
    # 明确的英文商号。能按案件编号和国家/项目边界恢复时直接修复；否则保留
    # 空值/人工复核，不再把说明文字显示成公司。
    if _is_non_company_customer_value(current_company):
        fallback = _subject_loose_company(
            row.get("邮件主题") or row.get("主题") or row.get("subject")
        )
        if fallback:
            row["客户公司名称"] = fallback
            row["客户"] = fallback
            row["客户提取来源"] = "主题案件编号边界恢复"
            current_company = fallback

    # 旧版阶段一可能已经把正文/附件解析结果落成空值或“待确认”。
    # 这类记录即使没有结构化附件证据，也能从正文明确标签和附件文件名
    # 安全回填；已有合法公司名不覆盖，避免把人工确认结果改掉。
    company_placeholder = {"待确认", "未知", "unknown", "none", "null", "空"}
    company_missing = (
        not current_company
        or current_company.casefold() in company_placeholder
        or _is_non_company_customer_value(current_company)
    )
    if company_missing:
        evidence_sources = (
            ("正文确定性回填", ("邮件正文原文", "邮件正文", "正文", "body_text", "邮件正文摘要(最多300字)")),
            ("主题确定性回填", ("邮件主题", "主题", "subject")),
            ("附件文件名确定性回填", ("附件名称",)),
        )
        repaired_company = ""
        repaired_source = ""
        for source_label, keys in evidence_sources:
            for key in keys:
                for candidate in _company_evidence_candidates(row.get(key)):
                    repaired_company = candidate
                    repaired_source = source_label
                    break
                if repaired_company:
                    break
            if repaired_company:
                break
        # 某些中间版本只保留了 JSON 附件索引，没有“附件名称”列；
        # 文件名仍是可信证据，但只读取安全索引中的 basename。
        if not repaired_company:
            indexed_names = [
                _text(item.get("filename"))
                for item in _json_list(row.get("附件文件索引"))
                if isinstance(item, dict) and _text(item.get("filename"))
            ]
            for filename in indexed_names:
                candidates = _company_evidence_candidates(filename)
                if candidates:
                    repaired_company = candidates[0]
                    repaired_source = "附件文件名确定性回填"
                    break
        if repaired_company:
            row["客户公司名称"] = repaired_company
            row["客户"] = repaired_company
            row["客户提取来源"] = repaired_source
            # 旧阶段一/LLM 校验提示可能仍保留“客户公司字段缺失”或
            # `company: 客户公司名称为空`。字段已由确定性证据回填后清理
            # 这两段矛盾提示，保留其它人工复核原因不变。
            stale_hint_parts = []
            for part in re.split(r"[；;]\s*", _text(row.get("人工复核提示"))):
                compact_part = re.sub(r"\s+", "", part)
                if "客户公司字段缺失" in compact_part or "company:客户公司名称为空" in compact_part:
                    continue
                if part.strip():
                    stale_hint_parts.append(part.strip())
            row["人工复核提示"] = "；".join(stale_hint_parts)
            current_company = repaired_company
    evidence = _json_list(row.get("附件证据"))
    if not evidence:
        return row
    source = _text(row.get("附件明细来源"))
    code = _text(row.get("客户编号"))
    company_candidates: List[str] = []
    sheet_company_candidates: List[Tuple[str, str, str]] = []
    agent_candidate = ""
    for attachment in evidence:
        for candidate, kind, locator in _attachment_sheet_company_candidates(attachment):
            if source and locator not in source:
                continue
            sheet_company_candidates.append((candidate, kind, locator))
        for record in attachment.get("records") or []:
            if not isinstance(record, dict):
                continue
            filename = _text(record.get("attachment_name") or attachment.get("filename"))
            sheet = _text(record.get("sheet_name") or "工作表")
            number = _text(record.get("row_number") or "?")
            locator = f"附件表格：{filename} / {sheet} 第{number}行"
            if source and locator not in source:
                continue
            raw = _text(record.get("raw_text"))
            parts = [_text(part) for part in raw.split("|") if _text(part)]
            if code and parts:
                code_index = next(
                    (index for index, part in enumerate(parts) if part == code), -1
                )
                if code_index >= 0:
                    if code_index > 0:
                        before = parts[code_index - 1]
                        if before and not re.search(r"\d", before) and len(before) <= 40:
                            agent_candidate = before
                    if code_index + 1 < len(parts):
                        after = parts[code_index + 1]
                        if _looks_like_company_name(after):
                            company_candidates.append(after)
            for part in parts:
                if _looks_like_company_name(part):
                    company_candidates.append(part)
    company_candidates = list(dict.fromkeys(company_candidates))
    candidate_pairs: List[Tuple[str, str]] = [
        (candidate, "附件表格中文公司名") for candidate in company_candidates
    ] + [
        (candidate, kind) for candidate, kind, _ in sheet_company_candidates
    ]
    unique_pairs: List[Tuple[str, str]] = []
    seen_candidates = set()
    for candidate, kind in candidate_pairs:
        compact = re.sub(r"\s+", "", candidate).casefold()
        if compact and compact not in seen_candidates:
            seen_candidates.add(compact)
            unique_pairs.append((candidate, kind))

    # 附件里可能是一个压缩包的多家公司。优先选择在本封邮件主题/正文中
    # 实际出现的候选；若来源已经精确到一行且只剩一个候选，则可直接采用。
    evidence_text = " ".join(
        _text(row.get(key))
        for key in ("邮件主题", "主题", "subject", "邮件正文", "邮件正文原文", "邮件正文摘要(最多300字)", "正文", "body_text")
    )
    compact_evidence = re.sub(r"\s+", "", evidence_text).casefold()
    matched_pairs = [
        pair for pair in unique_pairs
        if re.sub(r"\s+", "", pair[0]).casefold() in compact_evidence
    ]
    # A compact preview is not a complete applicant inventory. Its sole visible
    # candidate must not become an inferred sole company in the source workbook.
    incomplete_preview = any(attachment.get("preview_truncated") for attachment in evidence)
    usable_pairs = matched_pairs or (
        unique_pairs if len(unique_pairs) == 1 and (source or not incomplete_preview) else []
    )
    # 同一行同时有中文和英文名时，中文名优先；中文为空时才会落到英文名。
    usable_pairs.sort(key=lambda pair: (0 if pair[1] == "附件表格中文公司名" else 1, -len(pair[0])))
    replacement = next(
        (
            (candidate, kind)
            for candidate, kind in usable_pairs
            if candidate != current_company
        ),
        ("", ""),
    )
    if replacement[0] and not _looks_like_company_name(current_company):
        row["客户公司名称"] = replacement[0]
        row["客户"] = replacement[0]
        row["客户提取来源"] = replacement[1]
    if agent_candidate and not _text(row.get("代理")):
        row["代理"] = agent_candidate
        row["代理(可能空)"] = agent_candidate
        row["代理匹配方式"] = "附件表格代理列"
    return row


def _prepare_workbench_row(row: Dict[str, Any]) -> Dict[str, Any]:
    """保留无法确认公司的明细，但不把说明句显示成公司名称。

    ``一家公司``、``非中国公司``、表单字段标签等值不能作为客户主体。
    以前读取阶段直接丢弃整行，结果邮件和业务明细数量减少，操作人员也看不到
    需要人工补全的邮件。现在保留这条明细，清空展示用公司字段并标记为低置信度；
    原始值放在内部审计字段中，便于追溯而不会误导工单核对。
    """
    if not isinstance(row, dict):
        return row
    company = _text(row.get("客户公司名称") or row.get("客户") or row.get("company"))
    if company and _is_non_company_customer_value(company):
        row.setdefault("原始客户公司名称", company)
        row["客户公司名称"] = ""
        row["客户"] = ""
        row["company"] = ""
        row["客户提取来源"] = "未识别，待人工确认"
        row["人工复核提示"] = _text(row.get("人工复核提示")) or "邮件中未找到可确认的公司名称"
        row["置信度"] = "low"
    return row


def _is_non_company_customer_value(value: Any) -> bool:
    """判断历史导入的客户值是否是 EPR 表单标签而非主体名称。

    旧版本曾把 EPR 申请表中的“POA/注册资本/签字时间”等字段标签写成
    客户。这里仅在工作台读取时隐藏这些确定性脏记录，不删除数据库历史；
    重新运行阶段一后，正确的客户记录会正常进入队列。
    """
    text = _text(value)
    if not text:
        return False
    compact = re.sub(r"\s+", "", text).lower()
    hints = (
        "poa", "legalrepresentative", "legalperson", "legalpositions",
        "nameoflegalperson", "placeofsignature", "signingtime",
        "registrationcapital", "companyname", "companyaddress", "companybusiness",
        "plz", "postcode", "amazonlink", "shoplink", "e-mail", "email", "tel", "phone",
        "legrepresentativename", "companyregistrationnumber", "registrationnumber", "uscc",
        "营业执照", "公司名称", "公司中文名称", "公司英文名称", "中文公司名称", "英文公司名称",
        "公司中文名", "公司英文名", "公司地址", "公司注册", "注册资本", "法人",
        "公司成立日期", "成立日期", "签字", "签署", "职位", "联系信息", "联系人", "联系电话", "邮箱",
        "邮政编码", "邮编", "地址", "姓名", "身份证", "护照", "性别", "店铺链接",
        "平台信息", "服务内容", "服务的国家", "销售量", "预计销售", "说明", "备注", "请提供", "请选择", "填写",
        "注意事项", "不能提供", "提供", "请客户", "确认好", "注册类别", "翻译公司", "盖章",
        "回收公司", "适用国家", "所有国家", "产品图片或说明书", "资料列表", "不用提供",
        "保证有就可以", "否则不接单", "要求北爱公司", "非中国公司", "中国公司",
        "外国公司", "国公司", "不能与其它公司", "不能与其他公司",
    )
    if any(hint in compact for hint in hints):
        return True
    # 旧版阶段一会把正文引导语中的公司后缀截出来，例如
    # “示例9ad3384d有限公司……提交……名单”。
    # 这不是客户主体；重新解析前先从工作台显示层隐藏这类历史脏行，
    # 但不删除数据库中的原始记录，便于审计追溯。
    if re.match(r"^(?:以下|下面|下列|现将|本次)(?:为|是)?", text, re.I) and re.search(
        r"(?:提交|报送|发送|提供|列出|名单|新注册|申请)", text, re.I
    ):
        return True
    # 旧行里通常只剩公司后缀前的截断值（如“以下为某某有限公司”），
    # 后面的“提交名单”已经不在客户字段中，因此单独按前缀+公司后缀识别。
    if re.match(r"^(?:以下|下面|下列)(?:为|是)", text, re.I) and re.search(
        r"(?:有限责任公司|股份有限公司|集团有限公司|有限公司|责任公司|公司|企业)$",
        text,
        re.I,
    ):
        return True
    if "@" in text or re.search(r"https?://|www\.", text, re.I):
        return True
    if re.fullmatch(r"[+()\-\s\d]{6,}", text):
        return True
    if re.fullmatch(r"[A-Za-z]{1,8}[-_]?\d{4,}", text):
        return True
    if re.fullmatch(r"[0-9一二三四五六七八九十多几]*\s*家\s*(?:公司|主体|企业)", text, re.I):
        return True
    if re.fullmatch(
        r"[0-9一二三四五六七八九十多几]+\s*家\s*(?:公司|主体|企业)"
        r"(?:\s*[-—:：,，/／+＋].*)?",
        text,
        re.I,
    ):
        return True
    # 旧版表单解析会把法人姓名写进客户列（如 Huiming Wu）。没有公司后缀、
    # 仅由英文名组成的值没有足够主体证据，工作台不应继续展示为客户公司。
    if re.fullmatch(r"[A-Z][a-z]{1,24}(?:\s+[A-Z][a-z]{1,24}){1,3}", text) and not re.search(
        r"\b(?:inc|incorporated|ltd|limited|llc|llp|plc|corp|corporation|gmbh|ug|srl|bv|co)\.?$", text, re.I
    ):
        return True
    return len(text) > 120


def _mail_identity_tuple(row: Dict[str, Any]) -> Tuple[str, str, str]:
    """返回用于替换同一封邮件旧版本记录的稳定身份。"""
    sender = _text(row.get("发件人邮箱") or row.get("sender_email") or row.get("sender")).lower()
    raw_date = _text(row.get("发件日期") or row.get("date"))
    parsed_date = _mail_date_sort_value(raw_date)
    date = (
        parsed_date.strftime("%Y-%m-%d %H:%M:%S")
        if parsed_date != datetime.min
        else raw_date.replace("T", " ")
    )
    subject = _text(row.get("邮件主题") or row.get("主题") or row.get("subject"))
    return sender, date, subject


def _workorder_detail_tuple(row: Dict[str, Any]) -> Tuple[str, str, str, str, str, str]:
    """阶段二结果与人工回写结果的稳定业务明细身份。"""
    mail = _mail_identity_tuple(row)
    company = _text(row.get("客户公司名称") or row.get("客户") or row.get("company"))
    project = _text(row.get("标准化项目名称") or row.get("项目") or row.get("program"))
    request = _text(row.get("需求") or row.get("request"))
    return (*mail, company, project, request)


def _mail_number_from_row(row: Dict[str, Any]) -> str:
    """Return the persisted public mail number, with a deterministic legacy fallback."""
    existing = _text(row.get("mail_number") or row.get("邮件编号"))
    if existing:
        return existing
    mail_key = _text(row.get("_db_mail_key"))
    if mail_key:
        return f"MAIL-{mail_key.upper()}"
    identity = "|".join(_mail_identity_tuple(row))
    if not identity.strip("|"):
        return ""
    return f"MAIL-{hashlib.sha1(identity.encode('utf-8', errors='ignore')).hexdigest()[:24].upper()}"


def _mail_identity_aliases(mail: Dict[str, Any]) -> set[Tuple[str, ...]]:
    """返回跨阶段一主表/过滤日志的邮件去重别名。

    不同 Excel 输出可能对同一标题使用空格、标点或全角字符的不同写法，
    不能只用原始 ``(发件人, 时间, 标题)`` 判断，否则同一封邮件会同时出现在
    询单队列和过滤队列。邮件编号优先；标题别名用于兼容旧输出，日期只在同一
    发件人和同一标题下参与匹配，避免把同一发件人的不同邮件误合并。
    """
    sender = _text(mail.get("sender") or mail.get("发件人邮箱")).casefold()
    raw_date = mail.get("date") or mail.get("发件日期")
    parsed = _mail_date_sort_value(raw_date)
    date_key = parsed.strftime("%Y-%m-%d %H:%M:%S") if parsed != datetime.min else _text(raw_date).replace("T", " ")
    day_key = date_key[:10]
    subject = _text(mail.get("subject") or mail.get("邮件主题") or mail.get("主题"))
    normalized_subject = re.sub(
        r"[\W_]+", "", unicodedata.normalize("NFKC", subject).casefold(), flags=re.UNICODE
    )
    aliases: set[Tuple[str, ...]] = set()
    number = _text(mail.get("mail_number") or mail.get("邮件编号")).casefold()
    if number:
        aliases.add(("number", number))
    if sender and date_key and normalized_subject:
        aliases.add(("subject-time", sender, date_key, normalized_subject))
        aliases.add(("subject-day", sender, day_key, normalized_subject))
    elif sender and date_key:
        # 标题缺失时只能退化为发件人+时间，避免空标题的记录重复出现。
        aliases.add(("sender-time", sender, date_key))
    if not aliases and _text(mail.get("id")):
        aliases.add(("id", _text(mail.get("id"))))
    return aliases


def _detail_number_from_row(row: Dict[str, Any]) -> str:
    """Return the persisted public detail number, with a deterministic fallback."""
    existing = _text(row.get("detail_number") or row.get("明细编号"))
    if existing:
        return existing
    key = _text(row.get("_db_record_key") or row.get("_id"))
    if not key:
        key = hashlib.sha1(
            "|".join(_workorder_detail_tuple(row)).encode("utf-8", errors="ignore")
        ).hexdigest()[:24]
    return f"DETAIL-{key.upper()}"


def _obvious_business_issues(row: Dict[str, Any]) -> List[dict]:
    """Return only high-confidence company-name anomalies for legacy rows.

    Stage-one workbooks have used both internal keys (客户/项目/需求) and
    exported headers (客户公司名称/标准化项目名称/需求).  Passing an old
    exported row directly to ``inspect_row`` makes valid data look empty and
    incorrectly changes its status to 待人工复核.  We still re-check known bad
    company patterns, but do not infer missing values from header differences.
    """
    try:
        from modules.business_validator import inspect_row

        normalized = {
            "客户": _text(row.get("客户") or row.get("客户公司名称")),
            "项目": _text(row.get("项目") or row.get("标准化项目名称")),
            "需求": _text(row.get("需求")),
        }
        if not normalized["客户"]:
            return []
        issues = inspect_row(normalized)
        return [
            issue for issue in issues
            if _text(issue.get("code")).startswith("COMPANY_")
        ]
    except Exception:
        return []


def _semantic_program_suggestions(row: Dict[str, Any]) -> List[str]:
    """Extract normalized program candidates from a legacy suggestion."""
    for key in ("语义建议值", "语义校验建议", "语义校验原因", "人工复核提示"):
        value = _text(row.get(key))
        if not value:
            continue
        match = re.search(r"program\s*:\s*([^;；]+)", value, re.I)
        if match:
            candidate = match.group(1).strip()
            # 旧 LLM 输出偶尔把“国家/标准项目”写在同一个 program 值中，
            # 例如“德国/德国WEEE”；两部分都作为候选参与比对。
            return [part.strip() for part in re.split(r"[/／]", candidate) if part.strip()]
    return []


def _hide_superseded_review_rows(rows: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """Hide clearly superseded legacy LLM rows without deleting their source.

    Some old runs wrote one valid primary row and several ``人工补全`` rows
    generated from the same mail.  Those review rows are marked
    ``PROJECT_COVERAGE_MISSING`` and explicitly recommend the already-valid
    primary project.  They are stale correction candidates, not additional
    businesses.  Keep them in Excel/SQLite for audit, but do not count them in
    the active workbench queue.
    """
    by_mail: Dict[Tuple[str, str, str], List[Dict[str, Any]]] = {}
    for row in rows:
        by_mail.setdefault(_mail_identity_tuple(row), []).append(row)

    result: List[Dict[str, Any]] = []
    for row in rows:
        if _text(row.get("_source")) != "人工补全":
            result.append(row)
            continue
        if _text(row.get("语义校验状态")).lower() != "invalid":
            result.append(row)
            continue
        issue_codes = _text(row.get("语义问题编号"))
        if "PROJECT_COVERAGE_MISSING" not in issue_codes:
            result.append(row)
            continue

        company = _text(row.get("客户公司名称") or row.get("客户") or row.get("company"))
        project = _text(row.get("标准化项目名称") or row.get("项目") or row.get("program"))
        suggestions = _semantic_program_suggestions(row)
        primary_rows = [
            candidate for candidate in by_mail.get(_mail_identity_tuple(row), [])
            if _text(candidate.get("_source")) == "待查名单"
            and _text(candidate.get("客户公司名称") or candidate.get("客户") or candidate.get("company")) == company
            and _text(candidate.get("语义校验状态")).lower() not in {"invalid", "uncertain"}
        ]
        primary_projects = {
            _text(candidate.get("标准化项目名称") or candidate.get("项目") or candidate.get("program"))
            for candidate in primary_rows
        }
        if any(suggestion in primary_projects for suggestion in suggestions):
            continue
        # If a legacy row has no parseable suggestion, do not guess: keep it for
        # manual review rather than silently dropping a potentially real project.
        result.append(row)
    return result


def _safe_path(path: Optional[str], fallback: Path) -> Path:
    # 兼容旧版 session/config 中残留的开发机绝对路径。显式传入但尚未
    # 生成的临时路径必须原样保留（测试/用户刚选择的空结果文件不能被
    # fallback 偷换成 output 里的历史文件）；只有路径为空时才使用默认值。
    if not path:
        return resolve_runtime_path(None, fallback)
    candidate = Path(str(path)).expanduser()
    if not candidate.is_absolute():
        candidate = APP_ROOT / candidate
    try:
        if candidate.is_file() or candidate.is_dir():
            return candidate.resolve()
    except OSError:
        pass
    if candidate.is_absolute() and candidate.name:
        for sibling in (
            APP_ROOT / "data" / candidate.name,
            APP_ROOT / "storage" / "imported" / candidate.name,
            APP_ROOT / "output" / candidate.name,
            APP_ROOT / candidate.name,
        ):
            try:
                if sibling.exists():
                    return sibling.resolve()
            except OSError:
                continue
    return candidate.resolve()


def _stable_id(source: str, row_number: int, row: Dict[str, Any]) -> str:
    raw = "|".join(
        [
            source,
            str(row_number),
            _text(row.get("发件人邮箱")),
            _text(row.get("发件日期")),
            _text(row.get("邮件主题")),
            _text(row.get("客户公司名称")),
            _text(row.get("标准化项目名称")),
        ]
    )
    return hashlib.sha1(raw.encode("utf-8", errors="ignore")).hexdigest()[:16]


def _read_sheet(path: Path, sheet_name: str, source: str) -> List[Dict[str, Any]]:
    if not path.exists():
        return []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
    except Exception:
        return []
    if not rows:
        return []
    headers = [_text(v) for v in rows[0]]
    result: List[Dict[str, Any]] = []
    for row_number, values in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") for v in values):
            continue
        item = {headers[i]: values[i] if i < len(values) else "" for i in range(len(headers)) if headers[i]}
        item["_source"] = source
        item["_row_number"] = row_number
        item["_id"] = _stable_id(source, row_number, item)
        result.append(item)
    return result


def _load_json(path: Path) -> Dict[str, Any]:
    try:
        with path.open("r", encoding="utf-8") as fh:
            value = json.load(fh)
        return value if isinstance(value, dict) else {}
    except (FileNotFoundError, json.JSONDecodeError, OSError):
        return {}


def _save_json(path: Path, value: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    with tmp.open("w", encoding="utf-8") as fh:
        json.dump(value, fh, ensure_ascii=False, indent=2)
    os.replace(tmp, path)


def _mail_date_sort_value(value: Any) -> datetime:
    """把邮件日期转为排序键；无法解析的旧数据排在最后。"""
    text = _text(value).replace("T", " ")
    if not text:
        return datetime.min
    for candidate in (text, text[:19], text[:10]):
        try:
            parsed = datetime.fromisoformat(candidate)
            # 邮件头有时带 +00:00，有时是本地时间；排序键统一为无时区
            # 的 UTC 值，避免历史/新数据混排时触发 naive/aware 比较异常。
            if parsed.tzinfo is not None:
                parsed = parsed.astimezone(timezone.utc).replace(tzinfo=None)
            return parsed
        except ValueError:
            continue
    return datetime.min


def _date_filter_key(value: Any) -> str:
    """返回用于邮件日期筛选的 YYYY-MM-DD；空值保持为空。"""
    text = _text(value).replace("T", " ")
    if not text:
        return ""
    candidate = text[:10]
    try:
        return datetime.fromisoformat(candidate).date().isoformat()
    except ValueError:
        raise ValueError(f"日期格式不正确：{value!r}，应为 YYYY-MM-DD")


def _mail_in_date_range(mail: Dict[str, Any], date_start: str = "", date_end: str = "") -> bool:
    """按邮件发件日期筛选；有日期范围时，无有效发件日期的邮件不进入导出。"""
    start = _date_filter_key(date_start)
    end = _date_filter_key(date_end)
    if start and end and start > end:
        start, end = end, start
    if not start and not end:
        return True
    date_key = _date_filter_key(mail.get("date"))
    if not date_key:
        return False
    return (not start or date_key >= start) and (not end or date_key <= end)


def _history_result(status: str) -> str:
    return {
        "confirmed": "处理完成",
        "partial": "部分确认",
        "filtered": "已过滤归档",
        "workorder_completed": "工单核对完成",
        "review": "待人工复核",
        "returned": "已退回复核",
        "needs_info": "待补资料",
        "ready": "待确认",
    }.get(_text(status), "待处理")


def _history_is_complete(status: str) -> bool:
    """人工完成、过滤归档及阶段二已出明确结论的邮件进入完成总表。"""
    return _text(status) in {"confirmed", "filtered", "workorder_completed"}


def _is_weee_detail(detail: Dict[str, Any]) -> bool:
    """判断工作台明细是否走德国 WEEE 的独立确认链路。"""
    if not isinstance(detail, dict):
        return False
    weee = detail.get("weee") if isinstance(detail.get("weee"), dict) else {}
    # 只信阶段一/工作台明确标记的专项，避免把普通项目名中的“WEEE”
    # 当成德国品类链路，破坏旧数据的普通业务确认语义。
    return bool(weee.get("enabled"))


def _detail_confirmation_ready(detail: Dict[str, Any]) -> bool:
    """返回该明细是否可以进入阶段二和“已完成”视图。

    普通业务由业务字段确认按钮决定；德国 WEEE 使用品牌/品类确认按钮，
    不要求操作人员再重复点击普通业务确认。两条链路在明细级汇合。
    """
    if not isinstance(detail, dict):
        return False
    if _is_weee_detail(detail):
        weee = detail.get("weee") if isinstance(detail.get("weee"), dict) else {}
        return bool(weee.get("confirmed") or _text(weee.get("status")) == "confirmed")
    if detail.get("battery", {}).get("enabled"):
        return bool(detail["battery"].get("confirmed"))
    return _text(detail.get("status")) == "confirmed"


def _event_signature(detail_id: str, event: Dict[str, Any]) -> str:
    return "|".join([
        _text(detail_id), _text(event.get("at")), _text(event.get("action")),
        _text(event.get("reason")),
    ])


# ---- 项目名称表：给「新增项目」提供可选项目，避免人工手输造成阶段二匹配不上 ----

_PROJECT_CACHE: Dict[str, Tuple[float, List[Dict[str, str]]]] = {}
_PROJECT_LOCK = threading.Lock()


def _config_project_names() -> Optional[str]:
    """从 config.yaml 取项目名称表路径（不引入 YAML 依赖，只扫这一行）。"""
    try:
        text = CONFIG_PATH.read_text(encoding="utf-8", errors="ignore")
    except OSError:
        return None
    match = re.search(r"^\s*project_names:\s*(\S+)\s*$", text, re.MULTILINE)
    return match.group(1).strip().strip("'\"") if match else None


def project_table_path() -> Optional[Path]:
    """项目名称表位置：会话状态（GUI 手选） > config.yaml > data/project_names.xlsx。"""
    session = _load_json(SESSION_STATE)
    for candidate in (session.get("project_table_path"), _config_project_names(), "data/project_names.xlsx"):
        if not candidate:
            continue
        path = _safe_path(str(candidate), Path("data") / "project_names.xlsx")
        try:
            if path.is_file():
                return path
        except OSError:
            continue
    return None


def load_project_names() -> List[Dict[str, str]]:
    """读取项目名称表；按文件 mtime 缓存，避免每次刷新都读盘。"""
    path = project_table_path()
    if not path:
        return []
    try:
        mtime = path.stat().st_mtime
    except OSError:
        return []
    key = str(path)
    with _PROJECT_LOCK:
        cached = _PROJECT_CACHE.get(key)
        if cached and cached[0] == mtime:
            return cached[1]
    result: List[Dict[str, str]] = []
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = [r for r in ws.iter_rows(values_only=True) if any(v not in (None, "") for v in r)]
        wb.close()
    except Exception:
        return []
    if not rows:
        return []
    headers = [_text(v) for v in rows[0]]
    if "项目名称" in headers:
        i_name: Optional[int] = headers.index("项目名称")
    else:  # 没有表头时退化为第二列
        i_name = 1
    i_number = headers.index("项目编号") if "项目编号" in headers else None
    i_country = headers.index("国家") if "国家" in headers else None
    i_business = headers.index("业务类型") if "业务类型" in headers else None

    def cell(row: Iterable[Any], index: Optional[int]) -> str:
        row = list(row)
        return _text(row[index]) if index is not None and index < len(row) else ""

    for row in rows[1:]:
        name = cell(row, i_name)
        if not name:
            continue
        result.append({
            "name": name,
            "number": cell(row, i_number),
            "country": cell(row, i_country) or country_of_project(name),
            "business": cell(row, i_business),
        })
    with _PROJECT_LOCK:
        _PROJECT_CACHE[key] = (mtime, result)
    return result


class WorkbenchStore:
    """读取阶段一产物并叠加人工复核状态。"""

    def __init__(
        self,
        primary_path: Optional[str] = None,
        review_path: Optional[str] = None,
        filtered_path: Optional[str] = None,
        workorder_result_path: Optional[str] = None,
        state_path: Optional[str] = None,
        database_path: Optional[str] = None,
        test_mode: bool = False,
    ):
        primary_fallback = DEFAULT_PRIMARY if DEFAULT_PRIMARY.exists() else LEGACY_PRIMARY
        review_fallback = DEFAULT_REVIEW if DEFAULT_REVIEW.exists() else LEGACY_REVIEW
        filtered_fallback = DEFAULT_FILTERED if DEFAULT_FILTERED.exists() else LEGACY_FILTERED
        workorder_fallback = (
            DEFAULT_WORKORDER_RESULT
            if DEFAULT_WORKORDER_RESULT.exists()
            else LEGACY_WORKORDER_RESULT
        )
        self.primary_path = _safe_path(primary_path, primary_fallback)
        self.review_path = _safe_path(review_path, review_fallback)
        self.filtered_path = _safe_path(filtered_path, filtered_fallback)
        self.workorder_result_path = _safe_path(workorder_result_path, workorder_fallback)
        self.test_mode = bool(test_mode)
        if state_path:
            self.state_path = _safe_path(state_path, REVIEW_STATE)
        elif self.test_mode:
            # 测试会话不加载旧的“退回/确认/修改”等人工状态，也不覆盖正式状态。
            # 同一服务存活期间仍可刷新页面，继续当前测试。
            session_dir = APP_ROOT / "storage" / "workbench_test_sessions"
            token = f"{datetime.now():%Y%m%d_%H%M%S}_{uuid.uuid4().hex[:8]}"
            self.state_path = session_dir / f"workbench_test_{token}.json"
        else:
            self.state_path = REVIEW_STATE
        # 自动历史只属于正式工作台。单元测试/临时状态文件不能污染真实邮件台账。
        self.history_enabled = (
            not self.test_mode
            and self.state_path.resolve() == REVIEW_STATE.resolve()
        )
        # 正式工作台才启用共享数据库；测试/临时会话继续完全隔离，避免把测试数据
        # 写进正式台账。数据库默认放在当前 APP_ROOT 下，也支持独立启动时传入路径。
        self.database: Optional[WorkbenchDatabase] = None
        if self.history_enabled:
            db_fallback = self.state_path.parent / "workbench.db"
            self.database = WorkbenchDatabase(_safe_path(database_path, db_fallback))
        self._lock = WORKBENCH_DATA_LOCK
        self._history_cache_fingerprint = None
        self._history_cache_summary = None
        self._history_sync_pending = False
        self._history_sync_thread = None

    def _history_input_fingerprint(self) -> Tuple[Tuple[str, int, int], ...]:
        """只用输入/状态文件的修改信息判断历史是否需要重新归档。"""
        paths = [
            self.primary_path, self.review_path, self.filtered_path,
            self.workorder_result_path, self.state_path, HISTORY_STATE,
        ]
        fingerprint = []
        for path in paths:
            try:
                stat_result = Path(path).stat()
                fingerprint.append((str(Path(path).resolve()), int(stat_result.st_mtime_ns), int(stat_result.st_size)))
            except OSError:
                fingerprint.append((str(Path(path).resolve()), 0, 0))
        return tuple(fingerprint)

    @staticmethod
    def _history_response_summary(result: Dict[str, Any], refreshing: bool = False) -> Dict[str, Any]:
        """工作台只展示历史计数；完整历史留在本地台账，不随每次快照传输。"""
        return {
            "enabled": bool(result.get("enabled")),
            "message": _text(result.get("message")),
            "completed_count": int(result.get("completed_count", 0) or 0),
            "unfinished_count": int(result.get("unfinished_count", 0) or 0),
            "completed_path": _text(result.get("completed_path")),
            "unfinished_path": _text(result.get("unfinished_path")),
            "refreshing": bool(refreshing),
        }

    def _schedule_history_sync(self, mails: List[Dict[str, Any]], filtered_mails: List[Dict[str, Any]],
                               expected_fingerprint: Tuple[Tuple[str, int, int], ...]) -> None:
        """将确认/修改后的历史 Excel 汇总移出 /api/state 请求关键路径。"""
        if self._history_sync_pending or not self.history_enabled:
            return
        self._history_sync_pending = True

        def worker() -> None:
            try:
                with self._lock:
                    # 若同步排队时源文件或人工状态又发生变化，放弃旧快照；下一次
                    # /api/state 会以新指纹重新排队，避免历史回退到陈旧状态。
                    if self._history_input_fingerprint() != expected_fingerprint:
                        return
                    result = self._sync_persistent_history(mails, filtered_mails)
                    self._history_cache_summary = self._history_response_summary(result)
                    self._history_cache_fingerprint = self._history_input_fingerprint()
            finally:
                with self._lock:
                    self._history_sync_pending = False

        self._history_sync_thread = threading.Thread(target=worker, name="workbench-history-sync", daemon=True)
        self._history_sync_thread.start()

    def wait_background_tasks(self, timeout: float = 10) -> None:
        """退出时等待已开始的历史写入，避免中途截断归档文件。"""
        thread = getattr(self, "_history_sync_thread", None)
        if thread and thread is not threading.current_thread():
            thread.join(timeout)
            if thread.is_alive():
                raise RuntimeError("历史记录仍在保存，请稍后再次关闭工作台。")

    def _raw_records(self) -> List[Dict[str, Any]]:
        records = _read_sheet(self.primary_path, "工单待查", "待查名单")
        # 当前选择的阶段一输出是主数据；复查表是另一条待补全队列。
        records.extend(_read_sheet(self.review_path, "漏单复查", "人工补全"))
        if self.database is not None:
            records = self.database.unrefreshed_rows(records)
        # 兼容旧版阶段一：附件证据中的“代理 | 编号 | 公司中文名”
        # 可修复城市冒充公司、代理为空等确定性错误，再写入持久数据库。
        records = [
            _prepare_workbench_row(_repair_attachment_fields(dict(row)))
            for row in records
        ]
        # 仅用于当前队列展示的兼容过滤：保留原始 rows 进入数据库，
        # 但隐藏同一封邮件中被语义校验明确判定为“应改回已有主记录”的
        # 旧人工补全行，避免一封邮件被显示成多个虚假的业务明细。
        display_records = _hide_superseded_review_rows(records)
        if self.database is not None:
            # 先把本次阶段一输出增量写入数据库。数据库保存完整历史，
            # 但当前队列不能把同一封邮件在历次解析中产生的旧版本一起展示，
            # 否则一次修复会变成“7 个主体/32 条明细”这种重复污染。
            self.database.ingest_rows(
                records,
                dataset="active",
                source_path=f"{self.primary_path};{self.review_path}",
            )
            stored = [
                _prepare_workbench_row(_repair_attachment_fields(dict(row)))
                for row in self.database.read_rows("active")
            ]
            # 数据库保留完整导入历史，但当前工作台展示仍必须隐藏同一封邮件
            # 中已被有效主记录覆盖的旧版“人工补全”行。接口增量导入时通常没有
            # 当前 Excel 文件，后面的 ``if records`` 分支不会执行；因此这里
            # 先对数据库读取结果统一做展示层过滤，避免旧的 7 条项目重新出现。
            stored_display = _hide_superseded_review_rows(stored)
            if records:
                fresh = list(display_records)
                fresh_by_mail: Dict[Tuple[str, str, str], List[Dict[str, Any]]] = {}
                stored_by_mail: Dict[Tuple[str, str, str], List[Dict[str, Any]]] = {}
                for row in fresh:
                    fresh_by_mail.setdefault(_mail_identity_tuple(row), []).append(row)
                for row in stored_display:
                    stored_by_mail.setdefault(_mail_identity_tuple(row), []).append(row)

                # 同一封邮件优先使用本次输出；如果本次输出全是旧版无法
                # 识别的空主体占位行，则回退到数据库中最近保存的合法主体。
                # 这让用户无需先清库就能看到修复后的真实两条项目。
                selected: List[Dict[str, Any]] = []
                fresh_mail_keys = set(fresh_by_mail)
                stable_by_legacy: Dict[str, str] = {}
                for row in stored_display:
                    legacy = _text(row.get("_legacy_id"))
                    stable = _text(row.get("_db_record_key"))
                    if legacy and stable:
                        stable_by_legacy.setdefault(legacy, stable)
                for mail_key in dict.fromkeys([*fresh_by_mail.keys(), *stored_by_mail.keys()]):
                    fresh_rows = fresh_by_mail.get(mail_key, [])
                    stored_rows = stored_by_mail.get(mail_key, [])
                    fresh_with_company = [
                        row for row in fresh_rows
                        if _text(row.get("客户公司名称") or row.get("客户") or row.get("company"))
                    ]
                    stored_with_company = [
                        row for row in stored_rows
                        if _text(row.get("客户公司名称") or row.get("客户") or row.get("company"))
                    ]
                    if fresh_with_company:
                        chosen = fresh_with_company
                    elif stored_with_company:
                        chosen = stored_with_company
                    else:
                        chosen = fresh_rows or stored_rows

                    # 将同一行映射回数据库稳定 ID，保证之前已保存的人工
                    # 修改仍能叠加到本次新输出；新行则保留本次原始 ID。
                    for row in chosen:
                        if row in fresh_rows:
                            legacy = _text(row.get("_id"))
                            stable = stable_by_legacy.get(legacy)
                            if stable:
                                row["_legacy_id"] = legacy
                                row["_id"] = stable
                    selected.extend(chosen)

                # 本次文件未覆盖的历史邮件仍保留在工作台，可按邮件日期
                # 继续处理；同一邮件只保留一个版本，不把当前邮件的旧版本
                # 再追加一次。
                if stored_display:
                    selected_keys = {
                        _text(row.get("_db_record_key")) or _text(row.get("_id"))
                        for row in selected
                    }
                    for row in stored_display:
                        if _mail_identity_tuple(row) in fresh_mail_keys:
                            continue
                        key = _text(row.get("_db_record_key")) or _text(row.get("_id"))
                        if key and key not in selected_keys:
                            selected.append(row)
                return selected
            if stored_display:
                return stored_display
        # 测试/临时会话没有数据库覆盖层，也必须使用与正式工作台相同的
        # 未识别公司清空与待复核规则，避免把发件方说明句继续显示成客户。
        return list(display_records)

    def _state(self) -> Dict[str, Any]:
        return _load_json(self.state_path)

    def _status(self, row: Dict[str, Any], saved: Dict[str, Any]) -> str:
        if saved.get("status") in {"review", "ready", "confirmed", "returned", "needs_info"}:
            return saved["status"]
        # 对旧的阶段一输出也重新执行一次确定性业务闸门，避免工作台继续把
        # “在其他欧盟国家或第三国设立的公司”等申请表说明句显示为可确认。
        if _obvious_business_issues(row):
            return "review"
        source = row.get("_source")
        confidence = _text(row.get("置信度")).lower()
        needs = _text(row.get("人工复核提示")) or _text(row.get("待人工补全"))
        semantic = _text(row.get("语义校验状态")).lower()
        attachment_count_check = _text(row.get("附件表格数量校验"))
        if source == "人工补全" or needs or attachment_count_check == "需人工确认" or semantic in {"uncertain", "invalid", "llm调用失败/未完成"} or confidence in {"low", "medium"} or not _text(row.get("代理")):
            return "review"
        return "ready"

    @staticmethod
    def _base_fields(row: Dict[str, Any]) -> Dict[str, str]:
        project = _text(row.get("标准化项目名称"))
        # 国家必须从项目名里解析出「国家」本身，不能按空格硬切
        # （"奥地利WEEE" 无空格，旧写法会让国家列显示成 "奥地利WEEE"）。
        country = country_of_project(project)
        return {
            "agent": _text(row.get("代理")) or _text(row.get("代理(可能空)")),
            "customer_code": _text(row.get("客户编号")),
            "company": _text(row.get("客户公司名称")),
            "country": country,
            "program": project,
            "request": _text(row.get("需求")),
        }

    def _record(self, row: Dict[str, Any], state: Dict[str, Any]) -> Dict[str, Any]:
        rid = row["_id"]
        saved_records = state.get("records", {}) if isinstance(state.get("records"), dict) else {}
        saved = saved_records.get(rid, {})
        # 数据库首次接管已有 Excel 时，使用旧行号生成的 ID 迁移历史人工状态，
        # 后续新操作则统一写入数据库稳定 ID。
        if not isinstance(saved, dict) or not saved:
            legacy_id = _text(row.get("_legacy_id"))
            saved = saved_records.get(legacy_id, {}) if legacy_id else {}
        if not isinstance(saved, dict):
            saved = {}
        detail_number = _detail_number_from_row(row)
        mail_number = _mail_number_from_row(row)
        fields = self._base_fields(row)
        fields = _merge_saved_fields(fields, saved)
        status = self._status(row, saved)
        business_issues = _obvious_business_issues(row)
        business_hint = "；".join(
            f"{item.get('field', '')}: {item.get('reason', '')}"
            for item in business_issues
            if item.get("reason")
        )
        evidence = [
            {"field": "customer_code", "label": "客户编号", "initial": self._base_fields(row)["customer_code"], "accepted": fields["customer_code"], "source": "主题/正文中的编号实体"},
            {"field": "company", "label": "客户公司", "initial": self._base_fields(row)["company"], "accepted": fields["company"], "source": _text(row.get("客户提取来源")) or "邮件正文"},
            {"field": "agent", "label": "代理", "initial": self._base_fields(row)["agent"], "accepted": fields["agent"], "source": _text(row.get("代理匹配方式")) or "待确认"},
            {"field": "program", "label": "服务项目", "initial": self._base_fields(row)["program"], "accepted": fields["program"], "source": _text(row.get("项目原始值")) or "正文/附件"},
            {"field": "request", "label": "需求", "initial": self._base_fields(row)["request"], "accepted": fields["request"], "source": "正文/主题"},
        ]
        attachment_expected = _text(row.get("附件表格记录数"))
        attachment_output = _text(row.get("附件表格输出数"))
        attachment_check = _text(row.get("附件表格数量校验"))
        attachment_source = _text(row.get("附件明细来源"))
        attachment_evidence = _json_list(row.get("附件证据"))
        attachment_files = _json_list(row.get("附件文件索引"))
        weee_items = _with_weee_item_ids(_json_list(row.get("德国WEEE品类明细")))
        weee_enabled = (
            _text(row.get("德国WEEE专项")) == "是"
            and bool(re.search(r"WEEE", fields["program"], re.I))
        )
        if weee_enabled and not weee_items:
            weee_items = _recover_weee_items_from_cached_attachments(row, weee_items)
        weee_check = _json_object(row.get("德国WEEE品类核对"))
        weee_llm_review = _json_object(row.get("德国WEEE字段复检"))
        # WEEE 品牌/品类确认是独立于普通邮件字段的增量状态。人工确认后
        # 优先使用状态 JSON 中保存的项目快照，不会因刷新阶段一 Excel 把
        # 已处理的品牌、原始品类或人工映射类别覆盖掉。
        # 只要状态中明确保存过 WEEE 项目快照，就以快照为准；空列表也必须
        # 生效，否则删除最后一个品牌/品类后，刷新会从阶段一原始数据把它恢复。
        saved_weee_items = saved.get("weee_items") if isinstance(saved.get("weee_items"), list) else None
        if saved_weee_items is not None and "weee_items" in saved:
            saved_items = _with_weee_item_ids(saved_weee_items)
            if not _should_refresh_legacy_weee_snapshot(
                saved,
                saved_items,
                weee_items,
                fields["company"],
            ):
                weee_items = saved_items
        saved_weee_status = _text(saved.get("weee_status"))
        if saved_weee_status == "confirmed":
            for item in weee_items:
                item.setdefault("weee_confirmed", True)
        weee_confirmed = bool(
            saved_weee_status == "confirmed"
            or (weee_items and all(bool(item.get("weee_confirmed")) for item in weee_items))
        )
        # 普通字段确认不能让尚未确认品类的 WEEE 明细掉出所有待办筛选。
        # 仅修正展示状态，保留原始人工操作和专项确认资格。
        if weee_enabled and not weee_confirmed and status == "confirmed":
            status = "ready"
        battery = battery_review(fields["program"], saved,
                                 _json_list(row.get("德国电池法品类明细")), status)
        if battery["enabled"] and not battery["confirmed"] and status == "confirmed":
            status = "ready"
        # 兼容旧阶段一结果：历史行可能因为申请表模板中的“WEEE产品信息”
        # 被标成专项开启，但当前业务项目实际是“德国包装法”。工作台不应
        # 在非 WEEE 业务卡上显示 WEEE 品类待核对；以标准化项目字段作排他闸门。
        if attachment_expected or attachment_output or attachment_check or attachment_source:
            evidence.append({
                "field": "attachment_record_count",
                "label": "附件表格数量",
                "initial": " / ".join(value for value in (attachment_expected, attachment_output) if value) or "（未记录）",
                "accepted": attachment_check or "（未校验）",
                "source": attachment_source or "附件结构化解析",
            })
        if weee_enabled:
            evidence.append({
                "field": "weee_category",
                "label": "德国 WEEE 品类",
                "initial": "；".join(
                    f"{_text(item.get('brand')) + ' / ' if _text(item.get('brand')) else ''}{_text(item.get('category'))}"
                    for item in weee_items if _text(item.get("category")) or _text(item.get("brand"))
                ) or "（未提取到明确品类）",
                "accepted": _text(row.get("德国WEEE品类状态")) or "待工单核对",
                "source": "产品分类表 + 邮件/附件证据",
            })
        return {
            "id": rid,
            "detail_number": detail_number,
            "mail_number": mail_number,
            "source": row.get("_source", ""),
            "row_number": row.get("_row_number"),
            "status": status,
            "confirmation_ready": bool(
                weee_confirmed if weee_enabled else battery["confirmed"] if battery["enabled"] else status == "confirmed"
            ),
            "sender": _text(row.get("发件人邮箱")),
            "recipient": _text(row.get("收件人")) or _text(row.get("收件人邮箱")),
            "date": _text(row.get("发件日期")),
            "operation_date": _text(row.get("_db_operation_date")) or _text(row.get("发件日期"))[:10],
            "subject": _text(row.get("邮件主题")),
            # 新版阶段一保留“邮件正文原文”；旧工作簿没有该列时兼容摘要。
            "body": _text(row.get("邮件正文原文")) or _text(row.get("邮件正文摘要(最多300字)")),
            "attachments": _text(row.get("附件名称")),
            "attachment_files": attachment_files,
            "fields": fields,
            "evidence": evidence,
            "confidence": _text(row.get("置信度")) or "unknown",
            "agent_match": _text(row.get("代理匹配方式")) or "待确认",
            "extraction": _text(row.get("智能提取方式")) or "规则提取",
            "review_hint": (
                _text(row.get("人工复核提示"))
                or _text(row.get("待人工补全"))
                or (f"确定性业务校验: {business_hint}" if business_hint else "")
            ),
            "attachment_audit": {
                "expected": attachment_expected,
                "output": attachment_output,
                "status": attachment_check,
                "source": attachment_source,
            },
            "attachment_evidence": attachment_evidence,
            "battery": battery,
            "weee": {
                "enabled": weee_enabled,
                "items": weee_items,
                "status": "confirmed" if weee_confirmed else (
                    "partial" if any(bool(item.get("weee_confirmed")) for item in weee_items)
                    else (_text(row.get("德国WEEE品类状态")) or "待工单核对")
                ),
                "confirmed": weee_confirmed,
                "confirmed_at": _text(saved.get("weee_confirmed_at")),
                "confirmation_reason": _text(saved.get("weee_confirmation_reason")),
                "comparison": weee_check,
                "reason": _text(row.get("德国WEEE专项说明")),
                "extraction_method": _text(row.get("德国WEEE字段提取方式")) or "规则提取",
                "llm_review": weee_llm_review,
            },
            "semantic": {
                "status": _text(row.get("语义校验状态")),
                "issue_codes": _text(row.get("语义问题编号")),
                "issues": _text(row.get("语义问题字段")),
                "evidence": _text(row.get("语义问题证据")),
                "current_values": _text(row.get("语义当前值")),
                "suggested_values": _text(row.get("语义建议值")),
                "suggestion": _text(row.get("语义校验建议")),
                "reason": _text(row.get("语义校验原因")),
                "model": _text(row.get("语义校验模型")),
            },
            "data_source": _text(row.get("数据来源")),
            "initial_row": {k: _text(v) for k, v in row.items() if not k.startswith("_")},
            "events": [
                {
                    **event,
                    "mail_number": _text(event.get("mail_number")) or mail_number,
                    "detail_number": _text(event.get("detail_number")) or detail_number,
                }
                for event in (saved.get("events", []) if isinstance(saved.get("events"), list) else [])
                if isinstance(event, dict)
            ],
        }

    @staticmethod
    def _deleted_map(state: Dict[str, Any]) -> Dict[str, Any]:
        deleted = state.get("deleted")
        return deleted if isinstance(deleted, dict) else {}

    @staticmethod
    def _added_map(state: Dict[str, Any]) -> Dict[str, Any]:
        added = state.get("added")
        return added if isinstance(added, dict) else {}

    def _added_record(self, item: Dict[str, Any], mail: Dict[str, Any], state: Dict[str, Any]) -> Dict[str, Any]:
        """把「人工新增的项目」包装成与阶段一明细同构的记录。

        人工新增的明细没有原始 Excel 行，因此 initial_row 由录入字段拼出，
        导出时能直接写回"工单待查"表头，阶段二仍可识别。
        """
        rid = _text(item.get("id"))
        saved = state.get("records", {}).get(rid, {}) if isinstance(state.get("records"), dict) else {}
        raw = item.get("fields") if isinstance(item.get("fields"), dict) else {}
        fields = {k: _text(raw.get(k)) for k in ("agent", "customer_code", "company", "country", "program", "request")}
        saved_fields = saved.get("fields") if isinstance(saved.get("fields"), dict) else {}
        fields.update({k: _text(v) for k, v in saved_fields.items() if k in fields})
        if not fields["country"]:
            fields["country"] = country_of_project(fields["program"])
        status = _text(saved.get("status")) or "ready"
        battery = battery_review(fields["program"], saved, [], status)
        if battery["enabled"] and not battery["confirmed"] and status == "confirmed":
            status = "ready"
        note = _text(item.get("note"))
        created = _text(item.get("created_at"))
        evidence = [
            {"field": "company", "label": "客户公司", "initial": fields["company"], "accepted": fields["company"], "source": "人工新增（原邮件人工判读）"},
            {"field": "agent", "label": "代理", "initial": fields["agent"], "accepted": fields["agent"], "source": "人工新增（原邮件人工判读）"},
            {"field": "program", "label": "服务项目", "initial": fields["program"], "accepted": fields["program"], "source": "人工新增（项目名称表）"},
            {"field": "request", "label": "需求", "initial": fields["request"], "accepted": fields["request"], "source": "人工新增（原邮件人工判读）"},
        ]
        initial_row = {
            "发件人邮箱": _text(mail.get("sender")),
            "发件日期": _text(mail.get("date")),
            "邮件主题": _text(mail.get("subject")),
            "邮件正文摘要(最多300字)": _text(mail.get("body")),
            "附件名称": _text(mail.get("attachments")),
            "代理": fields["agent"],
            "代理匹配方式": "人工指定",
            "客户编号": fields["customer_code"],
            "客户公司名称": fields["company"],
            "客户提取来源": "人工新增",
            "标准化项目名称": fields["program"],
            "项目原始值": fields["program"],
            "需求": fields["request"],
            "置信度": "人工录入",
            "智能提取方式": "人工新增",
            "数据来源": "人工新增",
            "人工复核提示": note,
        }
        events = [{"at": created, "action": "add_project", "label": "人工新增项目", "reason": note}]
        events.extend(saved.get("events") if isinstance(saved.get("events"), list) else [])
        detail_number = _detail_number_from_row({"_id": rid, "_db_record_key": rid})
        mail_number = _text(mail.get("mail_number")) or _mail_number_from_row({
            "sender": mail.get("sender"), "date": mail.get("date"), "subject": mail.get("subject"),
        })
        events = [
            {
                **event,
                "mail_number": _text(event.get("mail_number")) or mail_number,
                "detail_number": _text(event.get("detail_number")) or detail_number,
            }
            for event in events if isinstance(event, dict)
        ]
        return {
            "id": rid,
            "detail_number": detail_number,
            "mail_number": mail_number,
            "source": "人工新增",
            "row_number": None,
            "status": status,
            "confirmation_ready": battery["confirmed"] if battery["enabled"] else status == "confirmed",
            "battery": battery,
            "sender": _text(mail.get("sender")),
            "date": _text(mail.get("date")),
            "operation_date": created[:10] or _text(mail.get("date"))[:10],
            "subject": _text(mail.get("subject")),
            "body": _text(mail.get("body")),
            "attachments": _text(mail.get("attachments")),
            "fields": fields,
            "evidence": evidence,
            "confidence": "人工录入",
            "agent_match": "人工指定",
            "extraction": "人工新增",
            "review_hint": note,
            "semantic": {"status": "", "issues": "", "suggestion": "", "reason": "", "model": ""},
            "data_source": "人工新增",
            "initial_row": initial_row,
            "events": events,
            "created_at": created,
        }

    @staticmethod
    def _route_map(state: Dict[str, Any]) -> Dict[str, Any]:
        routes = state.get("mail_routes")
        return routes if isinstance(routes, dict) else {}

    @staticmethod
    def _history_mail_key(mail: Dict[str, Any]) -> str:
        """历史按邮件身份归档，不依赖阶段一 Excel 的行号。"""
        sender, date, subject = WorkbenchStore._mail_identity(
            mail.get("sender"), mail.get("date"), mail.get("subject")
        )
        raw = "|".join([sender, date, subject])
        return hashlib.sha1(raw.encode("utf-8", errors="ignore")).hexdigest()[:20]

    @staticmethod
    def _mail_identity(sender: Any, date: Any, subject: Any) -> Tuple[str, str, str]:
        """阶段一和阶段二按同一组邮件元数据关联，兼容 Excel 读出的日期对象。"""
        parsed = _mail_date_sort_value(date)
        date_key = parsed.strftime("%Y-%m-%d %H:%M:%S") if parsed != datetime.min else _text(date)
        return (_text(sender).lower(), date_key, _text(subject))

    def _merged_workorder_rows(self) -> List[Dict[str, Any]]:
        """读取阶段二结果并叠加工作台人工同步结果。"""
        rows = _read_sheet(self.workorder_result_path, "工单核对", "工单核对")
        if self.database is not None:
            self.database.ingest_rows(
                rows,
                dataset="workorder",
                source_path=str(self.workorder_result_path),
            )
            stored = self.database.read_rows("workorder")
            if stored:
                rows = stored
            # 工作台人工核查结果是阶段二结果的覆盖层。它不能直接写回原始
            # Excel（原始文件仍由阶段二输出），否则下一次刷新会把人工结论
            # 覆盖掉；按邮件+公司+项目+需求身份替换同一条结果。
            manual_rows = self.database.read_rows("workorder_manual")
            if manual_rows:
                manual_by_detail = {
                    _workorder_detail_tuple(row): row
                    for row in manual_rows
                    if isinstance(row, dict)
                }
                replaced = set()
                merged: List[Dict[str, Any]] = []
                for row in rows:
                    key = _workorder_detail_tuple(row)
                    override = manual_by_detail.get(key)
                    if override is not None:
                        # 自动阶段二若在人工同步之后再次完成查询，以较新的
                        # 查询时间为准；否则继续显示人工核查结果。
                        base_at = _mail_date_sort_value(row.get("查询时间戳"))
                        manual_at = _mail_date_sort_value(override.get("查询时间戳"))
                        merged.append(override if manual_at >= base_at else row)
                        replaced.add(key)
                    else:
                        merged.append(row)
                merged.extend(
                    row for key, row in manual_by_detail.items()
                    if key not in replaced
                )
                rows = merged
        return rows

    def _workorder_summaries(self) -> Dict[Tuple[str, str, str], Dict[str, Any]]:
        """读取阶段二结论，按邮件聚合“找到/未找到/待核对”。

        只有页面/RPA 已返回明确“是”或“漏单/否”时才算已经处理完成；查询失败、
        未查询和待确认不会被误归档为完成。
        """
        rows = self._merged_workorder_rows()
        grouped: Dict[Tuple[str, str, str], Dict[str, Any]] = {}
        for row in rows:
            key = self._mail_identity(
                row.get("发件人邮箱"), row.get("发件日期"),
                row.get("邮件主题") or row.get("主题"),
            )
            summary = grouped.setdefault(key, {
                "found": 0, "missing": 0, "pending": 0, "total": 0,
                "checked_at": "", "labels": [],
            })
            found = _text(row.get("是否已录单"))
            match_status = _text(row.get("匹配状态"))
            rpa_status = _text(row.get("RPA查询状态"))
            weee_status = _text(row.get("德国WEEE品类状态"))
            checked_at = _text(row.get("查询时间戳"))
            summary["total"] += 1
            if _text(row.get("德国WEEE专项")) == "是" and weee_status in {"未找到", "缺失"}:
                summary["missing"] += 1
                summary["labels"].append("WEEE品类未找到")
            elif _text(row.get("德国WEEE专项")) == "是" and weee_status in {"待人工核对", "待工单品类核对", "未比对"}:
                summary["pending"] += 1
                summary["labels"].append("WEEE品类待核对")
            elif found == "是":
                summary["found"] += 1
                summary["labels"].append("已找到")
            elif (
                "漏单" in match_status
                or (found == "否" and not any(token in rpa_status for token in ("失败", "未查询", "待")))
            ):
                summary["missing"] += 1
                summary["labels"].append("未找到（漏单）")
            else:
                summary["pending"] += 1
                summary["labels"].append("待核对")
            if _mail_date_sort_value(checked_at) >= _mail_date_sort_value(summary["checked_at"]):
                summary["checked_at"] = checked_at
        return grouped

    @staticmethod
    def _history_projection(mail: Dict[str, Any]) -> Dict[str, Any]:
        details = mail.get("details") if isinstance(mail.get("details"), list) else []
        companies = list(dict.fromkeys(
            _text((item.get("fields") or {}).get("company"))
            for item in details if isinstance(item, dict)
            and _text((item.get("fields") or {}).get("company"))
        ))
        projects = list(dict.fromkeys(
            _text((item.get("fields") or {}).get("program"))
            for item in details if isinstance(item, dict)
            and _text((item.get("fields") or {}).get("program"))
        ))
        detail_numbers = list(dict.fromkeys(
            _text(item.get("detail_number") or item.get("id"))
            for item in details if isinstance(item, dict)
            and _text(item.get("detail_number") or item.get("id"))
        ))
        status = _text(mail.get("status")) or "review"
        workorder = mail.get("workorder_summary") if isinstance(mail.get("workorder_summary"), dict) else {}
        if int(workorder.get("total") or 0) and not int(workorder.get("pending") or 0):
            status = "workorder_completed"
        workorder_result = ""
        if workorder:
            workorder_result = (
                f"已找到 {int(workorder.get('found') or 0)} 条；"
                f"未找到 {int(workorder.get('missing') or 0)} 条；"
                f"待核对 {int(workorder.get('pending') or 0)} 条"
            )
        return {
            "mail_number": _text(mail.get("mail_number") or mail.get("id")),
            "detail_numbers": "；".join(detail_numbers),
            "mail_date": _text(mail.get("date")),
            "sender": _text(mail.get("sender")),
            "recipient": _text(mail.get("recipient")),
            "subject": _text(mail.get("subject")),
            "attachments": _text(mail.get("attachments")),
            "source": _text(mail.get("source")),
            "status": status,
            "result": _history_result(status),
            "detail_count": len(details),
            "companies": "；".join(companies),
            "projects": "；".join(projects),
            "workorder_result": workorder_result,
            "workorder_found": int(workorder.get("found") or 0),
            "workorder_missing": int(workorder.get("missing") or 0),
            "workorder_pending": int(workorder.get("pending") or 0),
            "workorder_checked_at": _text(workorder.get("checked_at")),
        }

    @staticmethod
    def _history_events(mail: Dict[str, Any]) -> List[Dict[str, str]]:
        events: List[Dict[str, str]] = []
        for detail in mail.get("details") or []:
            if not isinstance(detail, dict):
                continue
            detail_id = _text(detail.get("id"))
            detail_number = _text(detail.get("detail_number")) or detail_id
            mail_number = _text(mail.get("mail_number")) or _text(mail.get("id"))
            for raw in detail.get("events") or []:
                if not isinstance(raw, dict):
                    continue
                events.append({
                    "signature": _event_signature(detail_id, raw),
                    "detail_id": detail_id,
                    "detail_number": detail_number,
                    "mail_number": mail_number,
                    "at": _text(raw.get("at")),
                    "action": _text(raw.get("action")),
                    "label": _text(raw.get("label")),
                    "reason": _text(raw.get("reason")),
                })
        for raw in mail.get("route_events") or []:
            if not isinstance(raw, dict):
                continue
            detail_id = f"mail:{_text(mail.get('id'))}"
            mail_number = _text(mail.get("mail_number")) or _text(mail.get("id"))
            events.append({
                "signature": _event_signature(detail_id, raw),
                "detail_id": detail_id,
                "detail_number": "",
                "mail_number": mail_number,
                "at": _text(raw.get("at")),
                "action": _text(raw.get("action")),
                "label": _text(raw.get("label")),
                "reason": _text(raw.get("reason")),
            })
        return sorted(events, key=lambda item: (item["at"], item["signature"]))

    def _write_history_total(self, path: Path, title: str, records: List[Dict[str, Any]]) -> None:
        """将持久历史按邮件日期导出为可直接打开的总表。"""
        path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        ws = wb.active
        ws.title = title
        headers = [
            "邮件编号", "业务明细编号", "邮件日期", "邮件标题", "发件人", "收件人", "业务明细数", "客户公司",
            "服务项目", "当前状态", "处理结果", "工单核对结果", "已找到数", "未找到数", "待核对数", "工单查询时间", "最后操作", "最后操作时间",
            "最近原因", "首次归档时间", "最后更新时间", "来源", "附件名称",
        ]
        ws.append(headers)
        header_fill = PatternFill("solid", fgColor="0F766E")
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
        for item in records:
            mail_date = _mail_date_sort_value(item.get("mail_date"))
            mail_value: Any = mail_date if mail_date != datetime.min else _text(item.get("mail_date"))
            ws.append([
                _text(item.get("mail_number")),
                _text(item.get("detail_numbers")),
                mail_value,
                _text(item.get("subject")),
                _text(item.get("sender")),
                _text(item.get("recipient")),
                int(item.get("detail_count") or 0),
                _text(item.get("companies")),
                _text(item.get("projects")),
                _text(item.get("status")),
                _text(item.get("result")),
                _text(item.get("workorder_result")),
                int(item.get("workorder_found") or 0),
                int(item.get("workorder_missing") or 0),
                int(item.get("workorder_pending") or 0),
                _text(item.get("workorder_checked_at")),
                _text(item.get("last_action")),
                _text(item.get("last_action_at")),
                _text(item.get("last_reason")),
                _text(item.get("first_seen_at")),
                _text(item.get("updated_at")),
                _text(item.get("source")),
                _text(item.get("attachments")),
            ])
        for row in ws.iter_rows(min_row=2, max_col=1):
            row[0].number_format = "yyyy-mm-dd hh:mm:ss"
        widths = [30, 62, 20, 46, 28, 28, 12, 34, 28, 18, 18, 34, 12, 12, 12, 20, 20, 20, 34, 20, 20, 16, 38]
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(path)

    def _sync_persistent_history(
        self, mails: List[Dict[str, Any]], filtered_mails: List[Dict[str, Any]]
    ) -> Dict[str, Any]:
        """同步邮件历史及完成/未完成总表；测试会话绝不写入正式留痕。"""
        if not self.history_enabled:
            return {
                "enabled": False,
                "message": "测试或临时会话：不读取或写入正式历史留痕",
                "completed": [],
                "unfinished": [],
                "completed_count": 0,
                "unfinished_count": 0,
                "completed_path": "",
                "unfinished_path": "",
            }

        state = _load_json(HISTORY_STATE)
        entries = state.setdefault("mails", {})
        if not isinstance(entries, dict):
            entries = {}
            state["mails"] = entries
        now = _now()

        normal_keys: set[Tuple[str, ...]] = set()
        for mail in mails:
            normal_keys.update(_mail_identity_aliases(mail))
        history_mails = list(mails)
        # 已在询单复核队列的过滤候选由其队列状态记录，避免同一封邮件双份入账。
        for filtered in filtered_mails:
            aliases = _mail_identity_aliases(filtered)
            if aliases & normal_keys:
                continue
            history_mails.append({
                "id": _text(filtered.get("id")),
                "mail_number": _text(filtered.get("mail_number")),
                "sender": _text(filtered.get("sender")),
                "date": _text(filtered.get("date")),
                "subject": _text(filtered.get("subject")),
                "attachments": _text(filtered.get("attachments")),
                "source": "过滤日志",
                "status": "filtered",
                "details": [],
            })

        for mail in history_mails:
            key = self._history_mail_key(mail)
            entry = entries.get(key) if isinstance(entries.get(key), dict) else {}
            projection = self._history_projection(mail)
            known_signatures = set(entry.get("event_signatures") or [])
            stored_events = entry.get("events") if isinstance(entry.get("events"), list) else []
            for event in self._history_events(mail):
                if event["signature"] not in known_signatures:
                    stored_events.append(event)
                    known_signatures.add(event["signature"])
            last_event = stored_events[-1] if stored_events else {}
            entry.update(projection)
            entry["id"] = key
            entry["first_seen_at"] = _text(entry.get("first_seen_at")) or now
            entry["updated_at"] = now
            entry["events"] = stored_events
            entry["event_signatures"] = list(known_signatures)
            entry["last_action"] = _text(last_event.get("label"))
            entry["last_action_at"] = _text(last_event.get("at"))
            entry["last_reason"] = _text(last_event.get("reason"))
            entries[key] = entry

        records = sorted(
            (item for item in entries.values() if isinstance(item, dict)),
            key=lambda item: (_mail_date_sort_value(item.get("mail_date")), _text(item.get("updated_at"))),
            reverse=True,
        )
        completed = [item for item in records if _history_is_complete(item.get("status", ""))]
        unfinished = [item for item in records if not _history_is_complete(item.get("status", ""))]
        state["updated_at"] = now
        _save_json(HISTORY_STATE, state)
        self._write_history_total(COMPLETED_HISTORY_XLSX, "处理完成", completed)
        self._write_history_total(UNFINISHED_HISTORY_XLSX, "未处理完成", unfinished)
        return {
            "enabled": True,
            "message": "正式历史已按邮件日期归档",
            "completed": completed,
            "unfinished": unfinished,
            "completed_count": len(completed),
            "unfinished_count": len(unfinished),
            "completed_path": str(COMPLETED_HISTORY_XLSX),
            "unfinished_path": str(UNFINISHED_HISTORY_XLSX),
        }

    def snapshot(self) -> Dict[str, Any]:
        with self._lock:
            state = self._state()
            deleted = self._deleted_map(state)
            added = self._added_map(state)
            routes = self._route_map(state)
            details = [self._record(row, state) for row in self._raw_records() if row["_id"] not in deleted]
            # 把阶段二刚完成的德国 WEEE 品类核对结果回填到对应业务明细，
            # 工作台刷新后即可看到“邮件品类 / 工单品类 / 逐项状态”。
            workorder_rows = self._merged_workorder_rows()
            workorder_by_detail = {
                _workorder_detail_tuple(row): row
                for row in workorder_rows
                if isinstance(row, dict)
            }
            for detail in details:
                weee = detail.get("weee") if isinstance(detail.get("weee"), dict) else {}
                if not weee.get("enabled"):
                    continue
                identity = _workorder_detail_tuple(detail.get("initial_row") or {})
                workorder_row = workorder_by_detail.get(identity)
                if workorder_row is None:
                    continue
                comparison = _json_object(workorder_row.get("德国WEEE品类核对"))
                if comparison:
                    weee["comparison"] = comparison
                if not weee.get("confirmed"):
                    weee["status"] = _text(workorder_row.get("德国WEEE品类状态")) or weee.get("status", "待工单核对")
                weee["reason"] = _text(workorder_row.get("德国WEEE专项说明")) or weee.get("reason", "")
                detail["weee"] = weee
            filtered = _read_sheet(self.filtered_path, "过滤日志", "过滤日志")
            if self.database is not None:
                filtered = self.database.unrefreshed_rows(filtered)
                self.database.ingest_rows(
                    filtered,
                    dataset="filtered",
                    source_path=str(self.filtered_path),
                )
                stored_filtered = self.database.read_rows("filtered")
                if stored_filtered:
                    filtered = stored_filtered
            # 过滤邮件单独提供给工作台的“已过滤邮件”视图。这里不沿用主队列
            # 的 settled 过滤规则：即使是已经确认的证书通知/非注册业务，仍要
            # 保留原始证据，方便人工追溯“为什么没有进入询单队列”。
            filtered_mails: List[Dict[str, Any]] = []
            for row in filtered:
                if row["_id"] in deleted:
                    continue
                filtered_id = f"filtered:{row['_id']}"
                route = routes.get(filtered_id) if isinstance(routes.get(filtered_id), dict) else {}
                # 已由人工转入询单复核的过滤邮件不能同时留在过滤页。
                if _text(route.get("route")) == "review":
                    continue
                reason = _text(row.get("过滤原因"))
                body = (
                    _text(row.get("正文原文"))
                    or _text(row.get("邮件正文原文"))
                    or _text(row.get("正文摘要(最多300字)"))
                    or _text(row.get("正文摘要"))
                    or _text(row.get("正文"))
                )
                filtered_mails.append({
                    "id": filtered_id,
                    "mail_number": _mail_number_from_row(row),
                    "sender": _text(row.get("发件人邮箱")),
                    "recipient": _text(row.get("收件人")) or _text(row.get("收件人邮箱")),
                    "date": _text(row.get("发件日期")),
                    "operation_date": _text(row.get("_db_operation_date")) or _text(row.get("处理时间戳"))[:10] or _text(row.get("发件日期"))[:10],
                    "subject": _text(row.get("主题")) or _text(row.get("邮件主题")),
                    "body": body,
                    "attachments": _text(row.get("附件名称")),
                    "reason": reason,
                    "intent_status": _text(row.get("意图LLM状态")),
                    "processed_at": _text(row.get("处理时间戳")),
                    "settled": _is_settled_filter(reason),
                    "source": "过滤日志",
                    "route_events": route.get("events") if isinstance(route.get("events"), list) else [],
                })
            # 过滤日志不再只作为右上角数量展示，而是进入同一复核队列；
            # 保留过滤原因，人工可以判断是否误过滤并补录字段。
            # 例外（只计数、不进队列）：证书/下号通知等已定论的过滤邮件，
            # 以及被人工明确删除的明细。
            for row in filtered:
                filtered_id = f"filtered:{row['_id']}"
                route = routes.get(filtered_id) if isinstance(routes.get(filtered_id), dict) else {}
                should_restore = _text(route.get("route")) == "review"
                if row["_id"] in deleted or (not should_restore and _is_settled_filter(row.get("过滤原因"))):
                    continue
                row = dict(row)
                row["发件人邮箱"] = row.get("发件人邮箱", "")
                row["收件人"] = row.get("收件人", row.get("收件人邮箱", ""))
                row["发件日期"] = row.get("发件日期", "")
                row["邮件主题"] = row.get("主题", "")
                row["邮件正文摘要(最多300字)"] = row.get("过滤原因", "")
                row["附件名称"] = row.get("附件名称", "")
                row["客户公司名称"] = ""
                row["标准化项目名称"] = ""
                row["需求"] = ""
                row["人工复核提示"] = row.get("过滤原因", "")
                row["数据来源"] = "过滤日志"
                details.append(self._record(row, state))
            groups: Dict[Tuple[str, str, str], Dict[str, Any]] = {}
            for item in details:
                key = (item["sender"], item["date"], item["subject"])
                if key not in groups:
                    groups[key] = {
                        "id": hashlib.sha1("|".join(key).encode("utf-8", errors="ignore")).hexdigest()[:16],
                        "mail_number": _text(item.get("mail_number")) or _mail_number_from_row(item),
                        "sender": item["sender"],
                        "recipient": item.get("recipient", ""),
                        "date": item["date"],
                        "operation_date": item.get("operation_date", "") or item["date"][:10],
                        "subject": item["subject"],
                        "body": item["body"],
                        "attachments": item["attachments"],
                        "details": [],
                        "source": item["source"],
                    }
                groups[key]["details"].append(item)
            # 人工新增的项目：可能挂到已有邮件下，也可能是阶段一完全没有产出的邮件
            known_ids = {g["id"] for g in groups.values()}
            for mail_id, bucket in added.items():
                if mail_id in known_ids or not isinstance(bucket, dict):
                    continue
                meta_raw = bucket.get("mail") if isinstance(bucket.get("mail"), dict) else {}
                meta = {k: _text(meta_raw.get(k)) for k in ("sender", "date", "subject", "body", "attachments")}
                key = (meta["sender"], meta["date"], meta["subject"])
                if key in groups:  # 同键邮件已在队列里，不要覆盖它原有明细
                    continue
                groups[key] = {
                    "id": mail_id,
                        "mail_number": _text(meta_raw.get("mail_number")) or _mail_number_from_row(meta_raw),
                        "sender": meta["sender"],
                        "recipient": _text(meta_raw.get("recipient")),
                        "date": meta["date"],
                        "operation_date": _text(meta_raw.get("operation_date")) or meta["date"][:10],
                        "subject": meta["subject"],
                    "body": meta["body"],
                    "attachments": meta["attachments"],
                    "details": [],
                    "source": "人工新增",
                }
            for mail_id, bucket in added.items():
                if not isinstance(bucket, dict):
                    continue
                target = next((g for g in groups.values() if g["id"] == mail_id), None)
                if target is None:
                    meta_raw = bucket.get("mail") if isinstance(bucket.get("mail"), dict) else {}
                    target = groups.get((_text(meta_raw.get("sender")), _text(meta_raw.get("date")), _text(meta_raw.get("subject"))))
                if target is None:
                    continue
                for raw_item in bucket.get("items") or []:
                    if not isinstance(raw_item, dict) or _text(raw_item.get("id")) in deleted:
                        continue
                    target["details"].append(self._added_record(raw_item, target, state))
            mails = [g for g in groups.values() if g["details"]]
            # 旧版本可能已经把 ECOPV 内部邮件写入主表或数据库（例如“人工补全”
            # 来源），它们不会再次经过阶段一过滤。这里补执行同一硬过滤，避免
            # 仅重启工作台后内部邮件仍出现在询单复核队列。
            internal_mails = [
                mail for mail in mails
                if _is_ecopv_internal_sender(mail.get("sender"))
            ]
            mails = [
                mail for mail in mails
                if not _is_ecopv_internal_sender(mail.get("sender"))
            ]
            # 统计口径以最终队列为准（含人工新增的明细，不含被删除的明细）
            all_details = [d for m in mails for d in m["details"]]
            # 邮件状态保留“部分确认”这一中间态；已确认明细在前端的
            # “已完成”视图单独展示，未确认明细仍留在待处理视图。
            for mail in mails:
                states = {d["status"] for d in mail["details"]}
                ready_details = [d for d in mail["details"] if _detail_confirmation_ready(d)]
                pending_details = [d for d in mail["details"] if not _detail_confirmation_ready(d)]
                mail["confirmed_detail_count"] = len(ready_details)
                mail["pending_detail_count"] = len(pending_details)
                if ready_details and pending_details:
                    mail["status"] = "partial"
                elif ready_details and not pending_details:
                    mail["status"] = "confirmed"
                elif "review" in states:
                    mail["status"] = "review"
                elif "returned" in states:
                    mail["status"] = "returned"
                elif "needs_info" in states:
                    mail["status"] = "needs_info"
                else:
                    mail["status"] = "ready"
                route = routes.get(mail["id"]) if isinstance(routes.get(mail["id"]), dict) else {}
                mail["route_events"] = route.get("events") if isinstance(route.get("events"), list) else []
                # 处理日期筛选必须绑定邮件实际发生日期，不能随着人工操作记录更新而漂移。
                # 人工操作时间仍保存在 route_events/details.events 中，供审计使用。
                mail["operation_date"] = _text(mail.get("date"))[:10]
                # 邮件级统计与“附件表格行数”分开，避免把主题里的数字、项目数
                # 或旧表单行误当成公司数量。工作台据此明确显示主体/项目/明细口径。
                mail["company_count"] = len({
                    _text((detail.get("fields") or {}).get("company"))
                    for detail in mail["details"]
                    if _text((detail.get("fields") or {}).get("company"))
                })
                mail["project_count"] = len({
                    _text((detail.get("fields") or {}).get("program"))
                    for detail in mail["details"]
                    if _text((detail.get("fields") or {}).get("program"))
                })
                mail["detail_count"] = len(mail["details"])

            # 人工将一封询单转为过滤后，原始阶段一行不删除，转而出现在过滤页；
            # 转回询单复核时只改路由状态，保留原始明细和完整操作记录。
            active_mails: List[Dict[str, Any]] = []
            manual_filtered: List[Dict[str, Any]] = []
            for mail in mails:
                route = routes.get(mail["id"]) if isinstance(routes.get(mail["id"]), dict) else {}
                if _text(route.get("route")) != "filtered":
                    active_mails.append(mail)
                    continue
                manual_filtered.append({
                    "id": mail["id"],
                    "mail_number": _text(mail.get("mail_number")),
                    "sender": mail["sender"],
                    "recipient": mail.get("recipient", ""),
                    "date": mail["date"],
                    "subject": mail["subject"],
                    "body": mail["body"],
                    "attachments": mail["attachments"],
                    "reason": _text(route.get("reason")) or "人工转为已过滤邮件",
                    "intent_status": "人工路由",
                    "processed_at": _text(route.get("at")),
                    "settled": False,
                    "source": "人工过滤",
                    "route_events": route.get("events") if isinstance(route.get("events"), list) else [],
                })
            # 将从旧主表/数据库发现的内部邮件转为过滤视图中的合成记录，
            # 保留主题、正文和附件证据，但不再回到询单队列。
            for mail in internal_mails:
                route = routes.get(mail["id"]) if isinstance(routes.get(mail["id"]), dict) else {}
                manual_filtered.append({
                    "id": f"internal:{mail['id']}",
                    "mail_number": _text(mail.get("mail_number")),
                    "sender": mail["sender"],
                    "recipient": mail.get("recipient", ""),
                    "date": mail["date"],
                    "subject": mail["subject"],
                    "body": mail["body"],
                    "attachments": mail["attachments"],
                    "reason": "ECOPV内部发件人邮箱",
                    "intent_status": "工作台统一过滤",
                    "processed_at": _now(),
                    "settled": True,
                    "source": "工作台统一过滤",
                    "route_events": route.get("events") if isinstance(route.get("events"), list) else [],
                })
            # 阶段一过滤日志是持久化累积表，旧版本可能在后续刷新中仍存在；
            # 同一封邮件如果已经出现在当前询单队列，就不能再在“已过滤邮件”中
            # 计数一次。这里按发件人+规范化日期+主题去重，避免出现“总邮件数
            # 275，但询单复核 121 + 已过滤 156 = 277”的重叠统计。
            current_mail_keys: set[Tuple[str, ...]] = set()
            for mail in active_mails:
                current_mail_keys.update(_mail_identity_aliases(mail))
            deduped_filtered: List[Dict[str, Any]] = []
            seen_filtered: set[Tuple[str, ...]] = set()
            overlap_removed = 0
            # 人工路由优先于阶段一旧过滤日志，保留人工操作人的原因和时间。
            manual_keys: set[Tuple[str, ...]] = set()
            for item in manual_filtered:
                aliases = _mail_identity_aliases(item)
                record_aliases = {alias for alias in aliases if alias[0] != "subject-day"}
                if aliases & current_mail_keys:
                    overlap_removed += 1
                    continue
                if record_aliases & seen_filtered:
                    continue
                if aliases:
                    seen_filtered.update(record_aliases)
                    manual_keys.update(aliases)
                else:
                    fallback = ("id", _text(item.get("id")))
                    if fallback in seen_filtered:
                        continue
                    seen_filtered.add(fallback)
                deduped_filtered.append(item)
            for item in filtered_mails:
                aliases = _mail_identity_aliases(item)
                record_aliases = {alias for alias in aliases if alias[0] != "subject-day"}
                if aliases & (current_mail_keys | manual_keys):
                    overlap_removed += 1
                    continue
                if record_aliases & seen_filtered:
                    continue
                if aliases:
                    seen_filtered.update(record_aliases)
                else:
                    fallback = ("id", _text(item.get("id")))
                    if fallback in seen_filtered:
                        continue
                    seen_filtered.add(fallback)
                deduped_filtered.append(item)
            filtered_mails = deduped_filtered
            mails = active_mails
            workorder_summaries = self._workorder_summaries()
            for mail in mails:
                mail["workorder_summary"] = workorder_summaries.get(
                    self._mail_identity(mail["sender"], mail["date"], mail["subject"]), {}
                )
            # 漏单单独形成可处理视图：它保留原邮件、附件和业务明细，
            # 但不从询单队列删除；处理完成后仍可在历史总表追溯。
            missing_mails = sorted(
                [
                    mail for mail in mails
                    if int((mail.get("workorder_summary") or {}).get("missing") or 0) > 0
                ],
                key=lambda mail: _mail_date_sort_value(mail.get("date")),
                reverse=True,
            )
            all_details = [detail for mail in mails for detail in mail["details"]]
            imap_summary = _read_imap_summary()
            counts = {
                "mails": len(mails),
                "details": len(all_details),
                "review": sum(1 for d in all_details if d["status"] == "review"),
                "ready": sum(1 for d in all_details if d["status"] == "ready"),
                "confirmed": sum(1 for d in all_details if _detail_confirmation_ready(d)),
                "returned": sum(1 for d in all_details if d["status"] == "returned"),
                "needs_info": sum(1 for d in all_details if d["status"] == "needs_info"),
                # 当前过滤页包含阶段一过滤日志及人工转入的邮件。
                "filtered": len(filtered_mails),
                "overlap_removed": overlap_removed,
                # 主询单队列与已过滤邮件已经在上面按邮件身份去重；
                # 因此这个总数是当前快照的互斥邮件总数，不把同一封邮件重复计算。
                "total_mails": len(mails) + len(filtered_mails),
                # 保留旧字段，兼容已有前端、缓存和导出调用。
                "unique_mails": len(mails) + len(filtered_mails),
                # 最近一次阶段一向 IMAP SEARCH 请求得到的原始数量，
                # 与当前工作台的互斥台账总数严格区分。
                "imap_read_total": imap_summary.get("imap_read_total", 0),
                "imap_parsed_total": imap_summary.get("parsed_total", 0),
            }
            current_history_fingerprint = self._history_input_fingerprint()
            if self._history_cache_summary is None:
                # 首次快照建立基线；后续确认和编辑只返回已缓存计数，历史文件
                # 在后台更新，不再让每个按钮都等待重建两本全量 Excel。
                history_result = self._sync_persistent_history(mails, filtered_mails)
                self._history_cache_summary = self._history_response_summary(history_result)
                self._history_cache_fingerprint = self._history_input_fingerprint()
                history = dict(self._history_cache_summary)
            elif current_history_fingerprint != self._history_cache_fingerprint:
                self._schedule_history_sync(mails, filtered_mails, current_history_fingerprint)
                history = dict(self._history_cache_summary)
                history["refreshing"] = self._history_sync_pending
            else:
                history = dict(self._history_cache_summary)
            table = project_table_path()
            return {
                "ok": True,
                "generated_at": _now(),
                "paths": {
                    "primary": str(self.primary_path),
                    "review": str(self.review_path),
                    "filtered": str(self.filtered_path),
                    "workorder_result": str(self.workorder_result_path),
                    "missing_workorder": str(
                        DEFAULT_MISSING_WORKORDER
                        if DEFAULT_MISSING_WORKORDER.exists()
                        else (LEGACY_MISSING_WORKORDER if LEGACY_MISSING_WORKORDER.exists() else DEFAULT_MISSING_WORKORDER)
                    ),
                    "database": str(self.database.path) if self.database is not None else "",
                    "project_table": str(table) if table else "",
                },
                "session": {
                    "mode": "test" if self.test_mode else "persistent",
                    "history_loaded": self.history_enabled,
                },
                "counts": counts,
                "imap_summary": imap_summary,
                "mails": mails,
                "missing_mails": missing_mails,
                "projects": load_project_names(),
                "filtered": [{k: _text(v) for k, v in row.items() if not k.startswith("_")} for row in filtered],
                "filtered_mails": filtered_mails,
                "history": history,
                "database": self.database_info(),
            }

    def _locate(self, state: Dict[str, Any], rid: str) -> Dict[str, str]:
        """按明细编号找出该条记录，用于删除时写审计摘要。"""
        for bucket in self._added_map(state).values():
            if not isinstance(bucket, dict):
                continue
            for item in bucket.get("items") or []:
                if isinstance(item, dict) and _text(item.get("id")) == rid:
                    fields = item.get("fields") if isinstance(item.get("fields"), dict) else {}
                    meta = bucket.get("mail") if isinstance(bucket.get("mail"), dict) else {}
                    return {
                        "detail_number": _text(item.get("detail_number")) or _detail_number_from_row({"_id": rid}),
                        "company": _text(fields.get("company")),
                        "program": _text(fields.get("program")),
                        "subject": _text(meta.get("subject")),
                        "source": "人工新增",
                    }
        for row in self._raw_records():
            if row["_id"] == rid:
                rec = self._record(row, state)
                return {
                    "detail_number": _text(rec.get("detail_number")) or _detail_number_from_row({"_id": rid}),
                    "company": rec["fields"].get("company", ""),
                    "program": rec["fields"].get("program", ""),
                    "subject": rec["subject"],
                    "source": rec["source"],
                }
        return {}

    def _effective_fields(self, state: Dict[str, Any], rid: str) -> Dict[str, str]:
        """返回操作前的当前字段值，用于生成可读的修改明细。"""
        records = state.get("records") if isinstance(state.get("records"), dict) else {}
        saved = records.get(rid) if isinstance(records, dict) else {}
        saved_fields = saved.get("fields") if isinstance(saved, dict) and isinstance(saved.get("fields"), dict) else {}
        for row in self._raw_records():
            if _text(row.get("_id")) == rid:
                base = dict(self._base_fields(row))
                return _merge_saved_fields(base, saved)
        for bucket in self._added_map(state).values():
            if not isinstance(bucket, dict):
                continue
            for item in bucket.get("items") or []:
                if isinstance(item, dict) and _text(item.get("id")) == rid:
                    raw = item.get("fields") if isinstance(item.get("fields"), dict) else {}
                    base = {key: _text(value) for key, value in raw.items()}
                    base.update({key: _text(value) for key, value in saved_fields.items()})
                    return base
        return {key: _text(value) for key, value in saved_fields.items()}

    @staticmethod
    def _changed_fields(before: Dict[str, Any], after: Dict[str, Any]) -> Dict[str, Dict[str, str]]:
        changed: Dict[str, Dict[str, str]] = {}
        for key in ("agent", "company", "country", "program", "request", "customer_code"):
            old = _text(before.get(key))
            new = _text(after.get(key))
            if old != new:
                changed[key] = {"before": old, "after": new}
        return changed

    @staticmethod
    def _append_change(state: Dict[str, Any], event: Dict[str, Any]) -> None:
        """把修改/新增作为独立增量留痕保存，便于导出而不依赖当前卡片是否仍可见。"""
        changes = state.setdefault("change_log", [])
        if not isinstance(changes, list):
            changes = []
            state["change_log"] = changes
        changes.append(event)
        # 防止长期运行的工作台状态文件无限增长；完整操作记录仍在明细 events/SQLite。
        if len(changes) > 5000:
            del changes[:-5000]

    @staticmethod
    def _change_log(state: Dict[str, Any]) -> List[Dict[str, Any]]:
        value = state.get("change_log")
        return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []

    def _add_project(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        mail_id = _text(payload.get("mail_id"))
        if not mail_id:
            raise ValueError("缺少所属邮件，无法新增项目")
        raw = payload.get("fields") if isinstance(payload.get("fields"), dict) else {}
        fields = {k: _text(raw.get(k)) for k in ("agent", "company", "country", "program", "request")}
        if not fields["company"] and not fields["program"]:
            raise ValueError("请至少填写客户公司或服务项目")
        if not fields["country"]:
            fields["country"] = country_of_project(fields["program"])
        meta_raw = payload.get("mail") if isinstance(payload.get("mail"), dict) else {}
        mail = {
            k: _text(meta_raw.get(k))
            for k in ("sender", "recipient", "date", "subject", "body", "attachments")
        }
        added = state.setdefault("added", {})
        bucket = added.setdefault(mail_id, {"mail": {}, "items": []})
        # 邮件元信息用于「整封邮件只剩人工新增明细」时仍能还原队列分组
        if any(mail.values()):
            bucket["mail"] = mail
        items = bucket.setdefault("items", [])
        for existing in items:
            old = existing.get("fields") if isinstance(existing.get("fields"), dict) else {}
            if (
                _text(old.get("company")) == fields["company"]
                and _text(old.get("program")) == fields["program"]
                and _text(old.get("country")) == fields["country"]
            ):
                raise ValueError("该邮件下已存在相同的公司+项目明细，未重复新增")
        created_at = _now()
        item_id = "add_" + hashlib.sha1(f"{mail_id}|{created_at}|{len(items)}".encode("utf-8")).hexdigest()[:12]
        detail_number = _detail_number_from_row({"_id": item_id, "_db_record_key": item_id})
        mail_number = _text(payload.get("mail_number")) or _mail_number_from_row(mail)
        note = _text(payload.get("reason"))
        items.append({"id": item_id, "fields": fields, "created_at": created_at, "note": note})
        self._append_change(state, {
            "kind": "added_detail",
            "at": created_at,
            "mail_id": mail_id,
            "mail_number": mail_number,
            "detail_number": detail_number,
            "record_id": item_id,
            "subject": _text(mail.get("subject")),
            "fields": dict(fields),
            "reason": note,
        })
        _save_json(self.state_path, state)
        warning = ""
        known = {p["name"] for p in load_project_names()}
        if known and fields["program"] and fields["program"] not in known:
            warning = f"服务项目「{fields['program']}」不在项目名称表中，阶段二可能匹配不到"
        return {
            "ok": True,
            "record_id": item_id,
            "mail_number": mail_number,
            "detail_number": detail_number,
            "change_type": "added_detail",
            "fields": dict(fields),
            "message": f"已新增项目：{fields['company'] or fields['program']}",
            "warning": warning,
        }

    def _delete_project(self, state: Dict[str, Any], rid: str) -> Dict[str, Any]:
        info = self._locate(state, rid)
        deleted = state.setdefault("deleted", {})
        deleted[rid] = {
            "at": _now(),
            "detail_number": info.get("detail_number", "") or _detail_number_from_row({"_id": rid}),
            "subject": info.get("subject", ""),
            "company": info.get("company", ""),
            "program": info.get("program", ""),
            "source": info.get("source", ""),
        }
        _save_json(self.state_path, state)
        label = info.get("company") or info.get("program") or info.get("subject") or rid
        return {"ok": True, "message": f"已删除项目：{label}"}

    def _delete_mail(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        """删除整封邮件下的当前业务明细，并保留整封邮件审计留痕。"""
        mail_id = _text(payload.get("mail_id"))
        if not mail_id:
            raise ValueError("缺少邮件，无法删除整封邮件")
        current = self.snapshot()
        # snapshot() 在正式模式会同步历史文件；重新读取一次，避免用进入
        # action 时的旧字典覆盖刚写入的历史状态。
        latest_state = self._state()
        state.clear()
        state.update(latest_state)
        mail = next(
            (item for item in current.get("mails", []) if _text(item.get("id")) == mail_id),
            None,
        )
        if not isinstance(mail, dict):
            raise ValueError("邮件不存在或已被删除")
        details = [item for item in mail.get("details") or [] if isinstance(item, dict)]
        if not details:
            raise ValueError("该邮件没有可删除的业务明细")
        deleted = state.setdefault("deleted", {})
        now = _now()
        for detail in details:
            rid = _text(detail.get("id"))
            if not rid:
                continue
            deleted[rid] = {
                "at": now,
                "detail_number": _text(detail.get("detail_number")) or _detail_number_from_row({"_id": rid}),
                "subject": _text(mail.get("subject")),
                "company": _text((detail.get("fields") or {}).get("company")),
                "program": _text((detail.get("fields") or {}).get("program")),
                "source": _text(detail.get("source")),
                "mail_id": mail_id,
                "mail_deleted": True,
            }
        deleted_mails = state.setdefault("deleted_mails", {})
        if not isinstance(deleted_mails, dict):
            deleted_mails = {}
            state["deleted_mails"] = deleted_mails
        deleted_mails[mail_id] = {
            "at": now,
            "mail_number": _text(mail.get("mail_number")),
            "subject": _text(mail.get("subject")),
            "sender": _text(mail.get("sender")),
            "date": _text(mail.get("date")),
            "detail_count": len(details),
            "reason": _text(payload.get("reason")) or "人工删除整封邮件",
        }
        self._append_change(state, {
            "kind": "deleted_mail",
            "at": now,
            "mail_id": mail_id,
            "mail_number": _text(mail.get("mail_number")),
            "subject": _text(mail.get("subject")),
            "detail_count": len(details),
            "reason": _text(payload.get("reason")) or "人工删除整封邮件",
        })
        _save_json(self.state_path, state)
        return {
            "ok": True,
            "message": f"已删除整封邮件：{_text(mail.get('subject')) or mail_id}",
            "mail_id": mail_id,
            "deleted_details": len(details),
        }

    def _delete_weee_item(self, state: Dict[str, Any], payload: Dict[str, Any], rid: str) -> Dict[str, Any]:
        """删除单个德国 WEEE 品牌/品类项目，并保留可追溯的操作记录。"""
        target_item_id = _text(payload.get("weee_item_id"))
        if not target_item_id:
            raise ValueError("缺少要删除的德国 WEEE 品牌/品类项目编号")

        records = state.setdefault("records", {})
        entry = records.setdefault(rid, {"fields": {}, "events": []})
        if not isinstance(entry, dict):
            entry = {"fields": {}, "events": []}
            records[rid] = entry
        events = entry.setdefault("events", [])
        if not isinstance(events, list):
            events = []
            entry["events"] = events

        # 前端把当前表单快照一并提交，保证尚未完成防抖保存的人工新增项目
        # 也能立即删除；没有提交快照时兼容从状态或阶段一原始数据恢复。
        raw_items = payload.get("weee_items")
        if isinstance(raw_items, list):
            current_items = _with_weee_item_ids(
                [item for item in raw_items if isinstance(item, dict)]
            )
        elif isinstance(entry.get("weee_items"), list):
            current_items = _with_weee_item_ids(entry.get("weee_items") or [])
        else:
            current_items = []
            for row in self._raw_records():
                if _text(row.get("_id")) == rid:
                    current_items = _with_weee_item_ids(
                        _json_list(row.get("德国WEEE品类明细"))
                    )
                    break

        removed = next(
            (item for item in current_items if _text(item.get("item_id")) == target_item_id),
            None,
        )
        if removed is None:
            raise ValueError("未找到要删除的德国 WEEE 品牌/品类项目，页面可能已刷新")
        remaining = [
            item for item in current_items
            if _text(item.get("item_id")) != target_item_id
        ]
        now = _now()
        entry["weee_items"] = remaining
        confirmed_items = [item for item in remaining if item.get("weee_confirmed")]
        entry["weee_status"] = (
            "confirmed" if remaining and len(confirmed_items) == len(remaining)
            else "partial" if confirmed_items
            else "pending"
        )
        deleted_ids = entry.setdefault("weee_deleted_item_ids", [])
        if not isinstance(deleted_ids, list):
            deleted_ids = []
            entry["weee_deleted_item_ids"] = deleted_ids
        if target_item_id not in deleted_ids:
            deleted_ids.append(target_item_id)
        reason = _text(payload.get("reason")) or "人工删除德国 WEEE 品牌/品类项目"
        events.append({
            "at": now,
            "action": "delete_weee_item",
            "label": "删除德国 WEEE 品牌与品类项目",
            "reason": reason,
            "mail_number": _text(payload.get("mail_number")),
            "detail_number": _text(payload.get("detail_number")) or _detail_number_from_row({"_id": rid}),
            "weee_item_id": target_item_id,
            "brand": _text(removed.get("brand")),
            "category": _text(removed.get("category") or removed.get("category_original")),
        })
        entry["updated_at"] = now
        _save_json(self.state_path, state)
        label = " / ".join(
            value for value in (
                _text(removed.get("brand")),
                _text(removed.get("category") or removed.get("category_original")),
            ) if value
        ) or target_item_id
        return {
            "ok": True,
            "message": f"已删除德国 WEEE 项目：{label}",
            "weee_status": entry["weee_status"],
            "item_count": len(remaining),
            "weee_item_id": target_item_id,
        }

    def _bulk_confirm(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        """Confirm selected mail groups without bypassing unresolved review states.

        The UI selects whole mail groups, while the persisted state is per detail
        row.  A mail is eligible only when every detail is already ``ready`` or
        ``confirmed``.  Review/returned/needs-info rows are reported back to the
        operator instead of being silently forced into ``confirmed``.
        """
        raw_ids = payload.get("mail_ids")
        if not isinstance(raw_ids, list):
            raise ValueError("批量确认需要邮件ID列表")
        mail_ids = list(dict.fromkeys(_text(value) for value in raw_ids if _text(value)))
        if not mail_ids:
            raise ValueError("请至少选择一封邮件")
        if len(mail_ids) > 200:
            raise ValueError("一次最多批量确认200封邮件")

        reason = _text(payload.get("reason")) or "批量人工确认：邮件字段已核对"
        current = {mail["id"]: mail for mail in self.snapshot().get("mails", [])}
        records = state.setdefault("records", {})
        confirmed_mails: List[str] = []
        confirmed_details = 0
        skipped: List[Dict[str, str]] = []
        changed = False
        for mail_id in mail_ids:
            mail = current.get(mail_id)
            if not mail:
                skipped.append({"mail_id": mail_id, "reason": "邮件不存在或已从当前队列移除"})
                continue
            details = mail.get("details") or []
            blocked = [detail for detail in details if detail.get("status") not in {"ready", "confirmed"}]
            if blocked:
                states = ", ".join(sorted({str(detail.get("status") or "未判定") for detail in blocked}))
                skipped.append({"mail_id": mail_id, "reason": f"含未解决明细：{states}"})
                continue
            pending = [detail for detail in details if detail.get("status") != "confirmed"]
            if not pending:
                skipped.append({"mail_id": mail_id, "reason": "该邮件已确认"})
                continue
            for detail in pending:
                rid = _text(detail.get("id"))
                if not rid:
                    continue
                entry = records.setdefault(rid, {"fields": {}, "events": []})
                if not isinstance(entry.get("events"), list):
                    entry["events"] = []
                entry["status"] = "confirmed"
                entry["events"].append({
                    "at": _now(),
                    "action": "bulk_confirm",
                    "label": "批量人工确认整理完成",
                    "reason": reason,
                    "mail_number": _text(mail.get("mail_number")),
                    "detail_number": _text(detail.get("detail_number")) or _text(detail.get("id")),
                })
                entry["updated_at"] = _now()
                confirmed_details += 1
                changed = True
            confirmed_mails.append(mail_id)

        if changed:
            _save_json(self.state_path, state)
        return {
            "ok": True,
            "message": f"已批量确认{len(confirmed_mails)}封邮件、{confirmed_details}条明细",
            "confirmed_mails": len(confirmed_mails),
            "confirmed_details": confirmed_details,
            "skipped": skipped,
            "skipped_count": len(skipped),
        }

    @staticmethod
    def _write_route_event(
        state: Dict[str, Any], mail_id: str, route: str, label: str, reason: str,
        mail_number: str = "",
    ) -> None:
        routes = state.setdefault("mail_routes", {})
        if not isinstance(routes, dict):
            routes = {}
            state["mail_routes"] = routes
        entry = routes.setdefault(mail_id, {"events": []})
        if not isinstance(entry, dict):
            entry = {"events": []}
            routes[mail_id] = entry
        events = entry.setdefault("events", [])
        if not isinstance(events, list):
            events = []
            entry["events"] = events
        event = {
            "at": _now(),
            "action": "move_to_filtered" if route == "filtered" else "move_to_review",
            "label": label,
            "reason": reason,
            "mail_number": _text(mail_number) or f"MAIL-{_text(mail_id).upper()}",
        }
        events.append(event)
        entry.update({"route": route, "at": event["at"], "reason": reason})

    def _move_mail_route(self, state: Dict[str, Any], payload: Dict[str, Any], route: str) -> Dict[str, Any]:
        mail_id = _text(payload.get("mail_id"))
        if not mail_id:
            raise ValueError("缺少邮件，无法变更队列")
        reason = _text(payload.get("reason"))
        if route == "filtered":
            label = "转为已过滤邮件"
            reason = reason or "人工确认：该邮件不进入询单复核"
        else:
            label = "转入询单复核"
            reason = reason or "人工恢复：需要进入询单复核"
        self._write_route_event(state, mail_id, route, label, reason, _text(payload.get("mail_number")))
        _save_json(self.state_path, state)
        return {"ok": True, "message": label}

    def _bulk_filter(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        raw_ids = payload.get("mail_ids")
        if not isinstance(raw_ids, list):
            raise ValueError("批量过滤需要邮件ID列表")
        mail_ids = list(dict.fromkeys(_text(value) for value in raw_ids if _text(value)))
        if not mail_ids:
            raise ValueError("请至少选择一封邮件")
        if len(mail_ids) > 200:
            raise ValueError("一次最多批量处理200封邮件")
        reason = _text(payload.get("reason")) or "批量人工确认：不进入询单复核"
        numbers = payload.get("mail_numbers") if isinstance(payload.get("mail_numbers"), dict) else {}
        for mail_id in mail_ids:
            self._write_route_event(state, mail_id, "filtered", "批量转为已过滤邮件", reason, _text(numbers.get(mail_id)))
        _save_json(self.state_path, state)
        return {"ok": True, "message": f"已将{len(mail_ids)}封邮件转入已过滤邮件", "count": len(mail_ids)}

    def _bulk_restore(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        """批量把已过滤邮件恢复到询单复核，并保留每封邮件的流转理由。"""
        raw_ids = payload.get("mail_ids")
        if not isinstance(raw_ids, list):
            raise ValueError("批量恢复需要邮件ID列表")
        mail_ids = list(dict.fromkeys(_text(value) for value in raw_ids if _text(value)))
        if not mail_ids:
            raise ValueError("请至少选择一封邮件")
        if len(mail_ids) > 200:
            raise ValueError("一次最多批量处理200封邮件")
        reason = _text(payload.get("reason")) or "批量人工确认：恢复询单复核"
        numbers = payload.get("mail_numbers") if isinstance(payload.get("mail_numbers"), dict) else {}
        for mail_id in mail_ids:
            self._write_route_event(state, mail_id, "review", "批量转入询单复核", reason, _text(numbers.get(mail_id)))
        _save_json(self.state_path, state)
        return {"ok": True, "message": f"已将{len(mail_ids)}封邮件转入询单复核", "count": len(mail_ids)}

    def _manual_workorder_result(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        """保存操作人员在工单系统核查后的单条结果。

        结果写入独立的 ``workorder_manual`` 覆盖层，不直接改写阶段二原始
        Excel；刷新工作台时按邮件+业务明细身份覆盖旧的“漏单/待核对”结论。
        """
        if self.database is None:
            raise ValueError("当前测试会话未启用持久数据库，不能同步工单结果")
        mail = payload.get("mail") if isinstance(payload.get("mail"), dict) else {}
        fields = payload.get("fields") if isinstance(payload.get("fields"), dict) else {}
        found = _text(payload.get("found"))
        if found not in {"是", "否", "待复核"}:
            raise ValueError("工单结果必须选择：是、否或待复核")
        sender = _text(mail.get("sender"))
        mail_date = _text(mail.get("date"))
        subject = _text(mail.get("subject"))
        company = _text(fields.get("company"))
        project = _text(fields.get("program"))
        if not sender or not mail_date or not subject or not company or not project:
            raise ValueError("缺少邮件、公司或项目身份，无法同步工单结果")
        now = _now()
        match_status = _text(payload.get("match_status")) or {
            "是": "人工确认-已找到",
            "否": "人工确认-漏单",
            "待复核": "人工确认-待复核",
        }[found]
        query_note = _text(payload.get("note"))
        result_row = {
            "发件人邮箱": sender,
            "收件人": _text(mail.get("recipient")),
            "发件日期": mail_date,
            "邮件主题": subject,
            "邮件正文摘要(最多300字)": _text(mail.get("body")),
            "附件名称": _text(mail.get("attachments")),
            "代理": _text(fields.get("agent")),
            "客户公司名称": company,
            "标准化项目名称": project,
            "需求": _text(fields.get("request")),
            "是否已录单": found,
            "工单编号": _text(payload.get("workorder_number")),
            "工单日期": _text(payload.get("workorder_date")),
            "下单日期": _text(payload.get("placed_date")),
            "匹配状态": match_status,
            "RPA查询状态": "人工查询同步",
            "查询方式": "人工同步",
            "模糊查询词": "",
            "查询时间戳": now,
            "人工核查备注": query_note,
        }
        self.database.ingest_rows(
            [result_row],
            dataset="workorder_manual",
            source_path="工作台人工工单同步",
        )
        record_id = _text(payload.get("record_id"))
        records = state.setdefault("records", {})
        if record_id:
            entry = records.setdefault(record_id, {"fields": {}, "events": []})
            if not isinstance(entry, dict):
                entry = {"fields": {}, "events": []}
                records[record_id] = entry
            events = entry.setdefault("events", [])
            if not isinstance(events, list):
                events = []
                entry["events"] = events
            events.append({
                "at": now,
                "action": "manual_workorder_result",
                "label": f"人工同步工单结果：{found}",
                "reason": query_note or match_status,
                "mail_number": _text(payload.get("mail_number")),
                "detail_number": _text(payload.get("detail_number")),
            })
            entry["workorder_result"] = result_row
            entry["updated_at"] = now
        _save_json(self.state_path, state)
        return {
            "ok": True,
            "message": f"已同步工单结果：{found}",
            "record_id": record_id,
            "mail_id": _text(payload.get("mail_id")),
            # 操作日志只保存结果摘要，不把正文/附件内容重复写入审计表。
            "result": {
                "found": found,
                "match_status": match_status,
                "workorder_number": _text(payload.get("workorder_number")),
                "query_time": now,
            },
        }

    def _queue_workorder_retry(self, state: Dict[str, Any], payload: Dict[str, Any]) -> Dict[str, Any]:
        """把漏单明细加入自动重查队列；下一次阶段二运行会强制跳过旧缓存。"""
        if not self.history_enabled:
            raise ValueError("当前测试会话不会写入正式自动重查队列")
        mail = payload.get("mail") if isinstance(payload.get("mail"), dict) else {}
        fields = payload.get("fields") if isinstance(payload.get("fields"), dict) else {}
        row = {
            "发件人邮箱": _text(mail.get("sender")),
            "发件日期": _text(mail.get("date")),
            "邮件主题": _text(mail.get("subject")),
            "代理": _text(fields.get("agent")),
            "客户公司名称": _text(fields.get("company")),
            "标准化项目名称": _text(fields.get("program")),
            "需求": _text(fields.get("request")),
        }
        key = "\x1f".join(_workorder_detail_tuple(row))
        if not all(_workorder_detail_tuple(row)):
            raise ValueError("缺少邮件、公司或项目身份，无法加入自动重查")
        now = _now()
        queue = _load_json(WORKORDER_RETRY_QUEUE)
        items = queue.get("items") if isinstance(queue.get("items"), list) else []
        items = [item for item in items if isinstance(item, dict) and _text(item.get("key")) != key]
        items.append({
            "key": key,
            "mail_id": _text(payload.get("mail_id")),
            "record_id": _text(payload.get("record_id")),
            "requested_at": now,
            "row": row,
            "status": "pending",
        })
        _save_json(WORKORDER_RETRY_QUEUE, {"updated_at": now, "items": items})
        record_id = _text(payload.get("record_id"))
        records = state.setdefault("records", {})
        if record_id:
            entry = records.setdefault(record_id, {"fields": {}, "events": []})
            if not isinstance(entry, dict):
                entry = {"fields": {}, "events": []}
                records[record_id] = entry
            events = entry.setdefault("events", [])
            if not isinstance(events, list):
                events = []
                entry["events"] = events
            events.append({
                "at": now,
                "action": "queue_workorder_retry",
                "label": "已加入工单自动重查队列",
                "reason": _text(payload.get("reason")) or "漏单重新进入工单核查",
                "mail_number": _text(payload.get("mail_number")),
                "detail_number": _text(payload.get("detail_number")),
            })
            entry["updated_at"] = now
        _save_json(self.state_path, state)
        return {
            "ok": True,
            "message": "已加入自动重查队列；下次运行阶段二时将强制重新查询",
            "queue_path": str(WORKORDER_RETRY_QUEUE),
            "requested_at": now,
        }

    def _finish_action(self, payload: Dict[str, Any], result: Dict[str, Any]) -> Dict[str, Any]:
        """兼容旧 JSON 状态的同时，把每次操作复制到 SQLite 留痕表。"""
        result = dict(result)
        mail_id = _text(payload.get("mail_id")) or _text(result.get("mail_id"))
        record_id = _text(payload.get("record_id")) or _text(result.get("record_id"))
        mail_number = _text(payload.get("mail_number")) or _text(result.get("mail_number"))
        detail_number = _text(payload.get("detail_number")) or _text(result.get("detail_number"))
        if not mail_number and mail_id:
            mail_number = f"MAIL-{mail_id.upper()}"
        if not detail_number and record_id:
            detail_number = f"DETAIL-{record_id.upper()}"
        audit = {
            "mail_number": mail_number,
            "detail_number": detail_number,
        }
        if isinstance(payload.get("mail_numbers"), dict):
            audit["mail_numbers"] = {
                _text(key): _text(value)
                for key, value in payload["mail_numbers"].items()
                if _text(key) and _text(value)
            }
        result["audit"] = audit
        if self.database is not None and result.get("ok") and _text(payload.get("action")) not in {"save_weee_draft", "save_battery_draft"}:
            try:
                self.database.record_operation(
                    _text(payload.get("action")),
                    record_id=record_id,
                    mail_id=mail_id,
                    reason=_text(payload.get("reason")),
                    result=result,
                    operation_date=_text(payload.get("operation_date")),
                )
            except Exception:
                # 数据库留痕不能阻断原有人工操作；错误仍可从服务日志/JSON 状态追溯。
                pass
        return result

    def ingest(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        """接收外部程序增量写入的数据，不要求启动 PyQt GUI。

        推荐传入 ``{"dataset":"active", "rows":[...]}``，rows 使用阶段一 Excel
        的表头。为便于其它脚本接入，也接受单条 ``row``。
        """
        if self.database is None:
            return {"ok": False, "error": "当前是测试/临时会话，未启用正式数据库"}
        dataset = _text(payload.get("dataset")) or "active"
        rows = payload.get("rows")
        if isinstance(rows, dict):
            rows = [rows]
        if not isinstance(rows, list):
            row = payload.get("row")
            rows = [row] if isinstance(row, dict) else []
        if not rows:
            raise ValueError("数据库导入需要 rows 或 row")
        result = self.database.ingest_rows(
            rows,
            dataset=dataset,
            source_path=_text(payload.get("source_path")) or "API导入",
        )
        return {"ok": True, **result, "message": f"已增量写入数据库：{result.get('rows', 0)} 条"}

    def operations(self, operation_date: str = "") -> Dict[str, Any]:
        if self.database is None:
            return {"ok": True, "enabled": False, "operations": []}
        return {
            "ok": True,
            "enabled": True,
            "operation_date": _text(operation_date),
            "operations": self.database.read_operations(operation_date),
        }

    def database_info(self) -> Dict[str, Any]:
        if self.database is None:
            return {"ok": True, "enabled": False, "path": "", "records": {}, "operations": 0}
        return {"ok": True, "enabled": True, **self.database.summary()}

    def action(self, payload: Dict[str, Any]) -> Dict[str, Any]:
        rid = _text(payload.get("record_id"))
        action = _text(payload.get("action"))
        request_id = _text(payload.get("_request_id") or payload.get("request_id"))
        allowed = {
            "edit", "confirm", "confirm_detail", "return", "needs_info", "agent_confirm", "reset",
            "confirm_weee", "confirm_weee_item", "save_weee_draft", "delete_weee_item",
            "confirm_battery_item", "save_battery_draft", "delete_battery_item",
            "add_project", "delete_project", "delete_mail", "bulk_confirm", "move_to_filtered",
            "move_to_review", "bulk_filter", "bulk_restore",
            "manual_workorder_result", "queue_workorder_retry",
        }
        mail_actions = {
            "add_project", "delete_mail", "bulk_confirm", "bulk_filter", "bulk_restore",
            "move_to_filtered", "move_to_review", "manual_workorder_result",
            "queue_workorder_retry",
        }
        if action not in allowed or (action not in mail_actions and not rid):
            raise ValueError("无效的工作台操作")
        with self._lock:
            state = self._state()
            request_results = state.get("request_results") if isinstance(state.get("request_results"), dict) else {}
            cached_result = request_results.get(request_id) if request_id else None
            if isinstance(cached_result, dict):
                return cached_result

            def finish(result: Dict[str, Any]) -> Dict[str, Any]:
                completed = self._finish_action(payload, result)
                if request_id and completed.get("ok"):
                    latest = self._state()
                    saved = latest.setdefault("request_results", {})
                    if not isinstance(saved, dict):
                        saved = {}
                        latest["request_results"] = saved
                    saved[request_id] = completed
                    if len(saved) > 500:
                        for old_key in list(saved)[:-500]:
                            saved.pop(old_key, None)
                    _save_json(self.state_path, latest)
                return completed

            if action == "add_project":
                return finish(self._add_project(state, payload))
            if action == "delete_project":
                return finish(self._delete_project(state, rid))
            if action == "delete_mail":
                return finish(self._delete_mail(state, payload))
            if action == "delete_weee_item":
                return finish(self._delete_weee_item(state, payload, rid))
            if action == "bulk_confirm":
                return finish(self._bulk_confirm(state, payload))
            if action == "bulk_filter":
                return finish(self._bulk_filter(state, payload))
            if action == "bulk_restore":
                return finish(self._bulk_restore(state, payload))
            if action == "manual_workorder_result":
                return finish(self._manual_workorder_result(state, payload))
            if action == "queue_workorder_retry":
                return finish(self._queue_workorder_retry(state, payload))
            if action == "move_to_filtered":
                return finish(self._move_mail_route(state, payload, "filtered"))
            if action == "move_to_review":
                return finish(self._move_mail_route(state, payload, "review"))
            records = state.setdefault("records", {})
            entry = records.setdefault(rid, {"fields": {}, "events": []})
            if not isinstance(entry.get("events"), list):
                entry["events"] = []
            reason = _text(payload.get("reason"))
            before_fields = self._effective_fields(state, rid)
            changed_fields: Dict[str, Dict[str, str]] = {}
            if action in {"save_battery_draft", "confirm_battery_item", "delete_battery_item"}:
                if not is_german_battery(before_fields.get("program")):
                    raise ValueError("电池品类只能保存到德国电池法业务")
                target_id = _text(payload.get("battery_item_id"))
                if action != "save_battery_draft" and not target_id:
                    raise ValueError("缺少电池品牌/品类项目编号")
                items = normalize_battery_items(
                    payload.get("battery_items"), entry.get("battery_items") or [],
                    confirm_id=target_id if action == "confirm_battery_item" else "",
                    delete_id=target_id if action == "delete_battery_item" else "")
                entry["battery_items"] = items
                complete = bool(items) and all(item["battery_confirmed"] for item in items)
                now = _now()
                entry["updated_at"] = now
                if action != "save_battery_draft":
                    entry["status"] = "confirmed" if complete else "ready"
                    entry["events"].append({"at": now, "action": action,
                        "label": "确认电池品牌/品类" if action == "confirm_battery_item" else "删除电池品牌/品类项",
                        "reason": reason, "battery_item_id": target_id,
                        "mail_number": _text(payload.get("mail_number")),
                        "detail_number": _text(payload.get("detail_number")) or rid})
                _save_json(self.state_path, state)
                return finish({"ok": True, "message": "电池品牌/品类已保存",
                    "battery": battery_review(before_fields["program"], entry, [], entry.get("status", "ready"))})
            if action == "save_weee_draft":
                normalized_items = _normalize_weee_draft_items(payload.get("weee_items"))
                now = _now()
                entry["weee_items"] = normalized_items
                confirmed_items = [item for item in normalized_items if item.get("weee_confirmed")]
                entry["weee_status"] = "confirmed" if len(confirmed_items) == len(normalized_items) else (
                    "partial" if confirmed_items else "pending"
                )
                # 草稿保存不改变普通业务明细状态，也不生成逐字输入的审计事件；
                # 最终点击“确认此品牌/品类”时仍会写入正式确认记录。
                entry["weee_draft_at"] = now
                entry["updated_at"] = now
                _save_json(self.state_path, state)
                return finish({
                    "ok": True,
                    "message": "德国 WEEE 品牌/品类草稿已保存",
                    "weee_status": entry["weee_status"],
                    "item_count": len(normalized_items),
                })
            if action == "edit":
                fields = payload.get("fields") or {}
                allowed = {"agent", "company", "country", "program", "request"}
                after_fields = dict(before_fields)
                after_fields.update({k: _text(v) for k, v in fields.items() if k in allowed})
                changed_fields = self._changed_fields(before_fields, after_fields)
                entry.setdefault("fields", {}).update({k: _text(v) for k, v in fields.items() if k in allowed})
                entry["status"] = "review"
                label = "人工修正字段"
            elif action in {"confirm", "confirm_detail"}:
                entry["status"] = "confirmed"
                label = "人工确认当前业务" if action == "confirm_detail" else "人工确认整理完成"
            elif action in {"confirm_weee", "confirm_weee_item"}:
                raw_items = payload.get("weee_items")
                if not isinstance(raw_items, list):
                    raise ValueError("德国 WEEE 品牌/品类数据格式不正确")
                category_names = {
                    "1": "热交换设备", "2": "屏幕和显示设备", "3": "灯具和光源",
                    "4": "大型设备", "5": "小型设备", "6": "小型信息和电信设备",
                }
                previous_items = {
                    _weee_item_id(item, index): item
                    for index, item in enumerate(entry.get("weee_items") or [])
                    if isinstance(item, dict)
                }
                target_item_id = _text(payload.get("weee_item_id"))
                if action == "confirm_weee_item" and not target_item_id:
                    raise ValueError("缺少要确认的德国 WEEE 项目编号")
                normalized_items: List[Dict[str, Any]] = []
                for index, raw_item in enumerate(raw_items):
                    if not isinstance(raw_item, dict):
                        continue
                    item = dict(raw_item)
                    item["item_id"] = _weee_item_id(item, index)
                    item["brand"] = _text(item.get("brand"))
                    item["category"] = _text(item.get("category") or item.get("category_original"))
                    item["category_original"] = item["category"]
                    category_class = _text(item.get("category_class"))
                    if category_class not in category_names:
                        category_class = ""
                    item["category_class"] = category_class
                    item["category_class_name"] = category_names.get(category_class, "")
                    # 选择了六类之一即表示操作人员已确认映射；未选择时保留
                    # 原有自动分类状态，方便后续继续人工核对而不是伪装成匹配。
                    if category_class:
                        item["category_class_status"] = "matched"
                    else:
                        item["category_class_status"] = _text(item.get("category_class_status")) or "unmatched"
                    previous = previous_items.get(item["item_id"]) or {}
                    was_confirmed = bool(item.get("weee_confirmed") or previous.get("weee_confirmed"))
                    complete = bool(item["brand"] and item["category"] and category_class)
                    if action == "confirm_weee":
                        item["weee_confirmed"] = complete
                    elif item["item_id"] == target_item_id:
                        item["weee_confirmed"] = complete
                    else:
                        item["weee_confirmed"] = was_confirmed
                    normalized_items.append(item)
                if not normalized_items:
                    raise ValueError("没有可保存的德国 WEEE 品牌/品类项目")
                now = _now()
                entry["weee_items"] = normalized_items
                all_confirmed = all(bool(item.get("weee_confirmed")) for item in normalized_items)
                entry["weee_status"] = "confirmed" if all_confirmed else (
                    "partial" if any(bool(item.get("weee_confirmed")) for item in normalized_items) else "pending"
                )
                entry["weee_confirmed_at"] = now
                entry["weee_confirmation_reason"] = reason or "人工确认德国 WEEE 品牌与品类"
                label = "确认德国 WEEE 单项品牌与品类" if action == "confirm_weee_item" else "确认德国 WEEE 品牌与品类"
                entry["events"].append({
                    "at": now,
                    "action": action,
                    "label": label,
                    "reason": reason or "人工确认德国 WEEE 品牌与品类",
                    "mail_number": _text(payload.get("mail_number")),
                    "detail_number": _text(payload.get("detail_number")) or _detail_number_from_row({"_id": rid}),
                    "weee_item_count": len(normalized_items),
                    "weee_item_id": target_item_id,
                    "weee_item_confirmed": next(
                        (bool(item.get("weee_confirmed")) for item in normalized_items if item.get("item_id") == target_item_id),
                        all_confirmed,
                    ),
                })
                entry["updated_at"] = now
                _save_json(self.state_path, state)
                return finish({
                    "ok": True,
                    "message": label,
                    "weee_status": entry["weee_status"],
                    "item_count": len(normalized_items),
                    "weee_item_id": target_item_id,
                    "weee_item_confirmed": next(
                        (bool(item.get("weee_confirmed")) for item in normalized_items if item.get("item_id") == target_item_id),
                        all_confirmed,
                    ),
                })
            elif action == "return":
                entry["status"] = "returned"
                label = "退回复核"
            elif action == "needs_info":
                entry["status"] = "needs_info"
                label = "标记待补资料"
            elif action == "agent_confirm":
                fields = payload.get("fields") or {}
                after_fields = dict(before_fields)
                if _text(fields.get("agent")):
                    entry.setdefault("fields", {})["agent"] = _text(fields["agent"])
                    after_fields["agent"] = _text(fields["agent"])
                changed_fields = self._changed_fields(before_fields, after_fields)
                entry["status"] = "ready" if _text(fields.get("agent")) else "review"
                label = "确认代理归属"
            else:
                records.pop(rid, None)
                _save_json(self.state_path, state)
                return finish({"ok": True, "message": "已撤销当前人工复核状态"})
            if changed_fields:
                self._append_change(state, {
                    "kind": "modified_detail",
                    "at": _now(),
                    "mail_id": _text(payload.get("mail_id")),
                    "mail_number": _text(payload.get("mail_number")),
                    "detail_number": _text(payload.get("detail_number")) or _detail_number_from_row({"_id": rid}),
                    "record_id": rid,
                    "subject": _text(payload.get("subject")),
                    "fields": changed_fields,
                    "reason": reason,
                })
            entry["events"].append({
                "at": _now(),
                "action": action,
                "label": label,
                "reason": reason,
                "mail_number": _text(payload.get("mail_number")),
                "detail_number": _text(payload.get("detail_number")),
                "change_type": "modified_detail" if changed_fields else "",
                "changed_fields": changed_fields,
            })
            entry["updated_at"] = _now()
            _save_json(self.state_path, state)
            result = {"ok": True, "message": label}
            if changed_fields:
                result.update({"change_type": "modified_detail", "changed_fields": changed_fields})
            return finish(result)

    def attachment_bytes(self, mail_id: str, name: str = "", token: str = "") -> Tuple[bytes, str, str, str]:
        """返回原始附件；原始文件缺失时返回由结构化证据重建的 xlsx 预览。

        返回值为 ``(内容, 文件名, MIME, 来源)``。token 只允许是缓存目录中的
        basename，不能让浏览器借此读取任意本机路径。
        """
        requested = _text(name)
        safe_token = os.path.basename(_text(token))
        cache_dir = (APP_ROOT / "cache" / "attachments").resolve()
        candidates: List[Path] = []
        if safe_token and safe_token not in {".", ".."}:
            candidate = (cache_dir / safe_token).resolve()
            try:
                candidate.relative_to(cache_dir)
            except ValueError:
                candidate = None
            if candidate is not None and candidate.is_file():
                candidates.append(candidate)

        snap = self.snapshot()
        mail = next((item for item in snap.get("mails", []) if _text(item.get("id")) == _text(mail_id)), None)
        if mail:
            wanted_key = Path(requested).name.casefold()
            for detail in mail.get("details") or []:
                for item in detail.get("attachment_files") or []:
                    item_name = _text(item.get("filename"))
                    item_token = os.path.basename(_text(item.get("token")))
                    if wanted_key and Path(item_name).name.casefold() != wanted_key:
                        continue
                    if item_token:
                        candidate = (cache_dir / item_token).resolve()
                        try:
                            candidate.relative_to(cache_dir)
                        except ValueError:
                            continue
                        if candidate.is_file():
                            candidates.append(candidate)
            unique = []
            seen = set()
            for candidate in candidates:
                if str(candidate) not in seen:
                    seen.add(str(candidate))
                    unique.append(candidate)
            candidates = unique
        if candidates:
            path = candidates[0]
            filename = requested or path.name.split("_", 1)[-1]
            mime = mimetypes.guess_type(filename)[0] or "application/octet-stream"
            return path.read_bytes(), filename, mime, "原始附件"

        # 旧工作簿没有保留附件本体，但可能保留了可恢复的 xlsx 行快照。
        evidence: List[dict] = []
        if mail:
            seen = set()
            for detail in mail.get("details") or []:
                for item in detail.get("attachment_evidence") or []:
                    item_name = _text(item.get("filename"))
                    if not item_name or (requested and Path(item_name).name.casefold() != Path(requested).name.casefold()):
                        continue
                    key = (item_name, json.dumps(item, ensure_ascii=False, sort_keys=True))
                    if key not in seen:
                        seen.add(key)
                        evidence.append(item)
        if not evidence:
            raise FileNotFoundError("原始附件未保留，当前数据也没有可重建的结构化预览")
        out_name = requested or _text(evidence[0].get("filename")) or "附件预览.xlsx"
        if not Path(out_name).suffix:
            out_name += ".xlsx"
        wb = Workbook()
        default = wb.active
        wb.remove(default)
        used_titles = set()
        for attachment in evidence:
            sheets = attachment.get("sheets") or []
            if not sheets and attachment.get("text"):
                sheets = [{"sheet_name": "文本预览", "rows": [{"row_number": 1, "cells": [attachment.get("text", "")] }]}]
            for sheet in sheets:
                title = _text(sheet.get("sheet_name")) or "工作表"
                title = re.sub(r"[\\/*?:\[\]]", "_", title)[:31] or "工作表"
                base = title
                suffix = 2
                while title in used_titles:
                    title = f"{base[:28]}_{suffix}"
                    suffix += 1
                used_titles.add(title)
                ws = wb.create_sheet(title)
                for row in sheet.get("rows") or []:
                    values = [row.get("row_number", "")]
                    values.extend(row.get("cells") or [])
                    ws.append(values)
                ws.freeze_panes = "A2"
                for column in ws.columns:
                    width = min(42, max(12, max((len(_text(cell.value)) for cell in column), default=0) + 2))
                    ws.column_dimensions[get_column_letter(column[0].column)].width = width
        if not wb.sheetnames:
            ws = wb.create_sheet("预览")
            ws.append(["说明"])
            ws.append(["当前仅能从邮件快照重建预览，原始附件文件未保存。"])
        buffer = io.BytesIO()
        wb.save(buffer)
        return buffer.getvalue(), out_name, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "结构化预览（非原始附件）"

    def _authorized_attachment_path(self, mail_id: str, token: str,
                                    name: str = "") -> Path:
        """只允许通过所选邮件自身的附件索引访问缓存原件。"""
        safe_token = os.path.basename(_text(token))
        if not safe_token or safe_token in {".", ".."}:
            raise FileNotFoundError("附件索引缺失或已过期")
        requested_name = Path(_text(name)).name.casefold()
        snap = self.snapshot()
        mail = next((item for item in snap.get("mails", [])
                     if _text(item.get("id")) == _text(mail_id)), None)
        if not mail:
            raise FileNotFoundError("找不到对应邮件")
        allowed = any(
            os.path.basename(_text(attachment.get("token"))) == safe_token
            and (not requested_name or Path(_text(attachment.get("filename"))).name.casefold() == requested_name)
            for detail in mail.get("details") or []
            for attachment in detail.get("attachment_files") or []
            if isinstance(attachment, dict)
        )
        if not allowed:
            raise FileNotFoundError("附件名称或索引与所选邮件不匹配")
        cache_dir = (APP_ROOT / "cache" / "attachments").resolve()
        path = (cache_dir / safe_token).resolve()
        try:
            path.relative_to(cache_dir)
        except ValueError as exc:
            raise FileNotFoundError("附件路径无效") from exc
        if not path.is_file():
            raise FileNotFoundError("缓存中的原始附件不存在")
        return path

    def attachment_preview(self, mail_id: str, token: str, name: str = "",
                           member: str = "", offset: int = 0,
                           page_size: int = 40) -> Dict[str, Any]:
        path = self._authorized_attachment_path(mail_id, token, name)
        filename = Path(_text(name)).name or path.name.split("_", 1)[-1]
        from utils.workbook_preview import preview_attachment
        return preview_attachment(path, filename, member, offset, page_size)

    def attachment_member_bytes(self, mail_id: str, token: str, name: str,
                                member: str) -> Tuple[bytes, str, str, str]:
        path = self._authorized_attachment_path(mail_id, token, name)
        filename = Path(_text(name)).name or path.name.split("_", 1)[-1]
        from utils.workbook_preview import download_archive_member
        payload, member_name = download_archive_member(path, filename, member)
        mime = mimetypes.guess_type(member_name)[0] or "application/octet-stream"
        return payload, member_name, mime, "原始附件"

    def export(self, date_start: str = "", date_end: str = "") -> Path:
        with self._lock:
            snap = self.snapshot()
            output_dir = MANUAL_OUTPUT
            output_dir.mkdir(parents=True, exist_ok=True)
            out = output_dir / f"workbench_reviewed_{datetime.now():%Y%m%d_%H%M%S}.xlsx"
            wb = Workbook()
            ws = wb.active
            # 主表沿用阶段一“工单待查”表头，因此可直接被阶段二候选列表识别。
            ws.title = "工单待查"
            source_rows = self._raw_records()
            original_headers: List[str] = []
            for row in source_rows:
                for key in row:
                    if not key.startswith("_") and key not in original_headers:
                        original_headers.append(key)
            required = ["邮件编号", "明细编号", "代理", "客户公司名称", "标准化项目名称", "需求"]
            # 旧版 initial_row 同时包含中文标准字段和内部字段（客户/项目/date 等）。
            # 阶段二会把标准字段映射成内部键名，因此内部同名字段必须改名，
            # 否则空的内部字段会覆盖标准字段，导致整批被判定为“客户为空”。
            from modules.excel_writer import HEADER_ALIASES

            used_headers = set(required)
            header_specs = [(header, header) for header in required]
            internal_aliases = set(HEADER_ALIASES.values())
            for source_header in original_headers:
                if source_header in required:
                    continue
                display_header = source_header
                if source_header in internal_aliases:
                    display_header = f"原始_{source_header}"
                if display_header in used_headers:
                    suffix = 2
                    candidate = f"{display_header}_{suffix}"
                    while candidate in used_headers:
                        suffix += 1
                        candidate = f"{display_header}_{suffix}"
                    display_header = candidate
                used_headers.add(display_header)
                header_specs.append((display_header, source_header))
            for header in ["德国电池法品类明细", "工作台状态", "工作台确认类型", "人工修改原因", "工作台最后操作时间"]:
                if header in used_headers:
                    continue
                header_specs.append((header, header))
                used_headers.add(header)
            headers = [display for display, _source in header_specs]
            ws.append(headers)
            selected_mails = [
                mail for mail in snap["mails"]
                if _mail_in_date_range(mail, date_start, date_end)
            ]
            export_detail_count = 0
            for mail in selected_mails:
                for detail in mail["details"]:
                    # 导出主表就是当前确认结果：普通业务按字段确认，
                    # 德国 WEEE 按品牌/品类确认。未确认明细不进入阶段二。
                    if not _detail_confirmation_ready(detail):
                        continue
                    events = detail.get("events") or []
                    reason = events[-1].get("reason", "") if events else ""
                    values = dict(detail.get("initial_row") or {})
                    weee = detail.get("weee") if isinstance(detail.get("weee"), dict) else {}
                    battery = detail.get("battery") or {}
                    if battery.get("enabled"):
                        values["德国电池法品类明细"] = json.dumps(battery.get("items") or [], ensure_ascii=False, separators=(",", ":"))
                    if weee.get("enabled"):
                        items = weee.get("items") if isinstance(weee.get("items"), list) else []
                        values["德国WEEE品类明细"] = json.dumps(items, ensure_ascii=False, separators=(",", ":"))
                        values["德国WEEE品类状态"] = "已确认"
                    values.update({
                        "邮件编号": mail.get("mail_number", ""),
                        "明细编号": detail.get("detail_number", detail.get("id", "")),
                        "代理": detail["fields"].get("agent", ""),
                        "客户公司名称": detail["fields"].get("company", ""),
                        "标准化项目名称": detail["fields"].get("program", ""),
                        "需求": detail["fields"].get("request", ""),
                        # 只有能进入阶段二的明细才写入主表，因此统一导出为
                        # confirmed；WEEE 不要求重复点击普通业务确认按钮。
                        "工作台状态": "confirmed",
                        "工作台确认类型": "德国WEEE品牌/品类" if weee.get("enabled") else "德国电池法品牌/品类" if battery.get("enabled") and battery.get("requires_review") else "普通业务",
                        "人工修改原因": reason,
                        "工作台最后操作时间": events[-1].get("at", "") if events else "",
                    })
                    ws.append([values.get(source, "") for _display, source in header_specs])
                    export_detail_count += 1
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for idx, header in enumerate(headers, start=1):
                ws.column_dimensions[chr(64 + idx) if idx <= 26 else "A"].width = min(42, max(14, len(header) + 4))
            log_ws = wb.create_sheet("人工操作记录")
            log_ws.append(["邮件编号", "明细编号", "邮件主题", "操作时间", "操作", "原因"])
            for mail in selected_mails:
                for detail in mail["details"]:
                    for event in detail.get("events") or []:
                        log_ws.append([mail.get("mail_number", ""), detail.get("detail_number", detail["id"]), mail["subject"], event.get("at", ""), event.get("label", ""), event.get("reason", "")])
            log_ws.freeze_panes = "A2"
            # 被人工删除的项目不写回主表，但要留痕，便于事后追溯/恢复
            deleted_rows = self._deleted_map(self._state())
            if deleted_rows:
                del_ws = wb.create_sheet("已删除明细")
                del_ws.append(["明细编号", "邮件主题", "客户公司", "服务项目", "来源", "删除时间"])
                for rid, info in deleted_rows.items():
                    info = info if isinstance(info, dict) else {}
                    del_ws.append([
                        info.get("detail_number") or _detail_number_from_row({"_id": rid}),
                        info.get("subject", ""),
                        info.get("company", ""),
                        info.get("program", ""),
                        info.get("source", ""),
                        info.get("at", ""),
                    ])
                del_ws.freeze_panes = "A2"
            # 变更单独拆表：主表保留当前最终值，下面两张表记录人工“改了什么”
            # 和“新增了什么”，避免操作人员只能从整行前后对比中猜差异。
            change_log = self._change_log(self._state())
            field_labels = {
                "agent": "代理", "company": "客户公司", "country": "国家",
                "program": "服务项目", "request": "具体业务", "customer_code": "客户编号",
            }
            modified_ws = wb.create_sheet("修改明细")
            modified_ws.append([
                "邮件编号", "明细编号", "修改时间", "邮件主题", "修改字段",
                "修改前", "修改后", "修改原因", "来源",
            ])
            added_ws = wb.create_sheet("新增明细")
            added_ws.append([
                "邮件编号", "明细编号", "新增时间", "邮件主题", "代理",
                "客户公司", "国家", "服务项目", "具体业务", "新增原因", "来源",
            ])
            for sheet in (modified_ws, added_ws):
                for cell in sheet[1]:
                    cell.font = Font(color="FFFFFF", bold=True)
                    cell.fill = PatternFill("solid", fgColor="0F766E")
                sheet.freeze_panes = "A2"
                sheet.auto_filter.ref = sheet.dimensions
            for change in change_log:
                kind = _text(change.get("kind"))
                mail_number = _text(change.get("mail_number"))
                detail_number = _text(change.get("detail_number"))
                at = _text(change.get("at"))
                subject = _text(change.get("subject"))
                reason = _text(change.get("reason"))
                if kind == "modified_detail":
                    fields = change.get("fields") if isinstance(change.get("fields"), dict) else {}
                    for key, values in fields.items():
                        values = values if isinstance(values, dict) else {}
                        modified_ws.append([
                            mail_number, detail_number, at, subject,
                            field_labels.get(key, key), _text(values.get("before")),
                            _text(values.get("after")), reason, "人工复核工作台",
                        ])
                elif kind == "added_detail":
                    fields = change.get("fields") if isinstance(change.get("fields"), dict) else {}
                    added_ws.append([
                        mail_number, detail_number, at, subject,
                        _text(fields.get("agent")), _text(fields.get("company")),
                        _text(fields.get("country")), _text(fields.get("program")),
                        _text(fields.get("request")), reason, "人工复核工作台",
                    ])
            for sheet in (modified_ws, added_ws):
                for column in sheet.columns:
                    values = [_text(cell.value) for cell in column]
                    width = min(50, max(14, max((len(value) for value in values), default=0) + 2))
                    sheet.column_dimensions[get_column_letter(column[0].column)].width = width
            wb.save(out)
            return out


class _Handler(BaseHTTPRequestHandler):
    server_version = "ECOPVWorkbench/1.0"

    def _store(self) -> WorkbenchStore:
        return self.server.workbench_store  # type: ignore[attr-defined]

    def _send(self, payload: Any, status: int = 200, content_type: str = "application/json; charset=utf-8") -> None:
        raw = payload if isinstance(payload, bytes) else json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-store")
        # 允许用户通过独立启动脚本打开 file:// 页面；仅接受浏览器的 null origin，
        # 不把本机数据服务开放给任意网站。
        if self.headers.get("Origin") == "null":
            self.send_header("Access-Control-Allow-Origin", "null")
        self.end_headers()
        try:
            self.wfile.write(raw)
        except (BrokenPipeError, ConnectionResetError, OSError):
            # 浏览器刷新/超时后主动断开连接不应让服务线程留下异常噪声。
            pass

    def _send_binary(
        self,
        payload: bytes,
        filename: str,
        content_type: str,
        source: str = "",
        download: bool = False,
    ) -> None:
        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(payload)))
        disposition = "attachment" if download else "inline"
        self.send_header("Content-Disposition", f"{disposition}; filename*=UTF-8''{quote(filename)}")
        self.send_header("Cache-Control", "no-store")
        if source:
            # HTTP 头只能使用 latin-1；来源说明在界面不需要靠该头展示，
            # 这里使用 ASCII 标识避免中文来源导致响应头编码异常。
            source_code = "original" if source == "原始附件" else "reconstructed-preview"
            self.send_header("X-ECOPV-Attachment-Source", source_code)
        if self.headers.get("Origin") == "null":
            self.send_header("Access-Control-Allow-Origin", "null")
        self.end_headers()
        try:
            self.wfile.write(payload)
        except (BrokenPipeError, ConnectionResetError, OSError):
            pass

    def do_OPTIONS(self) -> None:  # noqa: N802
        if self.headers.get("Origin") == "null":
            self.send_response(204)
            self.send_header("Access-Control-Allow-Origin", "null")
            self.send_header("Access-Control-Allow-Methods", "GET, POST, OPTIONS")
            self.send_header("Access-Control-Allow-Headers", "Content-Type")
            self.send_header("Content-Length", "0")
            self.end_headers()
            return
        self._send({"ok": False, "error": "不允许的跨域来源"}, 403)

    def do_GET(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        if parsed.path in {"/", "/index.html"}:
            try:
                self._send(HTML_PATH.read_bytes(), content_type="text/html; charset=utf-8")
            except OSError as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        if parsed.path == "/api/version":
            try:
                html = HTML_PATH.read_bytes()
                marker = re.search(rb'<meta name="ecopv-workbench-build" content="([^"]+)"', html)
                self._send({
                    "ok": True,
                    "build": WORKBENCH_BUILD,
                    "html_build": marker.group(1).decode("ascii", errors="replace") if marker else "missing",
                    "html_sha256": hashlib.sha256(html).hexdigest(),
                })
            except OSError as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        if parsed.path == "/api/state":
            try:
                self._send(self._store().snapshot())
            except Exception as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        if parsed.path == "/api/attachment-preview":
            query = parse_qs(parsed.query)
            try:
                payload = self._store().attachment_preview(
                    query.get("mail_id", [""])[0],
                    query.get("token", [""])[0],
                    query.get("name", [""])[0],
                    query.get("member", [""])[0],
                    int(query.get("offset", ["0"])[0] or 0),
                    int(query.get("limit", ["40"])[0] or 40),
                )
                self._send(payload)
            except FileNotFoundError as exc:
                self._send({"ok": False, "error": str(exc)}, 404)
            except (ValueError, TypeError) as exc:
                self._send({"ok": False, "error": str(exc)}, 422)
            except Exception as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        if parsed.path == "/api/attachment-member":
            query = parse_qs(parsed.query)
            try:
                payload, filename, content_type, source = self._store().attachment_member_bytes(
                    query.get("mail_id", [""])[0],
                    query.get("token", [""])[0],
                    query.get("name", [""])[0],
                    query.get("member", [""])[0],
                )
                self._send_binary(payload, filename, content_type, source, download=True)
            except FileNotFoundError as exc:
                self._send({"ok": False, "error": str(exc)}, 404)
            except (ValueError, TypeError) as exc:
                self._send({"ok": False, "error": str(exc)}, 422)
            except Exception as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        if parsed.path == "/api/attachment":
            query = parse_qs(parsed.query)
            mail_id = query.get("mail_id", [""])[0]
            name = query.get("name", [""])[0]
            token = query.get("token", [""])[0]
            try:
                payload, filename, content_type, source = self._store().attachment_bytes(mail_id, name, token)
                self._send_binary(
                    payload,
                    filename,
                    content_type,
                    source,
                    download=query.get("download", ["0"])[0] == "1",
                )
            except FileNotFoundError as exc:
                self._send({"ok": False, "error": str(exc)}, 404)
            except Exception as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        if parsed.path == "/api/health":
            self._send({"ok": True, "service": "workbench", "time": _now()})
            return
        if parsed.path == "/api/operations":
            operation_date = parse_qs(parsed.query).get("date", [""])[0]
            try:
                self._send(self._store().operations(operation_date))
            except Exception as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        if parsed.path == "/api/database":
            try:
                self._send(self._store().database_info())
            except Exception as exc:
                self._send({"ok": False, "error": str(exc)}, 500)
            return
        self._send({"ok": False, "error": "not found"}, 404)

    def do_POST(self) -> None:  # noqa: N802
        parsed = urlparse(self.path)
        length = int(self.headers.get("Content-Length", "0") or 0)
        try:
            payload = json.loads(self.rfile.read(length).decode("utf-8")) if length else {}
        except (ValueError, UnicodeDecodeError):
            self._send({"ok": False, "error": "请求不是有效 JSON"}, 400)
            return
        request_id = _text(self.headers.get("X-ECOPV-Request-ID"))
        if request_id and isinstance(payload, dict):
            payload["_request_id"] = request_id
        try:
            if parsed.path == "/api/action":
                self._send(self._store().action(payload))
            elif parsed.path == "/api/ingest":
                self._send(self._store().ingest(payload))
            elif parsed.path == "/api/export":
                date_start = _text(payload.get("date_start"))
                date_end = _text(payload.get("date_end"))
                out = self._store().export(date_start=date_start, date_end=date_end)
                date_label = "全部邮件日期"
                if date_start or date_end:
                    date_label = f"{date_start or '最早'} 至 {date_end or '最新'}"
                self._send({
                    "ok": True,
                    "path": str(out),
                    "message": f"已导出：{out.name}（邮件日期：{date_label}）",
                })
            else:
                self._send({"ok": False, "error": "not found"}, 404)
        except Exception as exc:
            self._send({"ok": False, "error": str(exc)}, 400)

    def log_message(self, format: str, *args: Any) -> None:
        return


class _ExclusiveThreadingHTTPServer(ThreadingHTTPServer):
    # Windows SO_REUSEADDR can otherwise bind two listeners to the same local port.
    allow_reuse_address = False

    def server_bind(self):
        if hasattr(socket, "SO_EXCLUSIVEADDRUSE"):
            self.socket.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
        super().server_bind()


class WorkbenchServer:
    def __init__(
        self,
        primary_path: Optional[str] = None,
        review_path: Optional[str] = None,
        filtered_path: Optional[str] = None,
        port: int = 8765,
        database_path: Optional[str] = None,
        test_mode: bool = False,
        host: str = "127.0.0.1",
        public_host: Optional[str] = None,
        run_controller=None,
    ):
        if run_controller is not None and host != "127.0.0.1":
            raise ValueError("网页运行控制台只允许绑定 127.0.0.1")
        self.run_controller = run_controller
        self.store = WorkbenchStore(
            primary_path,
            review_path,
            filtered_path=filtered_path,
            database_path=database_path,
            test_mode=test_mode,
        )
        self.port = 8765 if port is None else int(port)
        # 默认仍只监听本机；局域网部署时由启动参数显式传入 0.0.0.0。
        self.host = str(host or "127.0.0.1").strip() or "127.0.0.1"
        self.public_host = str(public_host or "").strip()
        self.httpd: Optional[ThreadingHTTPServer] = None
        self.thread: Optional[threading.Thread] = None

    def _url_host(self) -> str:
        if self.public_host:
            return self.public_host
        if self.host not in {"0.0.0.0", "::", ""}:
            return self.host
        # 仅用于启动提示和自动打开浏览器；实际监听仍使用 0.0.0.0。
        try:
            return socket.gethostbyname(socket.gethostname())
        except OSError:
            return "127.0.0.1"

    def start(self, open_browser: bool = True) -> str:
        if self.httpd:
            url = f"http://{self._url_host()}:{self.httpd.server_address[1]}/"
            if open_browser:
                webbrowser.open(url)
            return url
        handler_class = _Handler
        if self.run_controller is not None:
            from web_run_http import RunHandler
            handler_class = RunHandler
        try:
            self.httpd = _ExclusiveThreadingHTTPServer((self.host, self.port), handler_class)
        except OSError:
            # 端口被旧实例占用时仍允许 GUI 启动，但独立 file:// 入口需要使用
            # 新页面 URL，而不是直接双击 HTML 文件。
            self.httpd = _ExclusiveThreadingHTTPServer((self.host, 0), handler_class)
        self.httpd.workbench_store = self.store  # type: ignore[attr-defined]
        if self.run_controller is not None:
            self.run_controller.store = self.store
            self.httpd.run_controller = self.run_controller
        self.thread = threading.Thread(target=self.httpd.serve_forever, name="ecopv-workbench", daemon=True)
        self.thread.start()
        url = f"http://{self._url_host()}:{self.httpd.server_address[1]}/"
        if open_browser:
            webbrowser.open(url)
        return url

    def stop(self) -> None:
        if self.httpd:
            self.httpd.shutdown()
            self.httpd.server_close()
            self.store.wait_background_tasks()
            self.httpd = None
            self.thread = None
