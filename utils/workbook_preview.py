"""安全、按需读取邮件附件工作簿及 ZIP/RAR 内成员的只读预览。"""

from __future__ import annotations

import io
import base64
import json
import os
import stat
import tempfile
import zipfile
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

from utils.attachment_parser import (
    MAX_ARCHIVE_DEPTH,
    MAX_ARCHIVE_MEMBER_BYTES,
    MAX_ARCHIVE_TREE_BYTES,
    MAX_ARCHIVE_TREE_MEMBERS,
    _configure_rar_tool,
    _safe_archive_target,
    _validate_archive_members,
    _zip_member_name,
)


SUPPORTED_WORKBOOKS = {".xlsx", ".xls", ".xlsm"}
MAX_PREVIEW_ROWS = 100_000
MAX_PREVIEW_COLUMNS = 80
MAX_PAGE_SIZE = 60
MAX_CELL_CHARS = 4096
MAX_XLSX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_XLSX_MEMBERS = 10_000


class PreviewError(ValueError):
    """可直接显示给操作人员的预览错误。"""


def _safe_member_name(name: str) -> str:
    value = str(name or "").replace("\\", "/")
    if not value or _safe_archive_target(tempfile.gettempdir(), value) is None:
        raise PreviewError("压缩包成员路径不安全，已拒绝读取")
    return value


def _zip_is_symlink(info: zipfile.ZipInfo) -> bool:
    return stat.S_ISLNK(info.external_attr >> 16)


def _rar_is_symlink(info: Any) -> bool:
    flag = getattr(info, "is_symlink", False)
    try:
        return bool(flag() if callable(flag) else flag)
    except Exception:
        return False


def _member_token(chain: List[str]) -> str:
    payload = json.dumps(chain, ensure_ascii=False, separators=(",", ":")).encode("utf-8")
    return "chain1." + base64.urlsafe_b64encode(payload).decode("ascii").rstrip("=")


def _decode_member_token(value: str) -> Optional[List[str]]:
    if not str(value or "").startswith("chain1."):
        return None
    encoded = str(value).split(".", 1)[1]
    try:
        payload = base64.urlsafe_b64decode(encoded + "=" * (-len(encoded) % 4))
        chain = json.loads(payload.decode("utf-8"))
    except Exception as exc:
        raise PreviewError("压缩包成员定位信息无效") from exc
    if not isinstance(chain, list) or not chain or any(not isinstance(item, str) for item in chain):
        raise PreviewError("压缩包成员定位信息无效")
    return [_safe_member_name(item) for item in chain]


def _archive_manifest(
    path: Path,
    filename: str,
    *,
    depth: int = 0,
    chain: Optional[List[str]] = None,
    tree_budget: Optional[dict] = None,
) -> Dict[str, Any]:
    chain = list(chain or [])
    tree_budget = tree_budget if isinstance(tree_budget, dict) else {"members": 0, "bytes": 0}
    if depth > MAX_ARCHIVE_DEPTH:
        return {"filename": filename, "archive": True, "members": [], "workbook_count": 0,
                "member_count": 0, "notice": "嵌套压缩包层级超过预览安全上限。"}
    ext = Path(filename).suffix.lower()
    members: List[Dict[str, Any]] = []
    if ext == ".zip":
        try:
            with zipfile.ZipFile(path, "r") as archive:
                infos = archive.infolist()
                problem = _validate_archive_members(infos, tempfile.gettempdir(), _zip_member_name)
                if problem:
                    raise PreviewError(problem)
                for info in infos:
                    name = _zip_member_name(info)
                    if info.is_dir():
                        continue
                    unsafe_link = _zip_is_symlink(info)
                    suffix = Path(name).suffix.lower()
                    encrypted = bool(info.flag_bits & 0x1)
                    display_path = " → ".join(chain + [name])
                    members.append({
                        "path": display_path,
                        "member_token": _member_token(chain + [name]),
                        "filename": Path(name).name,
                        "size": int(info.file_size or 0),
                        "previewable": suffix in SUPPORTED_WORKBOOKS and not unsafe_link and not encrypted,
                        "downloadable": not unsafe_link and not encrypted,
                        "reason": (
                            "压缩包符号链接不允许读取" if unsafe_link else
                            "成员已加密，当前无法读取" if encrypted else
                            "此成员不是支持预览的工作簿" if suffix not in SUPPORTED_WORKBOOKS else ""
                        ),
                    })
        except PreviewError:
            raise
        except Exception as exc:
            raise PreviewError(f"ZIP 无法读取：{type(exc).__name__}: {exc}") from exc
    elif ext == ".rar":
        try:
            import rarfile
        except ImportError as exc:
            raise PreviewError("当前运行环境未安装 RAR 读取组件；原始 RAR 仍可下载") from exc
        try:
            _configure_rar_tool(rarfile)
            with rarfile.RarFile(str(path), "r") as archive:
                infos = archive.infolist()
                problem = _validate_archive_members(infos, tempfile.gettempdir())
                if problem:
                    raise PreviewError(problem)
                for info in infos:
                    name = _safe_member_name(getattr(info, "filename", ""))
                    if info.isdir():
                        continue
                    unsafe_link = _rar_is_symlink(info)
                    encrypted_check = getattr(info, "needs_password", False)
                    encrypted = bool(encrypted_check() if callable(encrypted_check) else encrypted_check)
                    suffix = Path(name).suffix.lower()
                    display_path = " → ".join(chain + [name])
                    members.append({
                        "path": display_path,
                        "member_token": _member_token(chain + [name]),
                        "filename": Path(name).name,
                        "size": int(getattr(info, "file_size", 0) or 0),
                        "previewable": suffix in SUPPORTED_WORKBOOKS and not unsafe_link and not encrypted,
                        "downloadable": not unsafe_link and not encrypted,
                        "reason": (
                            "压缩包链接成员不允许读取" if unsafe_link else
                            "成员已加密，当前无法读取" if encrypted else
                            "此成员不是支持预览的工作簿" if suffix not in SUPPORTED_WORKBOOKS else ""
                        ),
                    })
        except PreviewError:
            raise
        except Exception as exc:
            detail = str(exc).strip() or type(exc).__name__
            raise PreviewError(f"RAR 无法读取：{detail}；请确认已安装 7-Zip 或 UnRAR") from exc
    else:
        raise PreviewError("仅支持 ZIP 或 RAR 压缩包预览")
    tree_budget["members"] = int(tree_budget.get("members", 0)) + len(members)
    tree_budget["bytes"] = int(tree_budget.get("bytes", 0)) + sum(int(item.get("size", 0) or 0) for item in members)
    if tree_budget["members"] > MAX_ARCHIVE_TREE_MEMBERS or tree_budget["bytes"] > MAX_ARCHIVE_TREE_BYTES:
        raise PreviewError("压缩包及嵌套成员累计超过预览安全上限")

    # 嵌套 ZIP/RAR 也逐层列出；成员链 token 用于后续读取，展示路径只用于操作员识别。
    nested = [
        item for item in members
        if Path(item.get("filename", "")).suffix.lower() in {".zip", ".rar"}
        and item.get("downloadable")
    ]
    for item in nested:
        nested_name = str(item.get("filename") or "")
        nested_ext = Path(nested_name).suffix.lower()
        try:
            # The direct reader determines the container type from its filename;
            # passing only ".zip" makes pathlib treat it as a dotfile with no suffix.
            payload, _ = _read_archive_member_direct(
                path, filename, _decode_member_token(item["member_token"])[-1]
            )
            with tempfile.TemporaryDirectory() as temp_dir:
                nested_path = Path(temp_dir) / nested_name
                nested_path.write_bytes(payload)
                nested_manifest = _archive_manifest(
                    nested_path,
                    nested_name,
                    depth=depth + 1,
                    chain=_decode_member_token(item["member_token"]),
                    tree_budget=tree_budget,
                )
            members.extend(nested_manifest.get("members") or [])
        except PreviewError as exc:
            item["reason"] = str(exc)
        except Exception as exc:
            item["reason"] = f"嵌套压缩包无法读取：{type(exc).__name__}: {exc}"

    return {
        "filename": filename,
        "archive": True,
        "members": members,
        "workbook_count": sum(1 for item in members if item["previewable"]),
        "member_count": len(members),
        "notice": "工作簿逐份预览；不支持的、加密的或超出安全限制的成员会标出原因。",
    }


def _read_archive_member_direct(path: Path, archive_name: str, member_path: str) -> Tuple[bytes, str]:
    target = _safe_member_name(member_path)
    ext = Path(archive_name).suffix.lower()
    if ext == ".zip":
        try:
            with zipfile.ZipFile(path, "r") as archive:
                infos = archive.infolist()
                problem = _validate_archive_members(infos, tempfile.gettempdir(), _zip_member_name)
                if problem:
                    raise PreviewError(problem)
                info = next((item for item in infos if _zip_member_name(item) == target), None)
                if info is None or info.is_dir():
                    raise PreviewError("压缩包中找不到所选成员")
                if _zip_is_symlink(info):
                    raise PreviewError("压缩包符号链接不允许读取")
                if info.flag_bits & 0x1:
                    raise PreviewError("所选成员已加密，当前无法读取")
                if info.file_size > MAX_ARCHIVE_MEMBER_BYTES:
                    raise PreviewError("所选成员超过单文件安全上限")
                with archive.open(info, "r") as stream:
                    payload = stream.read(MAX_ARCHIVE_MEMBER_BYTES + 1)
                if len(payload) > MAX_ARCHIVE_MEMBER_BYTES:
                    raise PreviewError("所选成员解压后超过单文件安全上限")
                return payload, Path(target).name
        except PreviewError:
            raise
        except Exception as exc:
            raise PreviewError(f"ZIP 成员读取失败：{type(exc).__name__}: {exc}") from exc

    if ext == ".rar":
        try:
            import rarfile
            _configure_rar_tool(rarfile)
            with rarfile.RarFile(str(path), "r") as archive:
                infos = archive.infolist()
                problem = _validate_archive_members(infos, tempfile.gettempdir())
                if problem:
                    raise PreviewError(problem)
                info = next((item for item in infos if _safe_member_name(item.filename) == target), None)
                if info is None or info.isdir():
                    raise PreviewError("压缩包中找不到所选成员")
                if _rar_is_symlink(info):
                    raise PreviewError("压缩包链接成员不允许读取")
                encrypted_check = getattr(info, "needs_password", False)
                encrypted = bool(encrypted_check() if callable(encrypted_check) else encrypted_check)
                if encrypted:
                    raise PreviewError("所选成员已加密，当前无法读取")
                if info.file_size > MAX_ARCHIVE_MEMBER_BYTES:
                    raise PreviewError("所选成员超过单文件安全上限")
                with archive.open(info) as stream:
                    payload = stream.read(MAX_ARCHIVE_MEMBER_BYTES + 1)
                if len(payload) > MAX_ARCHIVE_MEMBER_BYTES:
                    raise PreviewError("所选成员解压后超过单文件安全上限")
                return payload, Path(target).name
        except PreviewError:
            raise
        except ImportError as exc:
            raise PreviewError("当前运行环境未安装 RAR 读取组件；原始 RAR 仍可下载") from exc
        except Exception as exc:
            detail = str(exc).strip() or type(exc).__name__
            raise PreviewError(f"RAR 成员读取失败：{detail}；请确认已安装 7-Zip 或 UnRAR") from exc
    raise PreviewError("仅支持从 ZIP 或 RAR 中读取成员")


def _read_archive_member(path: Path, archive_name: str, member_path: str) -> Tuple[bytes, str]:
    """读取直接或嵌套压缩包成员；旧版调用仍可传直接成员路径。"""
    chain = _decode_member_token(member_path)
    if chain is None:
        chain = [_safe_member_name(member_path)]
    current_source: Any = Path(path)
    current_archive_name = str(archive_name)
    payload = b""
    member_name = ""
    for index, member in enumerate(chain):
        current_ext = Path(current_archive_name).suffix.lower()
        if isinstance(current_source, Path):
            container_path = current_source
            if index > 0:
                raise PreviewError("嵌套压缩包成员定位失败")
            payload, member_name = _read_archive_member_direct(
                container_path, current_archive_name, member
            )
        else:
            with tempfile.TemporaryDirectory() as temp_dir:
                container_path = Path(temp_dir) / (f"nested{current_ext}" if current_ext else "nested.zip")
                container_path.write_bytes(current_source)
                payload, member_name = _read_archive_member_direct(
                    container_path, container_path.name, member
                )
        current_source = payload
        current_archive_name = member_name
    return payload, member_name


def _cell_text(value: Any) -> Tuple[str, bool]:
    if value is None:
        return "", False
    if hasattr(value, "isoformat") and callable(value.isoformat):
        try:
            text = value.isoformat(sep=" ")
        except TypeError:
            text = value.isoformat()
    else:
        text = str(value)
    if len(text) > MAX_CELL_CHARS:
        return text[:MAX_CELL_CHARS], True
    return text, False


def _validate_xlsx_container(source: Any) -> None:
    """拒绝工作簿内部的 ZIP 炸弹或异常庞大成员表。"""
    position = None
    if hasattr(source, "seek") and hasattr(source, "tell"):
        try:
            position = source.tell()
            source.seek(0)
        except Exception:
            position = None
    try:
        with zipfile.ZipFile(source, "r") as archive:
            infos = archive.infolist()
            if len(infos) > MAX_XLSX_MEMBERS:
                raise PreviewError("工作簿内部文件数量超过预览安全上限")
            total = 0
            for info in infos:
                total += int(info.file_size or 0)
                if total > MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise PreviewError("工作簿解压后超过 100 MiB 预览安全上限")
                if info.file_size and not info.compress_size:
                    raise PreviewError("工作簿内部存在异常压缩成员")
                if info.compress_size and info.file_size / info.compress_size > 10_000:
                    raise PreviewError("工作簿内部压缩比异常，已拒绝预览")
                if _zip_is_symlink(info):
                    raise PreviewError("工作簿内部包含链接成员，已拒绝预览")
    except PreviewError:
        raise
    except Exception as exc:
        raise PreviewError(f"工作簿容器无法读取：{type(exc).__name__}: {exc}") from exc
    finally:
        if position is not None:
            try:
                source.seek(position)
            except Exception:
                pass


def _xlsx_preview(source: Any, filename: str, offset: int, page_size: int) -> Dict[str, Any]:
    from openpyxl import load_workbook

    _validate_xlsx_container(source)
    workbook = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        sheets = []
        for worksheet in workbook.worksheets:
            try:
                dimension = str(worksheet.calculate_dimension())
            except Exception:
                dimension = ""
            bad_dimension = dimension.upper() in {"A1", "A1:A1"}
            if bad_dimension and hasattr(worksheet, "reset_dimensions"):
                worksheet.reset_dimensions()
            row_count = int(worksheet.max_row or 0)
            col_count = int(worksheet.max_column or 0)
            if bad_dimension or row_count <= 0 or col_count <= 0:
                # 有些导出的 XLSX 把维度错误写成 A1；有限扫描真实非空边界。
                row_count = 0
                col_count = 0
                for row_idx, values in enumerate(
                    worksheet.iter_rows(max_row=MAX_PREVIEW_ROWS + 1,
                                         max_col=MAX_PREVIEW_COLUMNS + 1,
                                         values_only=True), start=1
                ):
                    if any(value is not None for value in values):
                        row_count = row_idx
                        for col_idx, value in enumerate(values, start=1):
                            if value is not None:
                                col_count = max(col_count, col_idx)
                row_truncated = row_count > MAX_PREVIEW_ROWS
            else:
                row_truncated = row_count > MAX_PREVIEW_ROWS
            safe_rows = min(row_count, MAX_PREVIEW_ROWS)
            safe_columns = min(col_count, MAX_PREVIEW_COLUMNS)
            start_row = max(0, offset)
            end_row = min(safe_rows, start_row + page_size)
            rows = []
            truncated_cells = 0
            if start_row < end_row and safe_columns:
                for row_number, values in enumerate(
                    worksheet.iter_rows(min_row=start_row + 1, max_row=end_row,
                                         max_col=safe_columns, values_only=True),
                    start=start_row + 1,
                ):
                    cells = []
                    for value in values:
                        text, was_truncated = _cell_text(value)
                        cells.append(text)
                        truncated_cells += int(was_truncated)
                    rows.append({"row_number": row_number, "cells": cells})
            sheets.append({
                "sheet_name": worksheet.title,
                "row_count": row_count,
                "column_count": col_count,
                "offset": start_row,
                "next_offset": end_row if end_row < safe_rows else None,
                "rows": rows,
                "row_limit_reached": row_truncated,
                "column_limit_reached": col_count > MAX_PREVIEW_COLUMNS,
                "truncated_cells": truncated_cells,
                "complete": not row_truncated and col_count <= MAX_PREVIEW_COLUMNS,
            })
        return {"filename": filename, "archive": False, "sheets": sheets,
                "page_size": page_size, "max_rows": MAX_PREVIEW_ROWS,
                "max_columns": MAX_PREVIEW_COLUMNS}
    finally:
        workbook.close()


def _xls_preview(source: Any, filename: str, offset: int, page_size: int) -> Dict[str, Any]:
    try:
        import xlrd
    except ImportError as exc:
        raise PreviewError("当前运行环境未安装 XLS 读取组件") from exc
    if isinstance(source, (bytes, bytearray)):
        workbook = xlrd.open_workbook(file_contents=bytes(source), on_demand=True)
    else:
        workbook = xlrd.open_workbook(filename=str(source), on_demand=True)
    sheets = []
    try:
        for worksheet in workbook.sheets():
            row_count = int(worksheet.nrows)
            col_count = int(worksheet.ncols)
            safe_rows = min(row_count, MAX_PREVIEW_ROWS)
            safe_columns = min(col_count, MAX_PREVIEW_COLUMNS)
            start_row = max(0, offset)
            end_row = min(safe_rows, start_row + page_size)
            rows = []
            truncated_cells = 0
            for row_number in range(start_row, end_row):
                values = worksheet.row_values(row_number, end_colx=safe_columns)
                cells = []
                for value in values:
                    text, was_truncated = _cell_text(value)
                    cells.append(text)
                    truncated_cells += int(was_truncated)
                rows.append({"row_number": row_number + 1, "cells": cells})
            sheets.append({
                "sheet_name": worksheet.name,
                "row_count": row_count,
                "column_count": col_count,
                "offset": start_row,
                "next_offset": end_row if end_row < safe_rows else None,
                "rows": rows,
                "row_limit_reached": row_count > MAX_PREVIEW_ROWS,
                "column_limit_reached": col_count > MAX_PREVIEW_COLUMNS,
                "truncated_cells": truncated_cells,
                "complete": row_count <= MAX_PREVIEW_ROWS and col_count <= MAX_PREVIEW_COLUMNS,
            })
    finally:
        workbook.release_resources()
    return {"filename": filename, "archive": False, "sheets": sheets,
            "page_size": page_size, "max_rows": MAX_PREVIEW_ROWS,
            "max_columns": MAX_PREVIEW_COLUMNS}


def _workbook_preview(source: Any, filename: str, offset: int, page_size: int) -> Dict[str, Any]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_WORKBOOKS:
        raise PreviewError("只支持 XLS/XLSX/XLSM 工作簿预览")
    try:
        if suffix == ".xls":
            return _xls_preview(source, filename, offset, page_size)
        return _xlsx_preview(source, filename, offset, page_size)
    except PreviewError:
        raise
    except Exception as exc:
        raise PreviewError(f"工作簿无法读取：{type(exc).__name__}: {exc}") from exc


def preview_attachment(path: Path, filename: str, member_path: str = "",
                       offset: int = 0, page_size: int = 40) -> Dict[str, Any]:
    """预览工作簿或返回压缩包清单；表格按物理行分页读取。"""
    path = Path(path)
    if not path.is_file():
        raise PreviewError("原始附件文件不存在")
    suffix = Path(filename).suffix.lower()
    safe_page_size = max(1, min(int(page_size or 40), MAX_PAGE_SIZE))
    safe_offset = max(0, min(int(offset or 0), MAX_PREVIEW_ROWS))
    if suffix in {".zip", ".rar"}:
        if not member_path:
            return _archive_manifest(path, filename)
        payload, inner_name = _read_archive_member(path, filename, member_path)
        if Path(inner_name).suffix.lower() == ".xls":
            result = _workbook_preview(payload, inner_name, safe_offset, safe_page_size)
        else:
            result = _workbook_preview(io.BytesIO(payload), inner_name, safe_offset, safe_page_size)
        result["archive_filename"] = filename
        result["member_path"] = member_path
        return result
    if path.stat().st_size > 50 * 1024 * 1024:
        raise PreviewError("工作簿超过 50 MiB 的在线预览安全上限；请下载原件查看")
    return _workbook_preview(path, filename, safe_offset, safe_page_size)


def download_archive_member(path: Path, archive_name: str, member_path: str) -> Tuple[bytes, str]:
    """下载一个经过成员清单验证的压缩包内原始文件。"""
    return _read_archive_member(Path(path), archive_name, member_path)
