"""Local web run console. Reuses the desktop worker; no alternate parsing pipeline."""
from __future__ import annotations

import copy
import io
import json
import os
import re
import secrets
import threading
import zipfile
from collections import deque
from datetime import date, datetime, timedelta
from pathlib import Path

import yaml

BUILD = "2026.09.30-r11"
UPLOAD_LIMIT = 20 * 1024 * 1024
ROLES = {
    "agent": ("agent_table_path", "agent_emails", "data/agent_emails.xlsx"),
    "internal": ("internal_table_path", "internal_emails", "data/internal_email_cache.xlsx"),
    "project": ("project_table_path", "project_names", "data/project_names.xlsx"),
    "stage2": ("stage2_input_path", "", ""),
}


def atomic_text(path: Path, text: str):
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    temporary.write_text(text, encoding="utf-8")
    os.replace(temporary, path)


def load_json(path):
    try:
        value = json.loads(Path(path).read_text(encoding="utf-8"))
        return value if isinstance(value, dict) else {}
    except (OSError, ValueError):
        return {}


def validate_workbook(data: bytes, role: str):
    """Validate before retaining an uploaded file; never execute macros/formulas."""
    from openpyxl import load_workbook
    if role not in ROLES or not data or len(data) > UPLOAD_LIMIT:
        raise ValueError("请选择不超过 20 MB 的 .xlsx 表格")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > 2000 or sum(x.file_size for x in members) > 100 * 1024 * 1024:
                raise ValueError("表格解压后过大，请拆分参考表")
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if role == "stage2":
                if "工单待查" not in workbook.sheetnames:
                    raise ValueError("阶段二文件必须包含“工单待查”工作表")
                sheet = workbook["工单待查"]
                headers = {str(x or "").strip() for x in next(sheet.iter_rows(max_row=1, values_only=True), ())}
                if not {"代理", "客户公司名称", "标准化项目名称"}.issubset(headers):
                    raise ValueError("阶段二表缺少代理、客户公司名称或标准化项目名称列")
            else:
                sheet = workbook.active
            # Same positional conventions as the desktop loaders. Bound scanning of malformed inputs.
            rows = sheet.iter_rows(min_row=2, max_row=10001, max_col=32, values_only=True)
            for row in rows:
                if role == "agent" and row[0] and "@" in str(row[2] or ""):
                    return
                if role == "internal" and any("@" in str(x or "") for x in row):
                    return
                if role == "project" and row[0] and row[1]:
                    return
                if role == "stage2" and any(x is not None for x in row):
                    return
            raise ValueError("未找到有效数据行，请核对表头及参考表格式")
        finally:
            workbook.close()
    except (zipfile.BadZipFile, KeyError, OSError) as exc:
        raise ValueError("文件不是可读取的 Excel 工作簿") from exc


class DesktopWorkerAdapter:
    """Qt signals are delivered directly; completion waits for actual thread exit."""
    def __init__(self, kwargs, callbacks):
        from gui import WorkerThread
        from PyQt5.QtCore import Qt
        self.worker = WorkerThread(**kwargs)
        self.callbacks = callbacks
        for signal, name in (("log_signal", "log"), ("progress_signal", "progress"),
                             ("finished_signal", "result"), ("error_signal", "error"),
                             ("cancelled_signal", "cancelled")):
            getattr(self.worker, signal).connect(callbacks[name], type=Qt.DirectConnection)

    def start(self):
        self.worker.start()
        def wait_for_exit():
            self.worker.wait()
            self.callbacks["done"]()
        self.joiner = threading.Thread(target=wait_for_exit, name="web-run-completion", daemon=False)
        self.joiner.start()

    def stop(self):
        self.worker.stop("网页运行控制")


class RunController:
    def __init__(self, root, worker_factory=DesktopWorkerAdapter):
        self.root = Path(root).resolve()
        self.config_path = self.root / "config.yaml"
        self.session_path = self.root / "storage/web_run_session.json"
        self.lock = threading.RLock()
        self.worker_factory = worker_factory
        self.worker = None
        self.store = None
        self.csrf = secrets.token_urlsafe(32)
        self.busy = False
        self.status = "idle"
        self.progress = [0, 0]
        self.message = "就绪，可以选择日期开始解析"
        self.logs = deque(maxlen=2000)
        self.sequence = 0
        self.results = []
        self.files = {}
        self.cache_entries = {}
        self.config = self._read_config()
        previous = load_json(self.root / "session_state.json")
        previous.update(load_json(self.session_path))
        self.session = {
            "mode": previous.get("mode", "stage1"),
            "date_from": previous.get("date_from") or (date.today() - timedelta(days=30)).isoformat(),
            "date_to": previous.get("date_to") or date.today().isoformat(),
            "force_live_query": bool(previous.get("force_live_query", False)),
        }
        self.imports = {}
        for role, (session_key, config_key, fallback) in ROLES.items():
            configured = self.config.get("reference_tables", {}).get(config_key)
            value = previous.get(session_key) or configured or fallback
            resolved = self._existing_path(value)
            if resolved is None and fallback:
                resolved = self._existing_path(fallback)
            if resolved is not None:
                self.imports[role] = resolved
        self.catalog = {}
        self.refresh_catalog()

    def _read_config(self):
        if not self.config_path.is_file():
            raise ValueError("找不到 config.yaml，请在原程序目录启动网页版")
        config = yaml.safe_load(self.config_path.read_text(encoding="utf-8"))
        if not isinstance(config, dict):
            raise ValueError("config.yaml 格式错误")
        return config

    def _existing_path(self, value):
        if not value:
            return None
        path = Path(value)
        if not path.is_absolute():
            path = self.root / path
        return path.resolve() if path.is_file() else None

    def _relative(self, path):
        try:
            return str(Path(path).relative_to(self.root))
        except ValueError:
            return str(path)

    def _save_session(self):
        saved = dict(self.session)
        saved.update({ROLES[role][0]: self._relative(path) for role, path in self.imports.items()})
        atomic_text(self.session_path, json.dumps(saved, ensure_ascii=False, indent=2))

    def _idle(self):
        if getattr(self, "closing", False):
            raise ValueError("服务正在退出，请重新启动网页版")
        if self.busy:
            raise ValueError("任务正在运行，请等待结束或先点击停止")

    def _redact(self, text):
        value = str(text)
        for section in self.config.values():
            if not isinstance(section, dict):
                continue
            for key, secret in section.items():
                if re.search(r"password|secret|api_key|token", str(key), re.I) and isinstance(secret, str) and secret:
                    value = value.replace(secret, "[已隐藏]")
        for key in ("DEEPSEEK_API_KEY", "OPENAI_API_KEY"):
            if os.environ.get(key):
                value = value.replace(os.environ[key], "[已隐藏]")
        return value[:16000]

    def log(self, value):
        with self.lock:
            self.sequence += 1
            self.logs.append({"seq": self.sequence, "text": self._redact(value)})

    def snapshot(self, after=0):
        with self.lock:
            email = self.config.get("email", {})
            workorder = self.config.get("workorder", {})
            logs = [x for x in self.logs if x["seq"] > after][:300]
            return {"ok": True, "build": BUILD, "csrf": self.csrf, "busy": self.busy,
                    "status": self.status, "message": self.message, "progress": list(self.progress),
                    "session": dict(self.session), "cursor": logs[-1]["seq"] if logs else self.sequence,
                    "logs": logs,
                    "credentials": {"email_address": email.get("address", ""),
                                    "email_password_set": bool(email.get("password")),
                                    "workorder_username": workorder.get("username", ""),
                                    "workorder_password_set": bool(workorder.get("password"))},
                    "imports": {key: path.name for key, path in self.imports.items()},
                    "results": list(self.results), "catalog": copy.deepcopy(self.catalog)}

    def _file_item(self, path):
        path = Path(path).resolve()
        token = next((key for key, value in self.files.items() if value == path), None)
        if token is None:
            token = secrets.token_urlsafe(18)
            self.files[token] = path
        return {"token": token, "name": path.name, "label": self._relative(path),
                "url": "/api/run/download?token=" + token}

    def download(self, token):
        with self.lock:
            path = self.files.get(token)
            if path is None or not path.is_file() or path.suffix.lower() != ".xlsx":
                raise ValueError("下载文件不存在，请刷新结果列表")
            # Only files explicitly registered as output/stage-two workbooks can be downloaded.
            return path

    def refresh_catalog(self):
        with self.lock:
            self._idle()
            output = Path(self.config.get("output", {}).get("dir", "output"))
            if not output.is_absolute():
                output = self.root / output
            candidates = list(output.glob("workbench_reviewed*.xlsx"))
            candidates += list(output.glob("**/workbench_reviewed*.xlsx"))
            candidates += list(output.glob("**/to_workorder_list.xlsx"))
            if "stage2" in self.imports:
                candidates.append(self.imports["stage2"])
            candidates = sorted(set(candidates), key=lambda p: p.stat().st_mtime, reverse=True)[:100]
            caches = []
            self.cache_entries = {}
            for folder in sorted((self.root / "output/stage1_email/cache").glob("*"), reverse=True):
                if not folder.is_dir():
                    continue
                primary, review = folder / "to_workorder_list.xlsx", folder / "to_review_list.xlsx"
                if not primary.is_file() and not review.is_file():
                    continue
                token = secrets.token_urlsafe(18)
                self.cache_entries[token] = folder
                caches.append({"token": token, "name": folder.name})
            self.catalog = {"stage2": [self._file_item(p) for p in candidates], "caches": caches}
            return {"ok": True, "catalog": copy.deepcopy(self.catalog)}

    def save_settings(self, payload):
        with self.lock:
            self._idle()
            config = self._read_config()
            for key, section, field in (("email_address", "email", "address"),
                                        ("email_password", "email", "password"),
                                        ("workorder_username", "workorder", "username"),
                                        ("workorder_password", "workorder", "password")):
                value = payload.get(key)
                if value is None:
                    continue
                if not isinstance(value, str) or len(value) > 1024:
                    raise ValueError("账号配置格式不正确")
                if field == "password" and not value:
                    continue  # Blank means retain, never silently clear an existing password.
                config.setdefault(section, {})[field] = value.strip() if field != "password" else value
            atomic_text(self.config_path, yaml.safe_dump(config, allow_unicode=True, sort_keys=False))
            self.config = config
            return {"ok": True, "message": "账号配置已保存；未填写的密码保持不变"}

    def upload(self, role, name, data):
        with self.lock:
            self._idle()
            if Path(name).suffix.lower() != ".xlsx":
                raise ValueError("请上传 .xlsx 文件；旧版 .xls 请先另存为 .xlsx")
            validate_workbook(data, role)
            safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name)[-100:]
            target = self.root / "storage/web_imports" / (secrets.token_hex(8) + "_" + safe_name)
            target.parent.mkdir(parents=True, exist_ok=True)
            target.write_bytes(data)
            if role == "internal":
                from gui import _cache_internal_email_table
                # Match desktop import normalization, but only rewrite the uploaded copy.
                _cache_internal_email_table(str(target), cache_path=str(target))
            old_imports = dict(self.imports)
            self.imports[role] = target
            try:
                self._save_session()
            except Exception:
                self.imports = old_imports
                raise
            self.refresh_catalog()
            return {"ok": True, "message": "文件已校验并导入", "name": target.name,
                    "file": self._file_item(target) if role == "stage2" else None}

    def use_cache(self, token):
        with self.lock:
            self._idle()
            folder = self.cache_entries.get(token)
            if folder is None:
                raise ValueError("缓存选择已失效，请刷新列表后重选")
            files = [folder / name for name in ("to_workorder_list.xlsx", "to_review_list.xlsx", "filtered_mail_record.xlsx")]
            if not any(p.is_file() for p in files[:2]):
                raise ValueError("缓存文件已不存在")
            self._bind_results(files[0], files[1], "stage1")
            match = re.fullmatch(r"(\d{4}-\d{2}-\d{2})_to_(\d{4}-\d{2}-\d{2})", folder.name)
            if match:
                self.session.update(date_from=match[1], date_to=match[2])
            self._save_session()
            return {"ok": True, "message": "已选择缓存；复核页仍可查看历史邮件，请按日期筛选"}

    def _bind_results(self, primary, secondary, mode):
        paths = [Path(p) for p in (primary, secondary) if p and Path(p).is_file()]
        self.results = [self._file_item(p) for p in paths]
        if self.store is not None:
            with self.store._lock:
                if mode == "stage1":
                    self.store.primary_path = Path(primary)
                    self.store.review_path = Path(secondary)
                    self.store.filtered_path = Path(primary).parent / "filtered_mail_record.xlsx"
                else:
                    self.store.workorder_result_path = Path(primary)

    def _runtime_config(self):
        config = copy.deepcopy(self.config)
        # Same path keys/defaults as the desktop GUI; no business configuration is replaced.
        defaults = {("output", "dir"): "output", ("logging", "dir"): "logs",
                    ("email", "cache_dir"): "cache/mails",
                    ("workorder", "query_cache_path"): "storage/query_cache.json",
                    ("workorder", "login_state_path"): "storage/login_state.json",
                    ("llm", "audit_log_path"): "logs/llm_calls.jsonl"}
        for (section, field), fallback in defaults.items():
            value = Path(config.setdefault(section, {}).get(field) or fallback)
            config[section][field] = str(value if value.is_absolute() else self.root / value)
        config.setdefault("reference_tables", {})
        for role, (_, key, fallback) in ROLES.items():
            if key:
                config["reference_tables"][key] = str(self.imports.get(role) or self.root / fallback)
        return config

    def start(self, payload):
        with self.lock:
            self._idle()
            mode = payload.get("mode", "stage1")
            if mode not in {"stage1", "stage2", "all"}:
                raise ValueError("运行模式无效")
            start = end = None
            if mode != "stage2":
                try:
                    start = datetime.strptime(payload.get("date_from", ""), "%Y-%m-%d")
                    end = datetime.strptime(payload.get("date_to", ""), "%Y-%m-%d")
                except (ValueError, TypeError):
                    raise ValueError("请选择有效的开始和结束日期") from None
                if start > end:
                    raise ValueError("开始日期不能晚于结束日期")
                if any(role not in self.imports for role in ("agent", "internal", "project")):
                    raise ValueError("请先导入代理邮箱表、内部邮箱表和项目名称表")
            stage2 = None
            if mode == "stage2":
                stage2 = self.files.get(payload.get("input_token"))
                if stage2 is None:
                    raise ValueError("请选择或上传阶段二工单待查文件")
                if stage2.stat().st_size > UPLOAD_LIMIT:
                    raise ValueError("待查询文件超过 20 MB，请拆分后查询")
                validate_workbook(stage2.read_bytes(), "stage2")
            config = self._runtime_config()
            required = ["email"] if mode == "stage1" else ["workorder"] if mode == "stage2" else ["email", "workorder"]
            for section in required:
                account = "address" if section == "email" else "username"
                if not config.get(section, {}).get(account) or not config.get(section, {}).get("password"):
                    raise ValueError("请先填写并保存邮箱/工单账号及密码")
            config.setdefault("workorder", {})["force_live_query"] = bool(payload.get("force_live_query", False))
            self.session.update(mode=mode, force_live_query=config["workorder"]["force_live_query"])
            if start:
                self.session.update(date_from=start.date().isoformat(), date_to=end.date().isoformat())
            if stage2:
                self.imports["stage2"] = stage2
            self._save_session()
            self.busy, self.status, self.progress = True, "running", [0, 0]
            self.message = "任务已启动，正在准备运行"
            self.results = []
            self.log("开始运行：" + mode + (f"，邮件日期 {start:%Y-%m-%d} 至 {end:%Y-%m-%d}" if start else ""))
            kwargs = dict(config=config, mode=mode, date_from=start, date_to=end,
                          agent_email_path=str(self.imports.get("agent", "")),
                          internal_email_path=str(self.imports.get("internal", "")),
                          project_table_path=str(self.imports.get("project", "")),
                          prefer_internal_email_path=True,
                          stage2_input_path=str(stage2) if stage2 else None)
            callbacks = {"log": self.log, "progress": self._progress, "error": self._error,
                         "result": lambda a, b: self._result(a, b, mode),
                         "cancelled": self._cancelled, "done": self._done}
            try:
                self.worker = self.worker_factory(kwargs, callbacks)
                self.worker.start()
            except Exception as exc:
                self.busy = False
                self._error(exc)
                raise ValueError(self.message) from None
            return {"ok": True, "message": "任务已启动"}

    def _progress(self, current, total):
        with self.lock:
            self.progress = [current, total]

    def _error(self, error):
        with self.lock:
            self.status, self.message = "error", self._redact(error)
            self.log(self.message)

    def _cancelled(self, message):
        with self.lock:
            self.status, self.message = "stopped", self._redact(message)
            self.log(self.message)

    def _result(self, primary, secondary, mode):
        with self.lock:
            try:
                self._bind_results(primary, secondary, mode)
                self.status, self.message = "finishing", "结果已生成，正在收尾"
            except Exception as exc:
                self._error(exc)

    def _done(self):
        with self.lock:
            if self.status == "finishing":
                self.status, self.message = "completed", "运行完成，可下载结果或进入人工复核"
            elif self.status not in {"error", "stopped"}:
                self.status, self.message = "error", "任务已退出但未生成结果，请查看运行日志"
            self.busy = False
            self.log(self.message)
            try:
                self.refresh_catalog()
            except Exception as exc:
                self.log("结果列表刷新失败：" + str(exc))

    def stop(self):
        with self.lock:
            if self.busy and self.worker:
                self.worker.stop()
                self.status, self.message = "stopping", "已请求停止，等待当前安全步骤结束；不会强制中断写入"
            return {"ok": True, "message": self.message}
